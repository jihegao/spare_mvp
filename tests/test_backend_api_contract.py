from __future__ import annotations

import copy
import hashlib
import http.client
import json
import sqlite3
import sys
import tempfile
import threading
import time
from pathlib import Path
from typing import Any
import unittest
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.spare_mvp_backend.api import (
    BackendApi,
    BackendApiError,
    _lite_mesa_carry_list_result,
    _lite_mesa_downtime_event_snapshots,
    _lite_mesa_downtime_factors_result,
    _lite_mesa_mission_reliability_result,
    _lite_mesa_spare_shortfall_result,
    _normalize_lite_mesa_analysis_settings,
    _run_aircraft_support_v1_analysis_sample,
    _run_lite_mesa_analysis_sample_worker,
)
from src.spare_mvp_backend.http_server import create_backend_server
from src.spare_mvp_backend.modeling_import import modeling_import_to_project, validate_modeling_import_package
from src.spare_mvp_backend.project_payload import export_project_json, strip_project_sweep
from src.spare_mvp_backend.repository import ContractRepository, initialize_database
from src.spare_mvp_backend.run_service import RunService, RunServiceError
from src.spare_mvp_contract.adapter import AdapterError, SimulationAdapter
from src.spare_mvp_contract.downtime import format_simulation_minute
from src.spare_mvp_contract.task_reliability import (
    build_task_reliability_result_fields,
    format_reliability_percent,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
M7_MONTE_CARLO_ARTIFACT_KINDS = {
    "run_config",
    "input_project",
    "compiled_scenario",
    "sample_results",
    "aggregate_result",
    "result_summary",
    "metrics",
    "report",
    "log",
    "monte_carlo_base",
    "analysis_projection_spare_shortfall",
    "analysis_projection_carry_list",
    "analysis_projection_mission_reliability",
    "analysis_projection_downtime_factors",
    "visualization_state_series",
}


def small_aircraft_support_project(project_id: str) -> dict[str, Any]:
    return {
        "schema_version": "project-v0",
        "project_id": project_id,
        "project_version": "project-v0.1",
        "scenarioId": "aircraft-support-contract-demo",
        "activeModule": "sparePlanning",
        "projectInfo": {"name": "small current project", "baseCode": "SM", "summary": "small current project"},
        "airports": ["A"],
        "missionProfile": {"name": "small current mission", "durationHours": 1, "compositeTasks": [], "periodicTasks": []},
        "experiment": {"seed": 42},
        "basicMissions": [{
            "id": "basic-small",
            "name": "small sortie",
            "missionId": "basic-small",
            "minRequiredSorties": 1,
            "taskDurationMinutes": 30,
            "equipmentType": "J-15",
            "missionPhases": [
                {"id": "phase-sortie", "name": "sortie", "sequence": 1, "durationMinutes": 30}
            ],
        }],
        "combatUnit": {"members": [{"aircraftNo": "J15-001", "model": "J-15", "status": "ready", "airport": "A"}]},
        "products": [{"id": "product-whole-aircraft", "name": "whole aircraft", "model": "whole-aircraft", "kind": "whole"}],
        "components": [{
            "id": "whole-aircraft",
            "name": "whole aircraft",
            "productId": "product-whole-aircraft",
            "aircraftModel": "J-15",
            "productType": "whole",
            "quantity": 1,
            "failureRate": 0.01,
            "mtbfHours": 100,
            "meanRepairTimeMinutes": 30,
            "failureDistribution": {"distributionType": "exponential", "parameters": "lambda=0.01"},
            "repairDistribution": {"distributionType": "fixed", "parameters": "value=30"},
        }],
        "supportNodes": [{
            "id": "node-a",
            "name": "node A",
            "personnelCapacity": 1,
            "equipmentCapacity": 1,
            "inventory": {"product-whole-aircraft": 2},
        }],
        "supportActivities": [{"id": "corrective", "activityType": "corrective", "durationHours": 1, "jobs": []}],
        "supportOrganization": {},
        "reliabilityBlockDiagram": {
            "nodes": [{"id": "whole-aircraft", "name": "whole aircraft", "type": "system", "failureRate": 0.01}],
            "edges": [],
        },
        "modelingImportValidation": {"usedTables": {}, "disabledDomains": [], "warnings": []},
    }


def periodic_three_day_aircraft_support_project(project_id: str) -> dict[str, Any]:
    project = small_aircraft_support_project(project_id)
    project["missionProfile"]["durationHours"] = 24
    project["missionProfile"]["compositeTasks"] = [
        {
            "id": "composite-a",
            "name": "three-day composite",
            "taskItems": [
                {
                    "id": "task-a",
                    "basicMissionId": "basic-small",
                    "basicTaskName": "small sortie",
                    "firstWaveTime": "08:00",
                    "taskDurationMinutes": 30,
                    "equipmentQuantity": 1,
                    "equipmentType": "J-15",
                }
            ],
        }
    ]
    project["missionProfile"]["periodicTasks"] = [
        {
            "id": "periodic-three-day",
            "name": "three-day explicit rows",
            "periodDays": 7,
            "repeatCount": 1,
            "compositeTaskIds": ["composite-a"],
            "weekdayAssignments": {
                "monday": "composite-a",
                "tuesday": "composite-a",
                "sunday": "composite-a",
            },
            "mondayCompositeTaskId": "composite-a",
            "tuesdayCompositeTaskId": "composite-a",
            "sundayCompositeTaskId": "composite-a",
            "compositeTasks": [
                {"weekIndex": 1, "weekday": "mondayCompositeTaskId", "compositeTaskId": "composite-a"},
                {"weekIndex": 1, "weekday": "tuesdayCompositeTaskId", "compositeTaskId": "composite-a"},
                {"weekIndex": 1, "weekday": "wednesdayCompositeTaskId", "compositeTaskId": "composite-a"},
                {"weekIndex": 1, "weekday": "thursdayCompositeTaskId", "compositeTaskId": ""},
                {"weekIndex": 1, "weekday": "fridayCompositeTaskId", "compositeTaskId": ""},
                {"weekIndex": 1, "weekday": "saturdayCompositeTaskId", "compositeTaskId": ""},
                {"weekIndex": 1, "weekday": "sundayCompositeTaskId", "compositeTaskId": ""},
            ],
        }
    ]
    return project


class RecordingAdapter(SimulationAdapter):
    def __init__(self) -> None:
        super().__init__(REPO_ROOT)
        self.compile_calls: list[tuple[dict, str]] = []
        self.compile_runtime_configs: list[dict | None] = []
        self.run_calls: list[tuple[dict, int]] = []
        self.monte_carlo_run_calls: list[dict] = []

    def compile_scenario(self, project: dict, model_family: str = "aircraft_support_v1", runtime_config: dict | None = None) -> dict:
        self.compile_calls.append((copy.deepcopy(project), model_family))
        self.compile_runtime_configs.append(copy.deepcopy(runtime_config))
        return super().compile_scenario(project, model_family=model_family, runtime_config=runtime_config)

    def compile_scenario_with_gate(self, project: dict, model_family: str = "aircraft_support_v1", runtime_config: dict | None = None) -> dict:
        self.compile_calls.append((copy.deepcopy(project), model_family))
        self.compile_runtime_configs.append(copy.deepcopy(runtime_config))
        return super().compile_scenario_with_gate(project, model_family=model_family, runtime_config=runtime_config)

    def run_scenario(
        self,
        scenario: dict,
        output_dir: Path | str,
        steps: int = 3,
        run_id: str | None = None,
    ) -> dict[str, dict]:
        self.run_calls.append((copy.deepcopy(scenario), steps, run_id))
        return super().run_scenario(scenario, output_dir=output_dir, steps=steps, run_id=run_id)

    def run_monte_carlo_scenario(
        self,
        scenario: dict,
        output_dir: Path | str,
        steps: int = 3,
        run_id: str | None = None,
        **kwargs,
    ) -> dict[str, dict]:
        self.monte_carlo_run_calls.append(
            {
                "scenario": copy.deepcopy(scenario),
                "steps": steps,
                "run_id": run_id,
                "kwargs": copy.deepcopy(kwargs),
            }
        )
        return super().run_monte_carlo_scenario(
            scenario,
            output_dir=output_dir,
            steps=steps,
            run_id=run_id,
            **kwargs,
        )


class FailingRunAdapter(RecordingAdapter):
    def run_scenario(
        self,
        scenario: dict,
        output_dir: Path | str,
        steps: int = 3,
        run_id: str | None = None,
    ) -> dict[str, dict]:
        self.run_calls.append((copy.deepcopy(scenario), steps, run_id))
        raise AdapterError("executor_failed", "synthetic executor failure", run_id=run_id)


class BackendApiContractTest(unittest.TestCase):
    def setUp(self) -> None:
        self.connection = sqlite3.connect(":memory:")
        initialize_database(self.connection)
        self.repository = ContractRepository(self.connection)
        self.adapter = RecordingAdapter()
        self.tempdir = tempfile.TemporaryDirectory()
        self.api = BackendApi(self.repository, self.adapter, output_dir=Path(self.tempdir.name))

    def tearDown(self) -> None:
        self.connection.close()
        self.tempdir.cleanup()

    def test_lite_mesa_analysis_defaults_to_four_samples(self) -> None:
        settings = _normalize_lite_mesa_analysis_settings({})
        self.assertEqual(settings["samples"], 4)
        self.assertEqual(settings["parallelCores"], 1)
        self.assertEqual(settings["sampleTimeoutSeconds"], 60)
        self.assertEqual(settings["sessionTimeoutSeconds"], 300)
        self.assertNotIn("maxTimeWindow", _normalize_lite_mesa_analysis_settings({"maxTimeWindow": 1}))

    def test_save_project_accepts_committed_case_large_after_legacy_export_normalization(self) -> None:
        legacy = json.loads((REPO_ROOT / "exports" / "project-case-large.json").read_text(encoding="utf-8"))
        clean = export_project_json(legacy, target="aircraft_support_v1")

        saved = self.api.save_project(clean)

        self.assertEqual(saved["project_id"], legacy["project_id"])

    def test_lite_mesa_session_budget_scales_with_execution_waves_and_stays_bounded(self) -> None:
        parallel = _normalize_lite_mesa_analysis_settings({"samples": 24, "parallelCores": 4})
        serial = _normalize_lite_mesa_analysis_settings({"samples": 24, "parallelCores": 1})
        very_large = _normalize_lite_mesa_analysis_settings({"samples": 1000, "parallelCores": 1})

        self.assertEqual(parallel["sessionTimeoutSeconds"], 420)
        self.assertEqual(serial["sessionTimeoutSeconds"], 900)
        self.assertEqual(very_large["sessionTimeoutSeconds"], 900)

    def test_lite_mesa_sample_worker_reports_hard_timeout_diagnostics(self) -> None:
        def slow_sample(*_args: Any, **_kwargs: Any) -> dict[str, Any]:
            time.sleep(0.05)
            return {"metrics": {}}

        with mock.patch(
            "src.spare_mvp_backend.api._run_aircraft_support_v1_analysis_sample",
            side_effect=slow_sample,
        ):
            outcome = _run_lite_mesa_analysis_sample_worker(({}, 20260718, 3, False, 0.01))

        self.assertEqual(outcome["status"], "failed")
        self.assertEqual(outcome["failure"]["error"]["code"], "sample_timeout")
        self.assertEqual(outcome["failure"]["error"]["details"]["phase"], "model_execution")
        self.assertEqual(outcome["failure"]["error"]["details"]["timeout_seconds"], 0.01)

    def test_lite_mesa_all_timeout_response_keeps_empty_moments_and_execution_counts(self) -> None:
        failures = [
            {
                "sample_index": index,
                "seed": 20260718 + index,
                "error": {"code": "sample_timeout", "message": "synthetic timeout", "details": {}},
            }
            for index in range(2)
        ]
        diagnostics = [
            {"sample_index": index, "seed": 20260718 + index, "status": "failed", "elapsed_seconds": 60.0}
            for index in range(2)
        ]
        with mock.patch(
            "src.spare_mvp_backend.api._run_lite_mesa_analysis_samples",
            return_value=([], failures, 1, diagnostics),
        ):
            payload = self.api.run_lite_mesa_analysis(
                small_aircraft_support_project("project-lite-mesa-all-timeout"),
                analysis_type="mission_reliability",
                settings={"samples": 2, "seed": 20260718},
            )

        self.assertEqual(payload["status"], "blocked")
        self.assertEqual(payload["requested_sample_count"], 2)
        self.assertEqual(payload["completed_sample_count"], 0)
        self.assertEqual(payload["failed_sample_count"], 2)
        self.assertEqual(payload["metric_moments"]["total_sample_count"], 2)
        self.assertEqual(payload["metric_moments"]["successful_sample_count"], 0)
        self.assertEqual(payload["metric_moments"]["failed_sample_count"], 2)
        self.assertTrue(all(metric["mean"] is None for metric in payload["metric_moments"]["metrics"]))
        self.assertTrue(all(metric["sample_variance"] is None for metric in payload["metric_moments"]["metrics"]))

    def test_lite_mesa_sample_preserves_initial_life_trace_from_shared_model(self) -> None:
        project = small_aircraft_support_project("project-lite-mesa-pre-life-trace")
        scenario = self.api.adapter.compile_scenario(project)

        sample = _run_aircraft_support_v1_analysis_sample(
            scenario["simulation_inputs"],
            seed=20260721,
            sample_index=0,
        )

        self.assertTrue(sample["lifecycle_trace"])
        self.assertEqual(sample["lifecycle_trace"][0]["initial_life_state"], {
            "calendar_days": 0,
            "flight_hours": 0.0,
            "takeoff_landing_cycles": 0,
        })

    def test_lite_mesa_finite_extremes_do_not_crash_the_api_or_emit_nonfinite_moments(self) -> None:
        project = small_aircraft_support_project("project-lite-mesa-finite-extremes")

        def extreme_samples(inputs, *, base_seed, settings):
            samples = [
                _run_aircraft_support_v1_analysis_sample(
                    inputs,
                    seed=base_seed + sample_index,
                    sample_index=sample_index,
                )
                for sample_index in range(settings["samples"])
            ]
            samples[0]["metrics"]["repair_backlog"] = 1e308
            samples[1]["metrics"]["repair_backlog"] = -1e308
            return samples, [], 1, []

        with mock.patch(
            "src.spare_mvp_backend.api._run_lite_mesa_analysis_samples",
            side_effect=extreme_samples,
        ):
            payload = self.api.run_lite_mesa_analysis(
                project,
                analysis_type="mission_reliability",
                settings={"samples": 2, "seed": 20260718, "parallelCores": 1},
            )

        repair_backlog = next(
            metric for metric in payload["metric_moments"]["metrics"]
            if metric["metric_id"] == "repair_backlog"
        )
        self.assertEqual(payload["status"], "session_complete")
        self.assertEqual(payload["aggregate_metrics"]["repair_backlog"], 0)
        self.assertEqual(repair_backlog["mean"], 0)
        self.assertIsNone(repair_backlog["sample_variance"])
        self.assertEqual(repair_backlog["valid_sample_count"], 2)
        self.assertEqual(repair_backlog["invalid_reason"], "sample_variance_not_finite")

    def test_periodic_profile_empty_slots_survive_project_round_trip_and_compile(self) -> None:
        project = small_aircraft_support_project("project-periodic-profile-empty-slots")
        project["missionProfile"]["periodicProfileLists"] = {
            "week": [],
            "month": [{"id": "month-empty", "name": "空月", "weekProfileIds": [""] * 4}],
            "year": [{"id": "year-empty", "name": "空年", "monthProfileIds": [""] * 12}],
        }

        saved = self.api.save_project(project)
        stored = self.api.get_project(saved["project_id"])
        compiled = self.adapter.compile_scenario(stored, model_family="aircraft_support_v1")

        self.assertEqual(
            stored["missionProfile"]["periodicProfileLists"],
            project["missionProfile"]["periodicProfileLists"],
        )
        self.assertEqual(compiled["simulation_model"]["family"], "aircraft_support_v1")

    def test_aircraft_mission_reliability_analysis_history_is_immutable_and_project_scoped(self) -> None:
        first_project = small_aircraft_support_project("project-aircraft-reliability-history-a")
        second_project = small_aircraft_support_project("project-aircraft-reliability-history-b")
        self.api.save_project(first_project)
        self.api.save_project(second_project)
        first_snapshot = {
            "reliability_block_diagram": {"nodes": [{"id": "engine", "reliability": 0.98}]},
            "calculation_details": [{"node_id": "engine", "failure_probability": 0.02}],
        }

        first = self.api.save_aircraft_mission_reliability_analysis(
            first_project["project_id"],
            {
                "analysis_id": "analysis-client-stable-001",
                "aircraft_model": "J-15",
                "mission_profile_id": "profile-five-hour",
                "mission_profile_name": "五小时任务",
                "duration_hours": 5,
                "aircraft_reliability": 0.98,
                "snapshot": first_snapshot,
            },
            actor_user_id="user-basic",
        )
        first_snapshot["calculation_details"][0]["failure_probability"] = 1
        second = self.api.save_aircraft_mission_reliability_analysis(
            first_project["project_id"],
            {
                "aircraft_model": "J-15",
                "mission_profile_id": "profile-ten-hour",
                "mission_profile_name": "十小时任务",
                "duration_hours": 10,
                "aircraft_reliability": 0.91,
                "snapshot": {"calculation_details": [{"node_id": "engine", "reliability": 0.91}]},
            },
            actor_user_id="user-data",
        )

        history = self.api.list_aircraft_mission_reliability_analyses(first_project["project_id"])
        self.assertEqual(history["project_id"], first_project["project_id"])
        self.assertEqual([entry["analysis_id"] for entry in history["analyses"]], [second["analysis_id"], first["analysis_id"]])
        self.assertEqual(history["analyses"][1]["snapshot"]["calculation_details"][0]["failure_probability"], 0.02)
        self.assertEqual(first["analysis_id"], "analysis-client-stable-001")
        self.assertEqual(first["created_by"], "user-basic")
        self.assertTrue(first["created_at"])
        self.assertEqual(self.api.list_aircraft_mission_reliability_analyses(second_project["project_id"])["analyses"], [])

        with self.assertRaises(BackendApiError) as duplicate:
            self.api.save_aircraft_mission_reliability_analysis(
                first_project["project_id"],
                {
                    "analysis_id": first["analysis_id"],
                    "aircraft_model": "J-15",
                    "mission_profile_id": "profile-five-hour",
                    "mission_profile_name": "五小时任务",
                    "duration_hours": 5,
                    "aircraft_reliability": 0.5,
                    "snapshot": {},
                },
                actor_user_id="user-basic",
            )
        self.assertEqual(duplicate.exception.code, "analysis_already_exists")

    def test_aircraft_mission_reliability_analysis_requires_valid_summary_and_project(self) -> None:
        project = small_aircraft_support_project("project-aircraft-reliability-validation")
        self.api.save_project(project)
        valid = {
            "aircraft_model": "J-15",
            "mission_profile_id": "profile-one",
            "mission_profile_name": "任务一",
            "duration_hours": 5,
            "aircraft_reliability": 0.95,
            "snapshot": {},
        }
        with self.assertRaises(KeyError):
            self.api.save_aircraft_mission_reliability_analysis(
                "project-missing",
                valid,
                actor_user_id="user-basic",
            )
        with self.assertRaisesRegex(ValueError, "between 0 and 1"):
            self.api.save_aircraft_mission_reliability_analysis(
                project["project_id"],
                {**valid, "aircraft_reliability": 1.1},
                actor_user_id="user-basic",
            )

    def test_replace_project_requires_current_version_and_records_audit(self) -> None:
        project = small_aircraft_support_project("project-replace-contract")
        self.api.save_project(project)
        expected = self.repository.project_updated_at(project["project_id"])
        replacement = copy.deepcopy(project)
        replacement["projectInfo"] = {"name": "覆盖后的项目"}
        result = self.api.replace_project(
            project["project_id"],
            replacement,
            expected_updated_at=expected,
            actor_user_id="user-admin",
        )
        self.assertEqual(result["status"], "replaced")
        self.assertEqual(self.api.get_project(project["project_id"])["projectInfo"]["name"], "覆盖后的项目")
        self.assertEqual(self.repository.list_audit_events(resource_id=project["project_id"])[-1]["action"], "project.replace")
        with self.assertRaises(BackendApiError) as conflict:
            self.api.replace_project(project["project_id"], replacement, expected_updated_at=expected, actor_user_id="user-admin")
        self.assertEqual(conflict.exception.code, "project_version_conflict")

    def test_create_imported_project_is_atomic_and_never_overwrites_existing_project(self) -> None:
        project = small_aircraft_support_project("project-xlsx-create-only")
        created = self.api.create_imported_project(project, actor_user_id="user-data")

        self.assertEqual(created["status"], "created")
        self.assertTrue(created["audit_event_id"].startswith("audit-"))
        stored_before = self.api.get_project(project["project_id"])
        conflicting = copy.deepcopy(project)
        conflicting["projectInfo"]["name"] = "must not overwrite"
        with self.assertRaises(BackendApiError) as conflict:
            self.api.create_imported_project(conflicting, actor_user_id="user-data")

        self.assertEqual(conflict.exception.code, "project_already_exists")
        self.assertEqual(self.api.get_project(project["project_id"]), stored_before)
        audits = self.repository.list_audit_events(resource_id=project["project_id"])
        self.assertEqual([event["action"] for event in audits], ["project.import_xlsx.create"])

        invalid = small_aircraft_support_project("project-xlsx-invalid-reference")
        invalid["components"][0]["productId"] = "does-not-exist"
        with self.assertRaises(BackendApiError) as invalid_error:
            self.api.create_imported_project(invalid, actor_user_id="user-data")
        self.assertEqual(invalid_error.exception.code, "invalid_project")
        self.assertIn(
            "missing_product_reference",
            {error.get("code") for error in invalid_error.exception.details["errors"]},
        )
        with self.assertRaises(KeyError):
            self.api.get_project(invalid["project_id"])

    def test_arbitrary_period_completion_and_four_downtime_contract(self) -> None:
        daily_success = [{"day": day, "plannedWaves": 1, "successfulWaves": 1} for day in range(1, 15)]
        daily_failure = copy.deepcopy(daily_success[:7])
        daily_failure[3]["successfulWaves"] = 0
        reliability = _lite_mesa_mission_reliability_result(
            {"data": {"mission_success_probability": 0.91, "sortie_rate": 0.8}},
            [
                {"metrics": {"ready_rate": 1, "failed_sorties": 0, "simulation_days": 14}, "daily_mission_reliability": daily_success},
                {"metrics": {"ready_rate": 1, "failed_sorties": 1, "simulation_days": 7}, "daily_mission_reliability": daily_failure},
                {"metrics": {"ready_rate": 1, "failed_sorties": 0, "simulation_days": 3}, "daily_mission_reliability": daily_success[:3]},
            ],
            {"maxTimeWindow": 20},
        )
        self.assertEqual(reliability["total_samples"], 3)
        self.assertEqual(reliability["successful_samples"], 2)
        self.assertEqual(reliability["failed_samples"], 1)
        self.assertEqual(reliability["valid_samples"], 3)
        self.assertAlmostEqual(reliability["period_completion_probability"], 2 / 3)
        self.assertEqual(reliability["period_duration_days"], 14)
        self.assertEqual(reliability["metrics"], [
            ["出动架次率", "0.800"],
            ["波次成功率", "91%"],
            ["整周期任务可靠度", "66.7%"],
            ["任务周期", "14 天"],
            ["仿真总次数", "3"],
            ["成功次数", "2"],
        ])
        self.assertEqual(
            [field["key"] for field in reliability["result_fields"]],
            ["sortie_rate", "wave_success_rate", "period_completion_probability", "period_duration_days"],
        )
        downtime = _lite_mesa_downtime_factors_result(
            {"data": []},
            {
                "downtime_failure_events": 2, "downtime_failure_hours": 10,
                "downtime_equipment_shortage_events": 3, "downtime_equipment_shortage_hours": 5,
                "downtime_spare_shortage_events": 4, "downtime_spare_shortage_hours": 4,
                "downtime_preventive_events": 1, "downtime_preventive_hours": 1,
            },
            [],
            {"topN": 10},
        )
        self.assertEqual({row["label"] for row in downtime["rows"]}, {"装备故障", "保障设备短缺", "备件短缺", "预防性维修"})
        self.assertAlmostEqual(sum(row["duration_contribution"] for row in downtime["rows"]), 1.0)

        event_backed = _lite_mesa_downtime_factors_result(
            {"data": []},
            {"downtime_failure_events": 99, "downtime_failure_hours": 99},
            [
                {
                    "sample_index": 0,
                    "seed": 17,
                    "downtime_events": [
                        {
                            "event_id": "d-1", "factor": "spare_shortage", "tail_number": "A-1",
                            "start_minute": 0, "end_minute": 120, "duration_minutes": 120,
                            "details": {"spare_name": "LRU", "required_quantity": 1, "available_quantity": 0},
                        },
                        {
                            "event_id": "d-2", "factor": "preventive", "tail_number": "A-2",
                            "start_minute": 0, "end_minute": 60, "duration_minutes": 60,
                            "details": {"trigger_type": None},
                        },
                    ],
                },
                {
                    "sample_index": 1,
                    "seed": 18,
                    "downtime_events": [
                        {
                            "event_id": "d-1", "factor": "spare_shortage", "tail_number": "A-3",
                            "start_minute": 0, "end_minute": 60, "duration_minutes": 60,
                            "details": {"spare_name": "LRU", "required_quantity": 1, "available_quantity": 0},
                        }
                    ],
                },
            ],
            {"topN": 10},
        )
        self.assertEqual(len(event_backed["event_details"]), 3)
        self.assertEqual(event_backed["event_details"][0]["sample_index"], 0)
        self.assertEqual(len({event["event_id"] for event in event_backed["event_details"]}), 3)
        self.assertEqual(event_backed["event_details"][0]["source_event_id"], "d-1")
        by_reason = {row["reason"]: row for row in event_backed["rows"]}
        self.assertEqual(by_reason["spare_shortage"]["event_count"], 2)
        self.assertEqual(by_reason["spare_shortage"]["downtime_hours"], 3)
        self.assertEqual(by_reason["preventive"]["event_count"], 1)
        self.assertEqual(by_reason["failure"]["event_count"], 0)
        self.assertAlmostEqual(
            sum(row["downtime_hours"] for row in event_backed["rows"]),
            sum(event["duration_minutes"] for event in event_backed["event_details"]) / 60,
        )

    def test_downtime_event_contract_localizes_aliases_and_all_visible_times(self) -> None:
        result = _lite_mesa_downtime_factors_result(
            {"data": []},
            {},
            [
                {
                    "sample_index": 0,
                    "seed": 17,
                    "downtime_events": [
                        {
                            "event_id": "d-1",
                            "factor": "failure",
                            "tail_number": "J15-101",
                            "taskId": "internal-task-id",
                            "taskLabel": "昼间制空任务",
                            "phaseId": "internal-phase-id",
                            "phaseName": "故障诊断",
                            "start_time": 1439,
                            "end_time": 1505,
                            "duration_minutes": 66,
                            "description": "unavailable_after_failure",
                            "details": {
                                "failure_mode": "must-not-project",
                                "failureMode": "must-not-project-alias",
                                "failureTime": 1435,
                                "repairCompletedTime": 1505,
                            },
                        },
                        {
                            "event_id": "d-2",
                            "factor": "preventive",
                            "start_minute": 0,
                            "end_minute": 1,
                            "duration_minutes": 1,
                            "details": {},
                        },
                    ],
                }
            ],
            {"topN": 10},
        )

        failure = next(event for event in result["event_details"] if event["factor"] == "failure")
        task_external = next(event for event in result["event_details"] if event["factor"] == "preventive")
        self.assertEqual(failure["mission_id"], "internal-task-id")
        self.assertEqual(failure["mission_name"], "昼间制空任务")
        self.assertEqual(failure["mission_phase_id"], "internal-phase-id")
        self.assertEqual(failure["mission_phase_name"], "故障诊断")
        self.assertEqual(failure["task_phase_label"], "昼间制空任务；阶段：故障诊断")
        self.assertEqual(failure["start_time_label"], "DAY_1 23:59")
        self.assertEqual(failure["end_time_label"], "DAY_2 01:05")
        self.assertEqual(failure["details"]["failure_time_label"], "DAY_1 23:55")
        self.assertEqual(failure["details"]["repair_completed_time_label"], "DAY_2 01:05")
        self.assertNotIn("failure_mode", failure["details"])
        self.assertNotIn("failureMode", failure["details"])
        self.assertNotIn("must-not-project", str(result))
        self.assertEqual(failure["description"], "飞机J15-101装备发生故障，当前不可用并等待修复")
        self.assertEqual(task_external["task_phase_label"], "不在任务阶段")
        self.assertEqual(task_external["details"]["failure_time_label"], "暂无时间")
        self.assertEqual(format_simulation_minute(0), "DAY_1 00:00")
        self.assertEqual(format_simulation_minute(3905), "DAY_3 17:05")
        self.assertEqual(format_simulation_minute(None), "暂无时间")

    def test_rms_export_is_a_real_xlsx_workbook(self) -> None:
        from openpyxl import load_workbook
        from io import BytesIO

        download = self.api.export_rms_allocation_xlsx({
            "project_name": "RMS案例",
            "method": "equal",
            "generated_at": "2026-07-12T00:00:00Z",
            "rows": [{
                "level": "系统", "nodeName": "动力", "model": "SYS-001",
                "installationCount": 2, "runningRatio": 0.8,
                "failureRate": 0.000833333333, "mtbfHours": 1200, "mttrHours": 2.5,
                "allocationShare": 0.625, "status": "已分配",
            }],
        })
        self.assertEqual(download["filename"], "rms-allocation-result.xlsx")
        workbook = load_workbook(BytesIO(download["body"]), read_only=True)
        sheet = workbook["RMS分配结果"]
        self.assertEqual(
            [sheet.cell(5, column).value for column in range(1, 7)],
            ["层级", "节点", "运行比", "失效率", "MTBF(h)", "MTTR(h)"],
        )
        self.assertEqual(
            [sheet.cell(6, column).value for column in range(1, 7)],
            ["系统", "动力", 0.8, 0.000833333333, 1200, 2.5],
        )

    def _fixture(self, name: str) -> dict:
        return json.loads((REPO_ROOT / "tests" / "fixtures" / name).read_text(encoding="utf-8"))

    def _public_import_template(self, name: str) -> dict:
        return json.loads((REPO_ROOT / "public" / "import-templates" / name).read_text(encoding="utf-8"))

    def _run_side_effect_counts(self) -> dict[str, int]:
        return {
            table: self.connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in ("scenarios", "simulation_runs", "result_summaries", "artifact_manifests")
        } | {
            "artifact_files": len([path for path in Path(self.tempdir.name).glob("**/*") if path.is_file()])
        }

    def _artifact_by_kind(self, manifest: dict, kind: str) -> dict:
        return next(artifact for artifact in manifest["artifacts"] if artifact["kind"] == kind)

    def _state_series_artifact(self, run_id: str) -> dict:
        return self._artifact_by_kind(self.api.get_run_artifacts(run_id), "visualization_state_series")

    def _artifact_payload(self, manifest: dict, kind: str) -> dict[str, Any]:
        artifact = self._artifact_by_kind(manifest, kind)
        return json.loads((Path(self.api.output_dir) / artifact["path"]).read_text(encoding="utf-8"))

    def _write_artifact_payload(self, manifest: dict, kind: str, payload: dict[str, Any]) -> None:
        artifact = self._artifact_by_kind(manifest, kind)
        (Path(self.api.output_dir) / artifact["path"]).write_text(
            json.dumps(payload, ensure_ascii=False, sort_keys=True),
            encoding="utf-8",
        )

    def _submit_successful_run(self) -> dict[str, Any]:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "m7 artifact download", "steps": 2})
        return self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

    def _create_imported_sample_project(self) -> dict[str, Any]:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])
        return self.api.create_project_from_modeling_import_as_system(import_package["importId"])

    def _submit_successful_aircraft_support_monte_carlo_run(
        self,
        *,
        analysis_type: str = "spare_shortfall",
    ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
        created = self._create_imported_sample_project()
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "current analysis result",
                "steps": 2,
                "projectJson": created["project"],
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 1,
                        "sweep": {
                            "failureRates": [0.05],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [2],
                        },
                    },
                    "spareShortfall": {"enabled": True},
                    "carryList": {"enabled": True},
                    "missionReliability": {"enabled": True},
                    "downtimeFactors": {"enabled": True},
                },
            },
        )
        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "analysis_type": analysis_type,
            }
        )
        return created, plan, run

    def _submit_successful_aircraft_support_monte_carlo_run_for_project(
        self,
        created: dict[str, Any],
        *,
        analysis_type: str,
        plan_name: str,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": plan_name,
                "steps": 2,
                "projectJson": created["project"],
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 1,
                        "sweep": {
                            "failureRates": [0.05],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [2],
                        },
                    },
                    "spareShortfall": {"enabled": True},
                    "carryList": {"enabled": True},
                    "missionReliability": {"enabled": True},
                    "downtimeFactors": {"enabled": True},
                },
            },
        )
        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "analysis_type": analysis_type,
            }
        )
        return plan, run

    def _persist_later_failed_analysis_run(
        self,
        success_run: dict[str, Any],
        *,
        run_id: str,
        analysis_type: str = "spare_shortfall",
        code: str = "executor_failed",
        message: str = "synthetic latest failure",
    ) -> dict[str, Any]:
        failed = copy.deepcopy(self.repository.get_run(success_run["run_id"]))
        failed.update(
            {
                "run_id": run_id,
                "status": "failed",
                "phase": "failed",
                "progress": 0,
                "artifact_manifest_id": f"artifact-manifest-{run_id}",
                "result_summary_id": None,
                "analysis_type": analysis_type,
                "error": {"code": code, "message": message, "details": {"analysis_type": analysis_type}},
                "started_at": "2099-01-01T00:00:00Z",
                "completed_at": "2099-01-01T00:00:01Z",
                "updated_at": "2099-01-01T00:00:01Z",
            }
        )
        failed["simulation_experiment_base"] = {
            **copy.deepcopy(failed.get("simulation_experiment_base") or {}),
            "run_id": run_id,
            "status": "failed",
            "analysis_type": analysis_type,
        }
        self.repository.upsert_run(failed)
        self.connection.execute(
            "UPDATE simulation_runs SET updated_at = ? WHERE run_id = ?",
            ("2099-01-01T00:00:01Z", run_id),
        )
        return failed

    def _reduced_scope_import_package_without_support_domain(self) -> dict[str, Any]:
        import_package = self._fixture("modeling_import_project.json")
        import_package["importId"] = "import-reduced-scope-no-support-domain"
        import_package["projectId"] = "project-reduced-scope-no-support-domain"
        import_package["usedTables"] = {
            "missionProfiles": True,
            "equipmentAssets": True,
            "reliabilityBlockDiagram": True,
            "supportResources": False,
            "supportActivities": False,
            "supportOrganization": False,
            "transportPolicies": False,
        }
        import_package["objects"] = copy.deepcopy(import_package["objects"])
        import_package["objects"].pop("supportResources", None)
        import_package["objects"].pop("supportActivities", None)
        import_package["objects"].pop("supportOrganization", None)
        return import_package

    def test_project_catalog_exposes_project_template_flag(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        project["projectInfo"] = {
            "name": "模板项目",
            "baseCode": "TPL",
            "summary": "项目数据层模板",
            "isTemplate": True,
        }
        saved = self.api.save_project(project)

        catalog = self.api.list_projects()
        entry = next(item for item in catalog["projects"] if item["project_id"] == saved["project_id"])

        self.assertEqual(entry["experiment_name"], "模板项目")
        self.assertEqual(entry["is_template"], True)

    def test_project_catalog_prefers_canonical_template_flag_after_unset_and_reload(self) -> None:
        project = small_aircraft_support_project("project-template-transition-001")
        project["isTemplate"] = True
        project["is_template"] = True
        project["projectInfo"] = {
            "name": "模板状态切换项目",
            "baseCode": "TPL",
            "summary": "模板状态持久化回归",
            "isTemplate": False,
            "is_template": True,
        }

        saved = self.api.save_project(project)
        reloaded = self.api.get_project(saved["project_id"])
        catalog = self.api.list_projects()
        entry = next(item for item in catalog["projects"] if item["project_id"] == saved["project_id"])

        self.assertEqual(reloaded["projectInfo"]["isTemplate"], False)
        self.assertEqual(entry["is_template"], False)

        reloaded["projectInfo"]["isTemplate"] = True
        reloaded["projectInfo"]["is_template"] = True
        reloaded["isTemplate"] = True
        reloaded["is_template"] = True
        self.api.save_project(reloaded)
        reset_entry = next(
            item for item in self.api.list_projects()["projects"]
            if item["project_id"] == saved["project_id"]
        )
        self.assertEqual(reset_entry["is_template"], True)

    def test_current_analysis_result_returns_only_valid_formal_projection(self) -> None:
        created, _plan, run = self._submit_successful_aircraft_support_monte_carlo_run()

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["analysis_type"], "spare_shortfall")
        self.assertEqual(current["status"], "completed")
        self.assertEqual(current["source"], "formal_backend")
        self.assertFalse(current["is_stale"])
        self.assertEqual(current["last_success_result"]["run_id"], run["run_id"])
        self.assertEqual(current["last_success_result"]["projection_type"], "spare_shortfall")
        self.assertEqual(current["internal_run_ref"]["run_id"], run["run_id"])
        self.assertEqual(current["internal_artifact_ref"]["kind"], "analysis_projection_spare_shortfall")
        self.assertIn("payload", current["last_success_result"])

    def test_current_analysis_result_fails_closed_on_projection_type_mismatch(self) -> None:
        created, _plan, run = self._submit_successful_aircraft_support_monte_carlo_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        payload = self._artifact_payload(manifest, "analysis_projection_spare_shortfall")
        payload["projection_type"] = "carry_list"
        self._write_artifact_payload(manifest, "analysis_projection_spare_shortfall", payload)

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["analysis_type"], "spare_shortfall")
        self.assertEqual(current["status"], "blocked")
        self.assertEqual(current["source"], "blocked")
        self.assertIsNone(current["last_success_result"])
        self.assertIn("projection_type mismatch", current["last_failure"]["message"])

    def test_current_analysis_result_fails_closed_when_projection_run_id_is_missing(self) -> None:
        created, _plan, run = self._submit_successful_aircraft_support_monte_carlo_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        payload = self._artifact_payload(manifest, "analysis_projection_spare_shortfall")
        payload.pop("run_id", None)
        self._write_artifact_payload(manifest, "analysis_projection_spare_shortfall", payload)

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["status"], "blocked")
        self.assertEqual(current["source"], "blocked")
        self.assertIsNone(current["last_success_result"])
        self.assertIn("projection run_id is required", current["last_failure"]["message"])

    def test_current_analysis_result_fails_closed_when_projection_model_family_is_missing(self) -> None:
        created, _plan, run = self._submit_successful_aircraft_support_monte_carlo_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        payload = self._artifact_payload(manifest, "analysis_projection_spare_shortfall")
        payload.pop("model_family", None)
        self._write_artifact_payload(manifest, "analysis_projection_spare_shortfall", payload)

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["status"], "blocked")
        self.assertEqual(current["source"], "blocked")
        self.assertIsNone(current["last_success_result"])
        self.assertIn("projection model_family is required", current["last_failure"]["message"])

    def test_current_analysis_result_ignores_successful_runs_for_other_analysis_type(self) -> None:
        created, _plan, spare_run = self._submit_successful_aircraft_support_monte_carlo_run(analysis_type="spare_shortfall")
        self._submit_successful_aircraft_support_monte_carlo_run_for_project(
            created,
            analysis_type="carry_list",
            plan_name="newer carry list current result",
        )

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["status"], "completed")
        self.assertEqual(current["last_success_result"]["run_id"], spare_run["run_id"])
        self.assertEqual(current["last_success_result"]["projection_type"], "spare_shortfall")

    def test_current_analysis_result_marks_previous_success_stale_after_latest_failure(self) -> None:
        created, _plan, run = self._submit_successful_aircraft_support_monte_carlo_run(analysis_type="spare_shortfall")
        failed = self._persist_later_failed_analysis_run(
            run,
            run_id="run-current-analysis-latest-failed",
            analysis_type="spare_shortfall",
        )

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["status"], "failed")
        self.assertTrue(current["is_stale"])
        self.assertEqual(current["last_success_result"]["run_id"], run["run_id"])
        self.assertEqual(current["last_failure"]["run_id"], failed["run_id"])
        self.assertEqual(current["last_failure"]["code"], "executor_failed")

    def test_current_analysis_result_marks_previous_success_stale_after_latest_projection_mismatch(self) -> None:
        created, _plan, old_run = self._submit_successful_aircraft_support_monte_carlo_run(analysis_type="spare_shortfall")
        _new_plan, new_run = self._submit_successful_aircraft_support_monte_carlo_run_for_project(
            created,
            analysis_type="spare_shortfall",
            plan_name="newer broken projection",
        )
        manifest = self.api.get_run_artifacts(new_run["run_id"])
        payload = self._artifact_payload(manifest, "analysis_projection_spare_shortfall")
        payload["projection_type"] = "carry_list"
        self._write_artifact_payload(manifest, "analysis_projection_spare_shortfall", payload)

        current = self.api.get_current_analysis_result(created["savedProject"]["project_id"], "spare_shortfall")

        self.assertEqual(current["status"], "blocked")
        self.assertTrue(current["is_stale"])
        self.assertEqual(current["last_success_result"]["run_id"], old_run["run_id"])
        self.assertEqual(current["last_failure"]["run_id"], new_run["run_id"])
        self.assertIn("projection_type mismatch", current["last_failure"]["message"])

    def test_current_carry_list_parameters_reach_adapter_only_for_carry_list(self) -> None:
        created = self._create_imported_sample_project()
        saved = created["savedProject"]
        plan_config = {
            "name": "current carry list parameter propagation",
            "steps": 2,
            "projectJson": created["project"],
            "analysisRequests": {
                "largeSample": {
                    "enabled": True,
                    "samples": 1,
                    "sweep": {
                        "failureRates": [0.05],
                        "spareMultipliers": [1.0],
                        "supportCapacities": [2],
                    },
                },
                "carryList": {
                    "enabled": True,
                    "scenarioOverrides": {
                        "sparesBySupportPoint": {
                            "carrier_deck": {"filter": "critical", "maxItems": 8}
                        },
                        "missionDurationMinutes": 240,
                    },
                    "carryListConfig": {
                        "missionConfidenceTarget": 0.92,
                    },
                },
            },
        }
        carry_plan = self.api.create_experiment_plan(saved["project_id"], plan_config)
        self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": carry_plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "analysis_type": "carry_list",
            }
        )
        spare_plan = self.api.create_experiment_plan(saved["project_id"], {**plan_config, "name": "spare ignores carry params"})
        self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": spare_plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "analysis_type": "spare_shortfall",
            }
        )

        carry_config = self.adapter.monte_carlo_run_calls[-2]["kwargs"]["monte_carlo_config"]
        spare_config = self.adapter.monte_carlo_run_calls[-1]["kwargs"]["monte_carlo_config"]

        self.assertEqual(carry_config["scenarioOverrides"]["missionDurationMinutes"], 240)
        self.assertEqual(
            carry_config["scenarioOverrides"]["sparesBySupportPoint"],
            {"carrier_deck": {"filter": "critical", "maxItems": 8}},
        )
        self.assertEqual(carry_config["carryListConfig"]["missionConfidenceTarget"], 0.92)
        self.assertNotIn("scenarioOverrides", spare_config)
        self.assertNotIn("carryListConfig", spare_config)

    def test_current_backend_flow_persists_complete_run_chain(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")

        validation = self.api.validate_project(project)
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "contract current", "steps": 4})
        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        self.assertTrue(validation["ok"])
        self.assertEqual(saved["project_id"], "project-aircraft-support-contract-001")
        self.assertEqual(snapshot["project_id"], saved["project_id"])
        self.assertEqual(snapshot["project_version"], "project-v0.1")
        self.assertEqual(plan["project_id"], saved["project_id"])
        self.assertEqual(plan["config"]["steps"], 4)
        self.assertRegex(run["run_id"], r"^run-scenario-aircraft-support-contract-demo-[0-9a-f]{12}-\d{4}$")
        self.assertEqual(run["project_id"], "project-aircraft-support-contract-001")
        self.assertRegex(run["scenario_id"], r"^scenario-aircraft-support-contract-demo-[0-9a-f]{12}-\d{4}$")
        self.assertEqual(run["result_summary_id"], f"result-{run['run_id']}")
        self.assertEqual(run["artifact_manifest_id"], f"artifact-manifest-{run['run_id']}")
        self.assertEqual(run["status"], "succeeded")

        self.assertEqual(len(self.adapter.compile_calls), 1)
        self.assertEqual(self.adapter.compile_calls[0], (export_project_json(project), "aircraft_support_v1"))
        self.assertEqual(len(self.adapter.run_calls), 1)
        self.assertEqual(self.adapter.run_calls[0][0]["scenario_id"], run["scenario_id"])
        self.assertEqual(self.adapter.run_calls[0][1], 4)
        self.assertEqual(self.adapter.run_calls[0][2], run["run_id"])

        stored_run = self.api.get_run(run["run_id"])
        result = self.api.get_run_result(run["run_id"])
        manifest = self.api.get_run_artifacts(run["run_id"])
        chain = self.api.get_run_chain(run["run_id"])

        self.assertEqual(stored_run["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(stored_run["status"], "succeeded")
        self.assertEqual(result["result_id"], run["result_summary_id"])
        self.assertEqual(result["run_id"], run["run_id"])
        self.assertIn("mission_success_rate", result["metrics"])
        self.assertEqual(manifest["artifact_manifest_id"], run["artifact_manifest_id"])
        self.assertEqual(manifest["run_id"], run["run_id"])
        self.assertEqual(chain["project_id"], run["project_id"])
        self.assertEqual(chain["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(chain["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(chain["scenario_id"], run["scenario_id"])
        self.assertEqual(chain["run_id"], run["run_id"])
        self.assertEqual(chain["result_summary_id"], run["result_summary_id"])
        self.assertEqual(chain["artifact_manifest_id"], run["artifact_manifest_id"])

    def test_save_project_rejects_runtime_monte_carlo_config_in_project_payload(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        project["monteCarlo"] = {
            "failureRates": [0.05],
            "spareMultipliers": [1.0],
            "supportCapacities": [2],
        }
        project.setdefault("missionProfile", {})["monteCarlo"] = {
            "failureRates": [0.08],
        }
        project["missionProfile"]["analysisRequests"] = {
            "largeSample": {
                "enabled": True,
                "samples": 5,
                "sweep": {
                    "failureRates": [0.07],
                },
            },
        }
        project["analysisRequests"] = {
            "largeSample": {
                "enabled": True,
                "samples": 3,
                "sweep": {
                    "failureRates": [0.05],
                    "spareMultipliers": [1.0],
                    "supportCapacities": [2],
                },
            },
        }
        project["stopPolicy"] = {"mode": "or", "conditions": [{"type": "duration"}]}
        project["missionProfile"]["stopPolicy"] = {"mode": "or", "conditions": [{"type": "failure"}]}

        validation = self.api.validate_project(project)

        self.assertFalse(validation["ok"])
        self.assertEqual(
            sorted(error["path"] for error in validation["errors"] if error["code"] == "unsupported_project_runtime_config"),
            [
                "analysisRequests",
                "missionProfile.analysisRequests",
                "missionProfile.monteCarlo",
                "missionProfile.stopPolicy",
                "monteCarlo",
                "stopPolicy",
            ],
        )
        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_project(project)
        self.assertEqual(ctx.exception.code, "invalid_project")
        self.assertIn("unsupported_project_runtime_config", {error["code"] for error in ctx.exception.details["errors"]})

    def test_save_project_rejects_invalid_equipment_k_out_of_n(self) -> None:
        project = small_aircraft_support_project("project-invalid-k-out-of-n")
        project["components"][0]["quantity"] = 2
        project["components"][0]["kOutOfN"] = {"enabled": True, "n": 2, "k": 0}

        validation = self.api.validate_project(project)

        self.assertFalse(validation["ok"])
        k_errors = [error for error in validation["errors"] if error["code"] == "invalid_equipment_k_out_of_n"]
        self.assertEqual(k_errors[0]["path"], "components[0].kOutOfN.k")
        self.assertIn("1 ≤ k ≤ n", k_errors[0]["message"])
        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_project(project)
        self.assertEqual(ctx.exception.code, "invalid_project")

    def test_save_project_repairs_legacy_equipment_tree_integrity(self) -> None:
        project = small_aircraft_support_project("project-legacy-equipment-tree")
        project["equipment"] = {"model": "J-15", "wholeMachineModels": ["J-15"]}
        project["products"] = [{
            "id": "product-j15-avionics",
            "name": "航电系统",
            "failureDistribution": {"distributionType": "指数分布", "parameters": "mean=125, sigma=14"},
        }]
        project["components"] = [{
            "id": "j15-avionics",
            "name": "航电系统",
            "productId": "product-j15-avionics",
            "aircraftModel": "J-15",
            "parentId": "aircraft-root",
            "quantity": 1,
            "failureDistribution": {"distributionType": "指数分布", "parameters": "mean=125, sigma=14"},
        }]
        project["supportActivities"].append({
            "id": "preventive",
            "activityName": "50小时舰载机定检",
            "activityType": "预防性维修",
            "equipmentId": "j35-hydraulic",
            "jobs": [],
        })

        self.api.save_project(project)
        stored = self.api.get_project(project["project_id"])
        compile_codes = {
            issue["code"] for issue in self.api.adapter._aircraft_support_v1_compile_issues(stored)
        }

        self.assertEqual(stored["components"][1]["id"], "aircraft-root")
        self.assertEqual(stored["components"][0]["failureDistribution"]["distributionType"], "正态分布")
        self.assertEqual(stored["supportActivities"][1]["equipmentId"], "aircraft-root")
        self.assertTrue({
            "missing_component_parent",
            "invalid_component_failure_distribution",
            "missing_equipment_reference",
        }.isdisjoint(compile_codes))

    def test_save_project_rejects_unresolved_equipment_references(self) -> None:
        project = small_aircraft_support_project("project-unresolved-equipment-reference")
        project["supportActivities"][0].update({
            "activityName": "局部设备维修",
            "equipmentId": "missing-component",
        })

        validation = self.api.validate_project(project)

        self.assertFalse(validation["ok"])
        self.assertIn(
            "missing_equipment_reference",
            {error["code"] for error in validation["errors"]},
        )
        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_project(project)
        self.assertEqual(ctx.exception.code, "invalid_project")

    def test_save_project_migrates_basic_mission_support_activity_name_matching_legacy_name(self) -> None:
        project = small_aircraft_support_project("project-invalid-support-activity-name")
        project["basicMissions"][0]["supportActivityName"] = "Legacy display name"
        project["supportActivities"][0]["name"] = "Legacy display name"
        project["supportActivities"][0]["activityName"] = "Canonical support plan"

        validation = self.api.validate_project(project)

        self.assertTrue(validation["ok"])
        self.api.save_project(project)
        stored = self.api.get_project(project["project_id"])
        self.assertEqual(stored["basicMissions"][0]["supportActivityName"], "Canonical support plan")

    def test_get_project_migrates_legacy_basic_mission_support_activity_name(self) -> None:
        project = small_aircraft_support_project("project-loaded-legacy-support-activity-name")
        project["basicMissions"][0]["supportActivityName"] = "Legacy display name"
        project["supportActivities"][0].update({
            "name": "Legacy display name",
            "activityName": "Canonical support plan",
        })
        self.api.repository.upsert_project(project)

        loaded = self.api.get_project(project["project_id"])

        self.assertEqual(loaded["basicMissions"][0]["supportActivityName"], "Canonical support plan")

    def test_save_project_rejects_ambiguous_or_missing_legacy_support_activity_name_migration(self) -> None:
        for suffix, activities in (
            (
                "ambiguous-legacy-name",
                [
                    {"id": "support-a", "name": "Legacy display name", "activityName": "Canonical support plan", "jobs": []},
                    {"id": "support-b", "name": "Legacy display name", "activityName": "Other support plan", "jobs": []},
                ],
            ),
            (
                "missing-canonical-name",
                [{"id": "support-a", "name": "Legacy display name", "jobs": []}],
            ),
        ):
            with self.subTest(suffix=suffix):
                project = small_aircraft_support_project(f"project-{suffix}")
                project["basicMissions"][0]["supportActivityName"] = "Legacy display name"
                project["supportActivities"] = activities

                validation = self.api.validate_project(project)

                self.assertFalse(validation["ok"])
                relationship_errors = [
                    error for error in validation["errors"]
                    if error["code"] == "invalid_basic_mission_support_activity_name"
                ]
                self.assertEqual(relationship_errors[0]["path"], "basicMissions[0].supportActivityName")
                with self.assertRaises(BackendApiError) as ctx:
                    self.api.save_project(project)
                self.assertEqual(ctx.exception.code, "invalid_project")

    def test_save_project_rejects_duplicate_basic_mission_support_activity_name_matches(self) -> None:
        project = small_aircraft_support_project("project-duplicate-support-activity-name")
        project["basicMissions"][0]["supportActivityName"] = "Canonical support plan"
        project["supportActivities"][0]["activityName"] = "Canonical support plan"
        project["supportActivities"].append({
            **project["supportActivities"][0],
            "id": "duplicate-support-plan",
        })

        validation = self.api.validate_project(project)

        self.assertFalse(validation["ok"])
        relationship_errors = [
            error for error in validation["errors"]
            if error["code"] == "invalid_basic_mission_support_activity_name"
        ]
        self.assertEqual(relationship_errors[0]["path"], "basicMissions[0].supportActivityName")
        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_project(project)
        self.assertEqual(ctx.exception.code, "invalid_project")

    def test_save_project_persists_basic_mission_support_activity_name_after_duplicate_name_migration(self) -> None:
        project = small_aircraft_support_project("project-support-activity-name-migration")
        project["basicMissions"] = [
            {
                **copy.deepcopy(project["basicMissions"][0]),
                "id": f"j16-basic-{index}",
                "missionId": f"j16-basic-{index}",
                "name": f"J16 basic mission {index}",
                "supportActivityName": "J16基本方案",
            }
            for index in range(1, 5)
        ]
        project["supportActivities"] = [
            {
                "id": "j16-primary",
                "activityType": "使用保障",
                "planType": "使用保障方案",
                "aircraftModel": "J16",
                "activityName": "J16基本方案",
                "jobs": [],
            },
            {
                "id": "j16-duplicate-renamed",
                "activityType": "使用保障",
                "planType": "使用保障方案",
                "aircraftModel": "J16",
                "activityName": "J16基本方案（2）",
                "jobs": [],
            },
        ]

        validation = self.api.validate_project(project)
        self.assertTrue(validation["ok"])
        self.api.save_project(project)
        stored = self.api.get_project(project["project_id"])

        self.assertEqual(
            [mission["supportActivityName"] for mission in stored["basicMissions"]],
            ["J16基本方案"] * 4,
        )
        self.assertEqual(
            [activity["activityName"] for activity in stored["supportActivities"]],
            ["J16基本方案", "J16基本方案（2）"],
        )

    def test_get_project_strips_legacy_persisted_monte_carlo_payload(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        project["project_id"] = "project-legacy-mc"
        project["monteCarlo"] = {
            "failureRates": [0.05],
            "spareMultipliers": [1.0],
            "supportCapacities": [2],
        }
        project.setdefault("missionProfile", {})["monteCarlo"] = {"failureRates": [0.08]}
        project["missionProfile"]["analysisRequests"] = {
            "largeSample": {
                "enabled": True,
                "samples": 5,
                "sweep": {
                    "failureRates": [0.07],
                },
            },
        }
        project["stopPolicy"] = {"mode": "or", "conditions": [{"type": "duration"}]}
        project["missionProfile"]["stopPolicy"] = {"mode": "or", "conditions": [{"type": "failure"}]}
        self.api.repository.upsert_project(project)

        stored = self.api.get_project("project-legacy-mc")

        self.assertNotIn("monteCarlo", stored)
        self.assertNotIn("monteCarlo", stored["missionProfile"])
        self.assertNotIn("analysisRequests", stored["missionProfile"])
        self.assertNotIn("stopPolicy", stored)
        self.assertNotIn("stopPolicy", stored["missionProfile"])

    def test_save_project_strips_non_model_project_fields(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        project["project_id"] = "project-non-model-fields"
        project["deletedSupportResourceKeys"] = ["support-org:spare:legacy"]
        project["missionProfile"] = {
            "name": "model profile",
            "profileType": "legacy label",
            "repeatCycleHours": 6,
            "endCondition": "legacy end condition",
            "analysisRequests": {
                "largeSample": {
                    "enabled": True,
                    "samples": 5,
                },
            },
        }
        project["supportActivities"] = [
            {
                "id": "activity-1",
                "name": "Legacy display name",
                "activityType": "飞行前保障",
                "planType": "直接准备方案",
                "resourceId": "legacy-node",
                "requireDevices": 3,
                "requiredDevices": 2,
                "requiredPersonnel": 4,
                "jobs": [
                    {
                        "activityCode": "BA-001",
                        "workName": "电源车准备",
                        "durationMinutes": 15,
                        "predecessors": [],
                    },
                    {
                        "activityCode": "BA-002",
                        "workName": "通电检查",
                        "durationMinutes": 30,
                        "predecessors": ["BA-001"],
                        "maxRepairTimeMinutes": 999,
                        "meanRepairTimeMinutes": 888,
                        "mttrMinutes": 777,
                        "repairDistribution": {"distributionType": "固定值", "value": 999},
                    },
                ],
            },
        ]

        saved = self.api.save_project(project)
        stored = self.api.get_project(saved["project_id"])

        self.assertNotIn("deletedSupportResourceKeys", stored)
        self.assertNotIn("experiment", stored)
        self.assertNotIn("profileType", stored["missionProfile"])
        self.assertNotIn("repeatCycleHours", stored["missionProfile"])
        self.assertNotIn("endCondition", stored["missionProfile"])
        self.assertNotIn("analysisRequests", stored["missionProfile"])
        self.assertNotIn("requireDevices", stored["supportActivities"][0])
        self.assertNotIn("name", stored["supportActivities"][0])
        self.assertEqual(stored["supportActivities"][0]["resourceId"], "legacy-node")
        self.assertNotIn("requiredDevices", stored["supportActivities"][0])
        self.assertNotIn("requiredPersonnel", stored["supportActivities"][0])
        self.assertEqual(stored["supportActivities"][0]["activityName"], "Legacy display name")
        self.assertEqual(stored["supportActivities"][0]["planType"], "使用保障方案")
        self.assertNotIn("jobs", stored["supportActivities"][0])
        self.assertEqual(stored["supportActivities"][0]["activityCodes"], ["BA-001", "BA-002"])
        self.assertEqual(stored["supportActivities"][0]["predecessors"], {"BA-001": [], "BA-002": ["BA-001"]})
        self.assertEqual(
            [job["activityCode"] for job in stored["supportActivityJobs"]],
            ["BA-001", "BA-002"],
        )
        self.assertNotIn("predecessors", stored["supportActivityJobs"][1])
        self.assertNotIn("maxRepairTimeMinutes", stored["supportActivityJobs"][1])
        self.assertNotIn("meanRepairTimeMinutes", stored["supportActivityJobs"][1])
        self.assertNotIn("mttrMinutes", stored["supportActivityJobs"][1])
        self.assertNotIn("repairDistribution", stored["supportActivityJobs"][1])

    def test_save_project_preserves_distinct_legacy_support_jobs_with_duplicate_codes(self) -> None:
        project = small_aircraft_support_project("project-duplicate-support-activity-codes")
        project["supportActivities"] = [
            {
                "id": "ops-activity",
                "activityType": "使用保障",
                "jobs": [
                    {
                        "activityCode": "BA-001",
                        "workName": "使用保障准备",
                        "durationMinutes": 20,
                        "predecessors": [],
                    }
                ],
            },
            {
                "id": "preventive-activity",
                "activityType": "预防性维修",
                "jobs": [
                    {
                        "activityCode": "BA-001",
                        "workName": "定检准备",
                        "durationMinutes": 45,
                        "predecessors": [],
                    },
                    {
                        "activityCode": "BA-003",
                        "workName": "定检执行",
                        "durationMinutes": 60,
                        "predecessors": ["BA-001"],
                    }
                ],
            },
        ]

        saved = self.api.save_project(project)
        stored = self.api.get_project(saved["project_id"])

        self.assertEqual(stored["supportActivities"][0]["activityCodes"], ["BA-001"])
        self.assertEqual(stored["supportActivities"][1]["activityCodes"], ["BA-002", "BA-003"])
        self.assertEqual(stored["supportActivities"][1]["predecessors"], {"BA-002": [], "BA-003": ["BA-002"]})
        self.assertEqual(
            [(job["activityCode"], job["workName"], job["durationMinutes"]) for job in stored["supportActivityJobs"]],
            [
                ("BA-001", "使用保障准备", 20),
                ("BA-002", "定检准备", 45),
                ("BA-003", "定检执行", 60),
            ],
        )
        self.assertNotIn("jobs", stored["supportActivities"][0])
        self.assertNotIn("jobs", stored["supportActivities"][1])

    def test_run_service_submits_current_run_and_returns_status_envelope(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "m6 status", "steps": 2})

        service = RunService(self.repository, self.adapter, self.api.output_dir)
        submitted = service.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        status = service.get_run_status(submitted["run_id"])

        self.assertEqual(submitted["status"], "succeeded")
        self.assertEqual(submitted["phase"], "completed")
        self.assertEqual(submitted["progress"], 1)
        self.assertEqual(submitted["run_type"], "single")
        self.assertEqual(submitted["model_family"], "aircraft_support_v1")
        self.assertEqual(status["run_id"], submitted["run_id"])
        self.assertEqual(status["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(status["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(status["result_summary_id"], submitted["result_summary_id"])
        self.assertEqual(status["artifact_manifest_id"], submitted["artifact_manifest_id"])
        self.assertEqual(self.adapter.run_calls[0][1], 2)

    def test_run_service_augments_run_config_artifact_with_plan_and_snapshot_identity(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "m7 run config", "steps": 2})

        submitted = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        manifest = self.api.get_run_artifacts(submitted["run_id"])
        run_config_artifact = next(artifact for artifact in manifest["artifacts"] if artifact["kind"] == "run_config")
        payload = json.loads((Path(self.api.output_dir) / run_config_artifact["path"]).read_text(encoding="utf-8"))

        self.assertEqual(payload["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(payload["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(payload["project_id"], saved["project_id"])
        self.assertEqual(payload["run_id"], submitted["run_id"])

    def test_run_service_keeps_disk_artifact_manifest_in_sync_after_run_config_augmentation(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "m7 manifest sync", "steps": 2})

        submitted = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        repository_manifest = self.api.get_run_artifacts(submitted["run_id"])
        disk_manifest_path = Path(self.api.output_dir) / submitted["run_id"] / "artifact-manifest.json"
        disk_manifest = json.loads(disk_manifest_path.read_text(encoding="utf-8"))
        repository_run_config = self._artifact_by_kind(repository_manifest, "run_config")
        disk_run_config = self._artifact_by_kind(disk_manifest, "run_config")

        for key in ("sha256", "size_bytes", "path", "artifact_id"):
            self.assertEqual(disk_run_config[key], repository_run_config[key])
        for artifact in repository_manifest["artifacts"]:
            artifact_path = Path(self.api.output_dir) / artifact["path"]
            data = artifact_path.read_bytes()
            self.assertEqual(hashlib.sha256(data).hexdigest(), artifact["sha256"], artifact)
            self.assertEqual(len(data), artifact["size_bytes"], artifact)

    def test_backend_api_lists_runs_with_strict_include_deleted_flag_and_limit_clamp(self) -> None:
        runs = [self._submit_successful_run() for _ in range(3)]
        self.api.soft_delete_run(runs[0]["run_id"], actor_user_id="user-admin")

        hidden = self.api.list_runs({"include_deleted": "0", "limit": "500"})
        visible = self.api.list_runs({"include_deleted": "1", "limit": "500"})
        truthy_true = self.api.list_runs({"include_deleted": "True"})
        truthy_lower = self.api.list_runs({"include_deleted": "true"})
        truthy_bool = self.api.list_runs({"include_deleted": True})
        false_string = self.api.list_runs({"include_deleted": "yes"})
        limited = self.api.list_runs({"include_deleted": "1", "limit": "2"})

        self.assertNotIn(runs[0]["run_id"], [item["run_id"] for item in hidden["runs"]])
        self.assertNotIn(runs[0]["run_id"], [item["run_id"] for item in false_string["runs"]])
        self.assertIn(runs[0]["run_id"], [item["run_id"] for item in visible["runs"]])
        self.assertIn(runs[0]["run_id"], [item["run_id"] for item in truthy_true["runs"]])
        self.assertIn(runs[0]["run_id"], [item["run_id"] for item in truthy_lower["runs"]])
        self.assertIn(runs[0]["run_id"], [item["run_id"] for item in truthy_bool["runs"]])
        self.assertEqual(len(limited["runs"]), 2)

    def test_backend_api_resolves_artifact_download_with_hash_verification(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        artifact = manifest["artifacts"][0]

        download = self.api.get_run_artifact_download(
            run["run_id"],
            artifact["artifact_id"],
            actor_user_id="user-admin",
        )

        self.assertEqual(download["artifact"]["artifact_id"], artifact["artifact_id"])
        self.assertTrue(download["path"].is_file())
        self.assertEqual(download["content_type"], artifact["media_type"])
        self.assertEqual(hashlib.sha256(download["body"]).hexdigest(), artifact["sha256"])
        audit = self.repository.list_audit_events(resource_id=run["run_id"])
        self.assertEqual([event["action"] for event in audit], ["runs.artifact.download"])
        self.assertEqual(audit[0]["actor_user_id"], "user-admin")

    def test_successful_current_runs_publish_downloadable_visualization_state_series(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        branch_project = copy.deepcopy(project)
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        single_plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "m9 single state series", "steps": 2},
        )
        monte_carlo_plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "m9 mc state series",
                "steps": 2,
                "projectJson": branch_project,
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 3,
                        "sweep": {
                            "failureRates": [0.05],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [2],
                        },
                    }
                },
            },
        )
        submitted_runs = [
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": single_plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "single",
                }
            ),
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": monte_carlo_plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "monte_carlo",
                }
            ),
        ]

        for submitted in submitted_runs:
            with self.subTest(run_type=submitted["run_type"]):
                manifest = self.api.get_run_artifacts(submitted["run_id"])
                state_series_artifact = self._artifact_by_kind(manifest, "visualization_state_series")
                payload_path = Path(self.api.output_dir) / state_series_artifact["path"]
                payload_data = payload_path.read_bytes()
                payload = json.loads(payload_data.decode("utf-8"))
                download = self.api.get_run_artifact_download(
                    submitted["run_id"],
                    state_series_artifact["artifact_id"],
                    actor_user_id="user-admin",
                )

                self.assertEqual(state_series_artifact["schema_version"], "visualization-state-series-v0")
                self.assertEqual(state_series_artifact["media_type"], "application/json")
                self.assertEqual(state_series_artifact["source_run_id"], submitted["run_id"])
                self.assertEqual(state_series_artifact["source_result_summary_id"], submitted["result_summary_id"])
                self.assertEqual(state_series_artifact["source_scenario_id"], submitted["scenario_id"])
                self.assertEqual(hashlib.sha256(payload_data).hexdigest(), state_series_artifact["sha256"])
                self.assertEqual(len(payload_data), state_series_artifact["size_bytes"])
                self.assertEqual(download["body"], payload_data)
                self.assertEqual(payload["schema_version"], "visualization-state-series-v0")
                self.assertEqual(payload["run_id"], submitted["run_id"])
                self.assertEqual(payload["scenario_id"], submitted["scenario_id"])
                self.assertEqual(payload["scenario_version"], manifest["scenario_version"])
                self.assertEqual(payload["model_family"], "aircraft_support_v1")
                self.assertEqual(payload["artifact_manifest_id"], submitted["artifact_manifest_id"])
                self.assertEqual(payload["result_summary_id"], submitted["result_summary_id"])
                self.assertEqual(payload["run_config_artifact_id"], f"run_config-{submitted['run_id']}")
                self.assertEqual(payload["input_project_artifact_id"], f"input_project-{submitted['run_id']}")
                self.assertEqual(payload["compiled_scenario_artifact_id"], f"compiled_scenario-{submitted['run_id']}")
                self.assertGreater(len(payload["frames"]), 0)
                frame_steps = [frame["step"] for frame in payload["frames"]]
                self.assertEqual(frame_steps, sorted(frame_steps))
                self.assertEqual(len(frame_steps), len(set(frame_steps)))
                for frame in payload["frames"]:
                    self.assertEqual(frame["run_id"], submitted["run_id"])
                    self.assertEqual(frame["trace"]["run_id"], submitted["run_id"])
                    self.assertEqual(frame["trace"]["scenario_id"], submitted["scenario_id"])
                    self.assertEqual(frame["trace"]["result_summary_id"], submitted["result_summary_id"])
                    self.assertEqual(frame["trace"]["artifact_manifest_id"], submitted["artifact_manifest_id"])
                    self.assertIsInstance(frame["step"], int)
                    self.assertIsInstance(frame["aircraft_state"], dict)
                    self.assertIsInstance(frame["mission_state"], dict)
                    self.assertIsInstance(frame["resource_state"], dict)
                    self.assertIsInstance(frame["event_summary"], dict)
                    self.assertIsInstance(frame["aircraft"], list)
                    self.assertIsInstance(frame["missions"], list)
                    self.assertIsInstance(frame["resources"], list)
                    self.assertIsInstance(frame["events"], list)
                    self.assertGreater(len(frame["aircraft"]), 0)
                    self.assertGreater(len(frame["missions"]), 0)
                    self.assertGreater(len(frame["resources"]), 0)
                    self.assertGreater(len(frame["events"]), 0)
                    for event in frame["events"]:
                        self.assertTrue(event["event_id"])
                        self.assertEqual(event["run_id"], submitted["run_id"])
                        self.assertEqual(event["step"], frame["step"])
                        self.assertIsInstance(event["metric_refs"], list)
                    if submitted["run_type"] == "monte_carlo":
                        self.assertIsInstance(frame["sample_index"], int)
                        self.assertIsInstance(frame["sample_step"], int)

    def test_m9_2_subscribe_run_state_stream_reuses_visualization_state_series(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]
        artifact = self._state_series_artifact(run_id)
        payload = json.loads((Path(self.api.output_dir) / artifact["path"]).read_text(encoding="utf-8"))

        envelope = self.api.subscribe_run_state_stream(run_id)

        self.assertEqual(envelope["stream_id"], f"state-stream-{run_id}")
        self.assertEqual(envelope["run_id"], run_id)
        self.assertEqual(envelope["status"]["run_id"], run_id)
        self.assertEqual(envelope["status"]["status"], "succeeded")
        self.assertEqual(envelope["artifact_id"], artifact["artifact_id"])
        event_types = [event["event_type"] for event in envelope["events"]]
        self.assertEqual(event_types, ["run_status", *["state_frame"] * len(payload["frames"]), "artifact_ready"])
        self.assertEqual(envelope["events"][0]["payload"], envelope["status"])
        for index, event in enumerate(envelope["events"][1:-1]):
            source_frame = payload["frames"][index]
            self.assertEqual(event["event_type"], "state_frame")
            self.assertEqual(
                event["payload"],
                {
                    "schema_version": "visualization-state-frame-v0",
                    "stream_id": f"state-stream-{run_id}",
                    "run_id": run_id,
                    "artifact_id": artifact["artifact_id"],
                    "scenario_id": payload["scenario_id"],
                    "scenario_version": payload["scenario_version"],
                    "model_family": payload["model_family"],
                    "artifact_manifest_id": payload["artifact_manifest_id"],
                    "result_summary_id": payload["result_summary_id"],
                    "run_config_artifact_id": payload["run_config_artifact_id"],
                    "input_project_artifact_id": payload["input_project_artifact_id"],
                    "compiled_scenario_artifact_id": payload["compiled_scenario_artifact_id"],
                    "step": source_frame["step"],
                    "frame_index": index,
                    "frame_count": len(payload["frames"]),
                    "frame": source_frame,
                },
            )
        self.assertEqual(envelope["events"][-1]["event_type"], "artifact_ready")
        self.assertEqual(envelope["events"][-1]["payload"]["run_id"], run_id)
        self.assertEqual(envelope["events"][-1]["payload"]["artifact_id"], artifact["artifact_id"])
        self.assertEqual(envelope["events"][-1]["payload"]["kind"], "visualization_state_series")

        with self.assertRaises(KeyError):
            self.api.subscribe_run_state_stream("run-missing")

    def test_m9_3_run_control_cancel_is_backend_confirmed_and_audited(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]

        controlled = self.api.control_run(run_id, "cancel", actor_user_id="user-admin")

        self.assertEqual(controlled["run_id"], run_id)
        self.assertEqual(controlled["status"], "cancelled")
        self.assertEqual(controlled["phase"], "cancelled")
        self.assertEqual(controlled["progress"], 1)
        self.assertTrue(controlled["cancelled_at"])
        self.assertEqual(controlled["cancelled_by"], "user-admin")
        self.assertEqual(controlled["control"]["action"], "cancel")
        self.assertEqual(controlled["control"]["outcome"], "allowed")
        self.assertEqual(controlled["control"]["actor_user_id"], "user-admin")
        self.assertTrue(controlled["control"]["controlled_at"])

        stored = self.api.get_run(run_id)
        self.assertEqual(stored["status"], "cancelled")
        self.assertEqual(stored["cancelled_by"], "user-admin")
        events = self.repository.list_audit_events(resource_id=run_id)
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["action"], "runs.control.cancel")
        self.assertEqual(events[0]["outcome"], "allowed")
        self.assertEqual(events[0]["actor_user_id"], "user-admin")
        self.assertEqual(events[0]["details"]["status"], "cancelled")

    def test_m9_3_unsupported_run_control_fails_closed_and_audits_denial(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]
        before = self.api.get_run(run_id)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.control_run(run_id, "pause", actor_user_id="user-admin")

        self.assertEqual(ctx.exception.code, "unsupported_run_control")
        self.assertEqual(self.api.get_run(run_id), before)
        events = self.repository.list_audit_events(resource_id=run_id)
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["action"], "runs.control.pause")
        self.assertEqual(events[0]["outcome"], "denied")
        self.assertEqual(events[0]["actor_user_id"], "user-admin")
        self.assertEqual(events[0]["details"]["reason"], "unsupported_run_control")

        deleted = self.api.soft_delete_run(run_id, actor_user_id="user-admin")
        self.assertEqual(deleted["lifecycle_status"], "deleted")
        with self.assertRaises(BackendApiError) as deleted_ctx:
            self.api.control_run(run_id, "cancel", actor_user_id="user-admin")
        self.assertEqual(deleted_ctx.exception.code, "run_deleted")
        deleted_events = self.repository.list_audit_events(resource_id=run_id)
        self.assertEqual(deleted_events[-1]["action"], "runs.control.cancel")
        self.assertEqual(deleted_events[-1]["outcome"], "denied")
        self.assertEqual(deleted_events[-1]["details"]["reason"], "run_deleted")

    def test_m9_3_retry_blocks_stale_official_result_and_artifact_reads(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]
        self.assertEqual(self.api.get_run_result(run_id)["run_id"], run_id)
        self.assertEqual(self.api.get_run_artifacts(run_id)["run_id"], run_id)

        self.api.control_run(run_id, "cancel", actor_user_id="user-admin")
        retried = self.api.control_run(run_id, "retry", actor_user_id="user-admin")

        self.assertEqual(retried["status"], "queued")
        self.assertEqual(retried["phase"], "queued")
        self.assertEqual(retried["progress"], 0)
        self.assertIsNone(retried["result_summary_id"])
        self.assertIsNone(retried["artifact_manifest_id"])
        with self.assertRaises(KeyError):
            self.api.get_run_result(run_id)
        with self.assertRaises(KeyError):
            self.api.get_run_artifacts(run_id)

    def test_m9_3_retry_keeps_run_detail_refreshable_with_pending_manifest(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]

        self.api.control_run(run_id, "cancel", actor_user_id="user-admin")
        self.api.control_run(run_id, "retry", actor_user_id="user-admin")
        detail = self.api.get_run_detail(run_id)

        self.assertEqual(detail["run"]["status"], "queued")
        self.assertEqual(detail["run"]["phase"], "queued")
        self.assertIsNone(detail["result_summary"])
        self.assertEqual(detail["artifact_manifest"]["run_id"], run_id)
        self.assertEqual(
            detail["artifact_manifest"]["artifact_manifest_id"],
            f"artifact-manifest-{run_id}-retry-pending",
        )
        self.assertEqual(detail["artifact_manifest"]["status"], "pending")
        self.assertEqual(detail["artifact_manifest"]["artifacts"], [])
        self.assertEqual(detail["chain"]["artifact_manifest_id"], detail["artifact_manifest"]["artifact_manifest_id"])
        self.assertEqual(detail["download_base"], f"/api/runs/{run_id}/artifacts")

    def test_m9_3_cancel_after_retry_keeps_detail_refreshable_without_stale_outputs(self) -> None:
        submitted = self._submit_successful_run()
        run_id = submitted["run_id"]

        first_cancel = self.api.control_run(run_id, "cancel", actor_user_id="user-data")
        retried = self.api.control_run(run_id, "retry", actor_user_id="user-admin")
        self.assertIsNone(retried.get("cancelled_at"))
        self.assertIsNone(retried.get("cancelled_by"))
        second_cancel = self.api.control_run(run_id, "cancel", actor_user_id="user-admin")
        detail = self.api.get_run_detail(run_id)

        self.assertEqual(detail["run"]["status"], "cancelled")
        self.assertEqual(detail["run"]["phase"], "cancelled")
        self.assertTrue(first_cancel["cancelled_at"])
        self.assertTrue(second_cancel["cancelled_at"])
        self.assertEqual(second_cancel["cancelled_by"], "user-admin")
        self.assertEqual(detail["run"]["cancelled_at"], second_cancel["cancelled_at"])
        self.assertEqual(detail["run"]["cancelled_by"], "user-admin")
        self.assertIsNone(detail["run"]["result_summary_id"])
        self.assertIsNone(detail["run"]["artifact_manifest_id"])
        self.assertIsNone(detail["result_summary"])
        self.assertEqual(detail["artifact_manifest"]["run_id"], run_id)
        self.assertEqual(
            detail["artifact_manifest"]["artifact_manifest_id"],
            f"artifact-manifest-{run_id}-retry-pending",
        )
        self.assertEqual(detail["artifact_manifest"]["status"], "pending")
        self.assertEqual(detail["artifact_manifest"]["artifacts"], [])
        self.assertEqual(detail["chain"]["artifact_manifest_id"], detail["artifact_manifest"]["artifact_manifest_id"])
        with self.assertRaises(KeyError):
            self.api.get_run_result(run_id)
        with self.assertRaises(KeyError):
            self.api.get_run_artifacts(run_id)

    def test_m9_2_subscribe_run_state_stream_fails_closed_for_invalid_state_series(self) -> None:
        cases = (
            (
                "deleted",
                lambda run_id, manifest, artifact: self.api.soft_delete_run(run_id, actor_user_id="user-admin"),
                "run_deleted",
            ),
            (
                "missing",
                lambda run_id, manifest, artifact: (
                    manifest["artifacts"].remove(artifact),
                    self.repository.upsert_artifact_manifest(manifest),
                ),
                "visualization_state_series_missing",
            ),
            (
                "path_escape",
                lambda run_id, manifest, artifact: (
                    artifact.update({"path": "../escape.json"}),
                    self.repository.upsert_artifact_manifest(manifest),
                ),
                "artifact_path_escape",
            ),
            (
                "run_mismatch",
                self._write_mismatched_state_series_payload,
                "visualization_state_series_run_mismatch",
            ),
            (
                "malformed_json",
                lambda run_id, manifest, artifact: (
                    (Path(self.api.output_dir) / artifact["path"]).write_text("{", encoding="utf-8")
                ),
                "visualization_state_series_invalid",
            ),
            (
                "missing_frame_trace",
                self._remove_state_series_frame_trace,
                "visualization_state_series_invalid",
            ),
        )
        for label, mutate, expected_code in cases:
            with self.subTest(label):
                submitted = self._submit_successful_run()
                run_id = submitted["run_id"]
                manifest = self.api.get_run_artifacts(run_id)
                artifact = self._artifact_by_kind(manifest, "visualization_state_series")
                mutate(run_id, manifest, artifact)

                with self.assertRaises(BackendApiError) as ctx:
                    self.api.subscribe_run_state_stream(run_id)

                self.assertEqual(ctx.exception.code, expected_code)

    def _write_mismatched_state_series_payload(self, run_id: str, manifest: dict, artifact: dict) -> None:
        target = Path(self.api.output_dir) / artifact["path"]
        payload = json.loads(target.read_text(encoding="utf-8"))
        payload["run_id"] = f"{run_id}-other"
        target.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    def _remove_state_series_frame_trace(self, run_id: str, manifest: dict, artifact: dict) -> None:
        target = Path(self.api.output_dir) / artifact["path"]
        payload = json.loads(target.read_text(encoding="utf-8"))
        del payload["frames"][0]["trace"]
        target.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    def test_m9_2_http_state_stream_returns_sse_for_bearer_and_access_token_auth(self) -> None:
        server, thread, run_id, token = self._start_state_stream_server()
        try:
            for label, path, headers in (
                (
                    "bearer",
                    f"/api/runs/{run_id}/state-stream",
                    {"Authorization": f"Bearer {token}"},
                ),
                (
                    "access_token",
                    f"/api/runs/{run_id}/state-stream?access_token={token}",
                    {},
                ),
            ):
                with self.subTest(label):
                    response, body = self._http_request(server, "GET", path, headers=headers)

                    self.assertEqual(response.status, 200)
                    self.assertEqual(response.getheader("content-type"), "text/event-stream; charset=utf-8")
                    self.assertEqual(response.getheader("cache-control"), "no-cache")
                    self.assertIsNone(response.getheader("access-control-allow-origin"))
                    events = self._parse_sse_events(body.decode("utf-8"))
                    self.assertEqual(events[0][0], "run_status")
                    self.assertIn("state_frame", [event_type for event_type, _payload in events])
                    self.assertEqual(events[-1][0], "artifact_ready")
                    state_frame = next(payload for event_type, payload in events if event_type == "state_frame")
                    self.assertEqual(state_frame["schema_version"], "visualization-state-frame-v0")
                    self.assertEqual(state_frame["run_id"], run_id)
                    self.assertTrue(state_frame["artifact_id"].startswith("visualization_state_series-"))
                    self.assertEqual(state_frame["artifact_manifest_id"], f"artifact-manifest-{run_id}")
        finally:
            self._stop_http_server(server, thread)

    def test_m9_2_http_state_stream_requires_auth_and_maps_unknown_run_to_404(self) -> None:
        server, thread, run_id, token = self._start_state_stream_server()
        try:
            missing_auth, missing_body = self._http_request(server, "GET", f"/api/runs/{run_id}/state-stream")
            unknown, unknown_body = self._http_request(
                server,
                "GET",
                "/api/runs/run-missing/state-stream",
                headers={"Authorization": f"Bearer {token}"},
            )
            json_api_with_query_token, _ = self._http_request(
                server,
                "GET",
                f"/api/auth/session?access_token={token}",
            )

            self.assertEqual(missing_auth.status, 401)
            self.assertEqual(json.loads(missing_body.decode("utf-8"))["code"], "unauthorized")
            self.assertEqual(unknown.status, 404)
            self.assertEqual(json.loads(unknown_body.decode("utf-8"))["code"], "not_found")
            self.assertEqual(json_api_with_query_token.status, 401)
        finally:
            self._stop_http_server(server, thread)

    def _start_state_stream_server(self) -> tuple[Any, threading.Thread, str, str]:
        database_path = Path(self.tempdir.name) / "state-stream.sqlite3"
        connection = sqlite3.connect(database_path)
        try:
            initialize_database(connection)
            repository = ContractRepository(connection)
            api = BackendApi(repository, RecordingAdapter(), output_dir=Path(self.tempdir.name))
            token = api.login("admin", "admin")["session"]["token"]
            project = small_aircraft_support_project("project-aircraft-support-contract-001")
            saved = api.save_project(project)
            api.create_modeling_snapshot(saved["project_id"])
            plan = api.create_experiment_plan(saved["project_id"], {"name": "http state stream", "steps": 2})
            submitted = api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "single",
                }
            )
        finally:
            connection.close()
        server = create_backend_server(
            ("127.0.0.1", 0),
            repo_root=REPO_ROOT,
            database_path=database_path,
            output_dir=Path(self.tempdir.name),
        )
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        return server, thread, submitted["run_id"], token

    def _http_request(
        self,
        server: Any,
        method: str,
        path: str,
        *,
        headers: dict[str, str] | None = None,
    ) -> tuple[http.client.HTTPResponse, bytes]:
        connection = http.client.HTTPConnection(server.server_address[0], server.server_address[1], timeout=5)
        try:
            connection.request(method, path, headers=headers or {})
            response = connection.getresponse()
            return response, response.read()
        finally:
            connection.close()

    def _parse_sse_events(self, body: str) -> list[tuple[str, dict[str, Any]]]:
        events = []
        for block in body.strip().split("\n\n"):
            lines = block.splitlines()
            event_type = next(line.removeprefix("event: ").strip() for line in lines if line.startswith("event: "))
            data = next(line.removeprefix("data: ").strip() for line in lines if line.startswith("data: "))
            events.append((event_type, json.loads(data)))
        return events

    def _stop_http_server(self, server: Any, thread: threading.Thread) -> None:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)

    def test_backend_api_requires_explicit_actor_for_m7_lifecycle_and_download(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        artifact = manifest["artifacts"][0]
        cases = [
            ("archive", lambda: self.api.archive_run(run["run_id"])),
            ("delete", lambda: self.api.soft_delete_run(run["run_id"])),
            ("download", lambda: self.api.get_run_artifact_download(run["run_id"], artifact["artifact_id"])),
        ]

        for label, call in cases:
            with self.subTest(label):
                with self.assertRaises(BackendApiError) as ctx:
                    call()

                self.assertEqual(ctx.exception.code, "missing_actor")

        stored = self.api.get_run(run["run_id"])
        self.assertEqual(stored["lifecycle_status"], "active")
        self.assertEqual(self.repository.list_audit_events(resource_id=run["run_id"]), [])

    def test_backend_api_lifecycle_requires_admin_or_data_manager_actor(self) -> None:
        forbidden_run = self._submit_successful_run()

        for label, mutate in (
            ("archive", lambda: self.api.archive_run(forbidden_run["run_id"], actor_user_id="user-basic")),
            ("delete", lambda: self.api.soft_delete_run(forbidden_run["run_id"], actor_user_id="user-basic")),
        ):
            with self.subTest(label):
                with self.assertRaises(BackendApiError) as ctx:
                    mutate()

                self.assertEqual(ctx.exception.code, "forbidden")
                stored = self.api.get_run(forbidden_run["run_id"])
                self.assertEqual(stored["lifecycle_status"], "active")
                audit = self.repository.list_audit_events(resource_id=forbidden_run["run_id"])
                self.assertFalse(any(event["outcome"] == "allowed" for event in audit))

        admin_run = self._submit_successful_run()
        archived = self.api.archive_run(admin_run["run_id"], actor_user_id="user-admin")
        self.assertEqual(archived["lifecycle_status"], "archived")

        data_run = self._submit_successful_run()
        deleted = self.api.soft_delete_run(data_run["run_id"], actor_user_id="user-data")
        self.assertEqual(deleted["lifecycle_status"], "deleted")

    def test_experiment_plan_list_and_delete_soft_deletes_runs(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "visual replay cleanup", "steps": 2},
        )
        first_run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        second_run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        plans = self.api.list_experiment_plans(saved["project_id"])

        self.assertEqual([item["experiment_plan_id"] for item in plans["experiment_plans"]], [plan["experiment_plan_id"]])
        self.assertEqual(plans["experiment_plans"][0]["config"]["name"], "visual replay cleanup")
        self.assertEqual(plans["experiment_plans"][0]["run_count"], 2)
        self.assertEqual(
            [item["run_id"] for item in plans["experiment_plans"][0]["runs"]],
            [second_run["run_id"], first_run["run_id"]],
        )

        with self.assertRaises(BackendApiError) as forbidden_ctx:
            self.api.delete_experiment_plan(saved["project_id"], plan["experiment_plan_id"], actor_user_id="user-basic")
        self.assertEqual(forbidden_ctx.exception.code, "forbidden")

        deleted = self.api.delete_experiment_plan(
            saved["project_id"],
            plan["experiment_plan_id"],
            actor_user_id="user-data",
        )

        self.assertEqual(deleted["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(deleted["deleted"], True)
        self.assertEqual(deleted["soft_deleted_run_ids"], [first_run["run_id"], second_run["run_id"]])
        self.assertEqual(self.api.get_run(first_run["run_id"])["lifecycle_status"], "deleted")
        self.assertEqual(self.api.get_run(second_run["run_id"])["lifecycle_status"], "deleted")
        self.assertEqual(self.api.list_experiment_plans(saved["project_id"])["experiment_plans"], [])

    def test_backend_api_lifecycle_audit_failure_rolls_back_state(self) -> None:
        cases = [
            ("archive", lambda run_id: self.api.archive_run(run_id, actor_user_id="missing-user")),
            ("delete", lambda run_id: self.api.soft_delete_run(run_id, actor_user_id="missing-user")),
        ]
        for label, mutate in cases:
            with self.subTest(label):
                run = self._submit_successful_run()

                with self.assertRaises(sqlite3.IntegrityError):
                    mutate(run["run_id"])

                stored = self.api.get_run(run["run_id"])
                self.assertEqual(stored["lifecycle_status"], "active")
                self.assertEqual(self.repository.list_audit_events(resource_id=run["run_id"]), [])

    def test_backend_api_rejects_artifact_path_escape(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        manifest["artifacts"][0]["path"] = "../escape.json"
        self.repository.upsert_artifact_manifest(manifest)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.get_run_artifact_download(
                run["run_id"],
                manifest["artifacts"][0]["artifact_id"],
                actor_user_id="system",
            )

        self.assertEqual(ctx.exception.code, "artifact_path_escape")
        self.assertEqual(self.repository.list_audit_events(resource_id=run["run_id"]), [])

    def test_backend_api_rejects_artifact_hash_mismatch(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        artifact = manifest["artifacts"][0]
        target = Path(self.api.output_dir) / artifact["path"]
        target.write_text('{"tampered": true}\n', encoding="utf-8")

        with self.assertRaises(BackendApiError) as ctx:
            self.api.get_run_artifact_download(run["run_id"], artifact["artifact_id"], actor_user_id="system")

        self.assertEqual(ctx.exception.code, "artifact_hash_mismatch")
        self.assertEqual(self.repository.list_audit_events(resource_id=run["run_id"]), [])

    def test_backend_api_rejects_missing_artifact_file(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        artifact = manifest["artifacts"][0]
        target = Path(self.api.output_dir) / artifact["path"]
        target.unlink()

        with self.assertRaises(BackendApiError) as ctx:
            self.api.get_run_artifact_download(run["run_id"], artifact["artifact_id"], actor_user_id="system")

        self.assertEqual(ctx.exception.code, "artifact_missing")
        self.assertEqual(self.repository.list_audit_events(resource_id=run["run_id"]), [])

    def test_backend_api_rejects_deleted_run_artifact_download(self) -> None:
        run = self._submit_successful_run()
        manifest = self.api.get_run_artifacts(run["run_id"])
        artifact = manifest["artifacts"][0]
        self.api.soft_delete_run(run["run_id"], actor_user_id="user-admin")

        with self.assertRaises(BackendApiError) as ctx:
            self.api.get_run_artifact_download(run["run_id"], artifact["artifact_id"], actor_user_id="system")

        self.assertEqual(ctx.exception.code, "run_deleted")

    def test_run_service_submits_formal_monte_carlo_run_and_persists_projection_artifacts(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        branch_project = copy.deepcopy(project)
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "formal mc status",
                "steps": 2,
                "seed": branch_project["experiment"]["seed"],
                "projectJson": branch_project,
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 8,
                        "sweep": {
                            "failureRates": [0.06, 0.08],
                            "spareMultipliers": [0.75, 1.0],
                            "supportCapacities": [2, 3],
                        },
                    },
                    "spareShortfall": {"enabled": True},
                    "carryList": {"enabled": True},
                    "missionReliability": {"enabled": True},
                    "downtimeFactors": {"enabled": True},
                },
            },
        )
        service = RunService(self.repository, self.adapter, self.api.output_dir)

        submitted = service.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
            }
        )
        status = service.get_run_status(submitted["run_id"])
        stored_run = self.api.get_run(submitted["run_id"])
        manifest = self.api.get_run_artifacts(submitted["run_id"])
        kinds = {artifact["kind"] for artifact in manifest["artifacts"]}
        base_artifact = next(artifact for artifact in manifest["artifacts"] if artifact["kind"] == "monte_carlo_base")
        projection_artifacts = [
            artifact
            for artifact in manifest["artifacts"]
            if artifact["kind"].startswith("analysis_projection_")
        ]

        self.assertEqual(submitted["status"], "succeeded")
        self.assertEqual(submitted["run_type"], "monte_carlo")
        self.assertEqual(status["run_type"], "monte_carlo")
        self.assertEqual(stored_run["run_type"], "monte_carlo")
        self.assertEqual(status["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(kinds, M7_MONTE_CARLO_ARTIFACT_KINDS)
        self.assertEqual(len(projection_artifacts), 4)
        self.assertTrue(all(artifact["source_artifact_id"] == base_artifact["artifact_id"] for artifact in projection_artifacts))
        self.assertTrue(all(artifact["schema_version"] == "analysis-projection-v0" for artifact in projection_artifacts))

    def test_monte_carlo_run_uses_plan_large_sample_config(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        branch_project = copy.deepcopy(project)
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "canonical MC config",
                "steps": 4,
                "projectJson": branch_project,
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 5,
                        "sweep": {
                            "failureRates": [0.06, 0.08],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [2, 3],
                        },
                    }
                },
            },
        )

        status = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "mc_experiment_id": "mc-canonical-config",
            }
        )
        manifest = self.api.get_run_artifacts(status["run_id"])
        base_artifact = next(artifact for artifact in manifest["artifacts"] if artifact["kind"] == "monte_carlo_base")
        payload = json.loads((Path(self.tempdir.name) / base_artifact["path"]).read_text(encoding="utf-8"))

        self.assertEqual(payload["sample_count"], 5)
        self.assertEqual(payload["mc_experiment_id"], "mc-canonical-config")
        self.assertEqual(payload["sweep"]["failureRates"], [0.06, 0.08])
        self.assertEqual(payload["sweep"]["supportCapacities"], [2, 3])

    def test_run_service_hands_normalized_monte_carlo_config_to_adapter(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "adapter normalized mc config",
                "steps": 3,
                "parallelCores": 3,
                "projectJson": copy.deepcopy(project),
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 6,
                        "sweep": {
                            "failureRates": [0.05],
                            "spareMultipliers": [1.0, 1.2],
                            "supportCapacities": [2],
                        },
                    }
                },
            },
        )

        self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "mc_experiment_id": "mc-adapter-normalized",
            }
        )

        adapter_config = self.adapter.monte_carlo_run_calls[-1]["kwargs"].get("monte_carlo_config")
        self.assertIsNotNone(adapter_config)
        self.assertEqual(adapter_config["sample_count"], 6)
        self.assertEqual(adapter_config["parallel_cores"], 3)
        self.assertEqual(adapter_config["mc_experiment_id"], "mc-adapter-normalized")
        self.assertEqual(adapter_config["sweep"]["failureRates"], [0.05])
        self.assertEqual(adapter_config["sweep"]["spareMultipliers"], [1.0, 1.2])
        self.assertEqual(adapter_config["sweep"]["supportCapacities"], [2])

    def test_experiment_plan_rejects_invalid_parallel_cores_before_save(self) -> None:
        project = small_aircraft_support_project("project-invalid-parallel-cores")
        saved = self.api.save_project(project)
        before_plans = self.api.list_experiment_plans(saved["project_id"])["experiment_plans"]

        for value in (0, 1.5, 33, "bad"):
            with self.subTest(value=value), self.assertRaises(BackendApiError) as ctx:
                self.api.create_experiment_plan(
                    saved["project_id"],
                    {"name": "invalid parallel", "parallelCores": value},
                )
            self.assertEqual(ctx.exception.code, "bad_run_request")
            self.assertEqual(ctx.exception.details["field"], "parallelCores")

        self.assertEqual(self.api.list_experiment_plans(saved["project_id"])["experiment_plans"], before_plans)

    def test_lite_mesa_parallel_workers_preserve_seeded_sample_order_and_statistics(self) -> None:
        project = small_aircraft_support_project("project-lite-parallel-reproducible")
        serial = self.api.run_lite_mesa_analysis(
            project,
            analysis_type="mission_reliability",
            settings={"samples": 3, "seed": 20260717, "parallelCores": 1},
        )
        parallel = self.api.run_lite_mesa_analysis(
            project,
            analysis_type="mission_reliability",
            settings={"samples": 3, "seed": 20260717, "parallelCores": 3},
        )

        self.assertEqual(serial["worker_count"], 1)
        self.assertEqual(parallel["worker_count"], 3)
        self.assertEqual(serial["seed_list"], [20260717, 20260718, 20260719])
        self.assertEqual(parallel["seed_list"], serial["seed_list"])
        self.assertEqual(parallel["aggregate_metrics"], serial["aggregate_metrics"])

    def test_lite_mesa_rejects_invalid_parallel_cores_before_run(self) -> None:
        project = small_aircraft_support_project("project-lite-invalid-parallel")
        for value in (0, 1.5, 33, "bad"):
            with self.subTest(value=value), self.assertRaises(BackendApiError) as ctx:
                self.api.run_lite_mesa_analysis(
                    project,
                    analysis_type="mission_reliability",
                    settings={"samples": 2, "seed": 20260717, "parallelCores": value},
                )
            self.assertEqual(ctx.exception.code, "bad_run_request")
            self.assertEqual(ctx.exception.details["field"], "settings.parallelCores")

    def test_monte_carlo_run_rejects_request_level_samples_and_sweep(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        branch_project = copy.deepcopy(project)
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "reject request MC config",
                "steps": 4,
                "projectJson": branch_project,
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 4,
                        "sweep": {
                            "failureRates": [0.07],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [3],
                        },
                    }
                },
            },
        )
        before_counts = self._run_side_effect_counts()
        before_compile_calls = len(self.adapter.compile_calls)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "monte_carlo",
                    "sample_count": 99,
                    "samples": 99,
                    "sweep": {"supportCapacities": [9]},
                    "monte_carlo": {"samples": 99},
                }
            )

        self.assertEqual(ctx.exception.code, "bad_run_request")
        self.assertIn("ExperimentPlan.config.analysisRequests.largeSample", str(ctx.exception))
        self.assertIn("sample_count", ctx.exception.details.get("fields", []))
        self.assertIn("sweep", ctx.exception.details.get("fields", []))
        self.assertEqual(self._run_side_effect_counts(), before_counts)
        self.assertEqual(len(self.adapter.compile_calls), before_compile_calls)
        self.assertEqual(self.adapter.monte_carlo_run_calls, [])

    def test_monte_carlo_bad_request_wins_before_compile_gate_model_family(self) -> None:
        project = self._fixture("aircraft_support_v1_project.json")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "bad request before aviation compile gate",
                "steps": 1,
                "projectJson": copy.deepcopy(project),
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 4,
                        "sweep": {
                            "failureRates": [0.07],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [3],
                        },
                    }
                },
            },
        )
        before_counts = self._run_side_effect_counts()
        before_compile_calls = len(self.adapter.compile_calls)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aviation_support",
                    "run_type": "monte_carlo",
                    "sample_count": 99,
                }
            )

        self.assertEqual(ctx.exception.code, "bad_run_request")
        self.assertIn("ExperimentPlan.config.analysisRequests.largeSample", str(ctx.exception))
        self.assertNotEqual(ctx.exception.code, "unsupported_model_family")
        self.assertEqual(self._run_side_effect_counts(), before_counts)
        self.assertEqual(len(self.adapter.compile_calls), before_compile_calls)
        self.assertEqual(self.adapter.monte_carlo_run_calls, [])

    def test_monte_carlo_config_rejects_fractional_integer_fields(self) -> None:
        cases = [
            ("fractional samples", {"samples": 2.5}, "analysisRequests.largeSample.samples"),
            (
                "fractional support capacity",
                {"sweep": {"failureRates": [0.07], "spareMultipliers": [1.0], "supportCapacities": [2.5]}},
                "analysisRequests.largeSample.sweep.supportCapacities",
            ),
        ]
        for label, override, expected_field in cases:
            with self.subTest(label):
                project = small_aircraft_support_project("project-aircraft-support-contract-001")
                branch_project = copy.deepcopy(project)
                large_sample = {
                    "enabled": True,
                    "samples": 4,
                    "sweep": {
                        "failureRates": [0.07],
                        "spareMultipliers": [1.0],
                        "supportCapacities": [3],
                    },
                }
                if "samples" in override:
                    large_sample["samples"] = override["samples"]
                if "sweep" in override:
                    large_sample["sweep"] = override["sweep"]
                saved = self.api.save_project(project)
                self.api.create_modeling_snapshot(saved["project_id"])
                plan = self.api.create_experiment_plan(
                    saved["project_id"],
                    {
                        "name": f"reject {label}",
                        "steps": 4,
                        "projectJson": branch_project,
                        "analysisRequests": {"largeSample": large_sample},
                    },
                )
                before_counts = self._run_side_effect_counts()
                before_compile_calls = len(self.adapter.compile_calls)

                with self.assertRaises(BackendApiError) as ctx:
                    self.api.submit_run(
                        {
                            "project_id": saved["project_id"],
                            "experiment_plan_id": plan["experiment_plan_id"],
                            "model_family": "aircraft_support_v1",
                            "run_type": "monte_carlo",
                        }
                    )

                self.assertEqual(ctx.exception.code, "bad_run_request")
                self.assertEqual(
                    ctx.exception.details.get("field") or ctx.exception.details.get("field_path"),
                    expected_field,
                )
                self.assertEqual(self._run_side_effect_counts(), before_counts)
                self.assertEqual(len(self.adapter.compile_calls), before_compile_calls)
                self.assertEqual(self.adapter.monte_carlo_run_calls, [])

    def test_run_service_rejects_missing_model_family_on_canonical_submit(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "missing family", "steps": 2})
        service = RunService(self.repository, self.adapter, self.api.output_dir)

        with self.assertRaises(RunServiceError) as ctx:
            service.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "run_type": "single",
                }
            )

        self.assertEqual(ctx.exception.code, "bad_run_request")

    def test_run_service_rejects_retired_formal_model_families_before_compile(self) -> None:
        created = self._create_imported_sample_project()
        project = created["project"]
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "retired model family", "steps": 2, "projectJson": copy.deepcopy(project)},
        )
        service = RunService(self.repository, self.adapter, self.api.output_dir)

        before_counts = self._run_side_effect_counts()

        for model_family, run_type in (("aviation_support", "single"), ("aviation_support", "monte_carlo"), ("smoke", "single")):
            with self.subTest(model_family=model_family, run_type=run_type):
                with self.assertRaises(RunServiceError) as ctx:
                    service.submit_run(
                        {
                            "project_id": saved["project_id"],
                            "experiment_plan_id": plan["experiment_plan_id"],
                            "model_family": model_family,
                            "run_type": run_type,
                            "formal_run": True,
                        }
                    )

                self.assertEqual(ctx.exception.code, "retired_model_family")
                self.assertEqual(ctx.exception.details["model_family"], model_family)
                self.assertEqual(ctx.exception.details["replacement_model_family"], "aircraft_support_v1")
                self.assertIn(model_family, ctx.exception.details["retired_model_families"])

        self.assertEqual(self._run_side_effect_counts(), before_counts)
        self.assertEqual(self.adapter.compile_calls, [])
        self.assertEqual(self.adapter.run_calls, [])
        self.assertEqual(self.adapter.monte_carlo_run_calls, [])

    def test_run_service_submits_aircraft_support_v1_formal_single_run_and_exposes_behavior_scope(self) -> None:
        created = self._create_imported_sample_project()
        project = copy.deepcopy(created["project"])
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "m9.7.4 aircraft support single", "projectJson": copy.deepcopy(project)},
        )
        service = RunService(self.repository, self.adapter, self.api.output_dir)

        submitted = service.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
                "formal_run": True,
            }
        )
        status = service.get_run_status(submitted["run_id"])
        result = self.api.get_run_result(submitted["run_id"])
        manifest = self.api.get_run_artifacts(submitted["run_id"])
        state_artifact = self._artifact_by_kind(manifest, "visualization_state_series")
        report_artifact = self._artifact_by_kind(manifest, "report")
        state_payload = json.loads((Path(self.api.output_dir) / state_artifact["path"]).read_text(encoding="utf-8"))
        report_payload = json.loads((Path(self.api.output_dir) / report_artifact["path"]).read_text(encoding="utf-8"))

        self.assertEqual(submitted["status"], "succeeded")
        self.assertEqual(submitted["phase"], "completed")
        self.assertEqual(submitted["model_family"], "aircraft_support_v1")
        self.assertEqual(status["model_family"], "aircraft_support_v1")
        self.assertEqual(result["model_family"], "aircraft_support_v1")
        self.assertEqual(result["run_id"], submitted["run_id"])
        self.assertEqual(submitted["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(state_payload["model_family"], "aircraft_support_v1")
        self.assertEqual(state_payload["run_id"], submitted["run_id"])
        self.assertEqual(state_payload["scenario_id"], submitted["scenario_id"])
        self.assertIn("sortie_completion_rate", result["metrics"])
        scope = report_payload["m9_7_4_behavior_scope"]
        self.assertIn("supportResources[].quantity", scope["behavior_driving_fields"])
        self.assertIn("components[].failureDistribution", scope["behavior_driving_fields"])
        self.assertIn("transportPolicies[]", scope["behavior_driving_fields"])
        self.assertIn("missionProfile.periodicTasks", scope["behavior_driving_fields"])
        self.assertNotIn("reliabilityBlockDiagram", scope["behavior_driving_fields"])
        self.assertNotIn("experiment.steps", scope["behavior_driving_fields"])
        self.assertEqual(
            set(scope["fail_closed_fields"]),
            {
                "combatUnit.members[].preLifeCalendarDays",
                "combatUnit.members[].preLifeFlightHours",
                "combatUnit.members[].preLifeTakeoffLandingCount",
                "supportActivities[].aircraftModel",
                "supportActivities[].equipmentId",
                "supportActivities[].calendarDayInterval",
                "supportActivities[].runHourInterval",
                "supportActivities[].takeoffLandingInterval",
                "support_network.nodes[].organization_node_id",
                "support_network.organization_graph.nodes[]",
                "support_network.organization_graph.runtime_mode",
                "support_network.organization_graph.parent_edges[]",
                "support_network.organization_graph.lateral_edges[]",
                "support_network.organization_graph.transport_policies[]",
            },
        )
        self.assertEqual(scope["m9_7_4_coverage_hardening_fields"], [])
        self.assertEqual(len(self.adapter.compile_calls), 1)
        self.assertEqual(self.adapter.compile_calls[0][1], "aircraft_support_v1")
        self.assertEqual(len(self.adapter.run_calls), 1)
        self.assertEqual(self.adapter.run_calls[0][0]["simulation_model"]["family"], "aircraft_support_v1")

    def test_lite_mesa_analysis_runs_in_memory_without_formal_side_effects(self) -> None:
        self.assertFalse(hasattr(self.api, "run_independent_mesa_visualization"))
        before = self._run_side_effect_counts()

        payload = self.api.run_lite_mesa_analysis(
            small_aircraft_support_project("project-lite-mesa-contract"),
            analysis_type="mission_reliability",
            settings={"samples": 2, "seed": 20260705, "maxTimeWindow": 1},
        )

        self.assertEqual(payload["status"], "session_complete")
        self.assertEqual(payload["source"], "lite_mesa_aircraft_support_v1")
        self.assertEqual(payload["model_family"], "aircraft_support_v1")
        self.assertEqual(payload["analysis_type"], "mission_reliability")
        self.assertEqual(payload["project_id"], "project-lite-mesa-contract")
        self.assertEqual(payload["sample_count"], 2)
        self.assertEqual(payload["requested_sample_count"], 2)
        self.assertEqual(payload["completed_sample_count"], 2)
        self.assertEqual(payload["failed_sample_count"], 0)
        self.assertEqual(payload["metric_moments"]["variance_denominator"], "n-1")
        self.assertEqual(payload["metric_moments"]["total_sample_count"], 2)
        self.assertEqual(payload["metric_moments"]["successful_sample_count"], 2)
        self.assertEqual(payload["metric_moments"]["failed_sample_count"], 0)
        self.assertTrue(all("unit" in metric for metric in payload["metric_moments"]["metrics"]))
        self.assertEqual(payload["seed_list"], [20260705, 20260706])
        self.assertEqual(payload["sample_timeout_seconds"], 60)
        self.assertEqual(payload["session_timeout_seconds"], 180)
        self.assertEqual([item["status"] for item in payload["sample_diagnostics"]], ["ok", "ok"])
        self.assertGreaterEqual(payload["timings"]["compile_seconds"], 0)
        self.assertGreater(payload["timings"]["sample_execution_seconds"], 0)
        self.assertGreaterEqual(payload["timings"]["aggregation_seconds"], 0)
        self.assertGreaterEqual(payload["timings"]["projection_seconds"], 0)
        self.assertGreater(payload["timings"]["total_seconds"], 0)
        self.assertTrue(payload["rows"])
        self.assertEqual(payload["rows"], payload["wave_rows"])
        self.assertEqual(payload["wave_rows"][0]["dayIndex"], 1)
        self.assertEqual(payload["wave_rows"][0]["sampleIndex"], 0)
        self.assertEqual(payload["wave_rows"][0]["sampleLabel"], "样本 1")
        self.assertIn("meanMissionSuccessRate", payload["wave_rows"][0])
        self.assertNotIn("seed", payload["wave_rows"][0])
        self.assertEqual([field["key"] for field in payload["result_fields"]], [
            "sortie_rate", "wave_success_rate", "period_completion_probability", "period_duration_days"
        ])
        self.assertEqual(payload["metrics"], [
            *[[field["label"], field["display_value"]] for field in payload["result_fields"]],
            ["仿真总次数", "2"],
            ["成功次数", "0"],
        ])
        self.assertNotIn("任务剖面可靠性", json.dumps(payload["metrics"], ensure_ascii=False))
        self.assertEqual(payload["visualization_state_series"]["run_id"], payload["run_id"])
        self.assertEqual(len(payload["visualization_state_series"]["frames"]), 1)
        self.assertEqual(
            payload["visualization_state_series"]["organization_dispatch_summary"]["summary_scope"],
            "representative_sample",
        )
        self.assertEqual(self._run_side_effect_counts(), before)

    def test_lite_mesa_analysis_applies_scenario_composition_before_compile(self) -> None:
        project = strip_project_sweep(small_aircraft_support_project("project-lite-mesa-composed"))
        project["scenarioComposition"] = {
            "schemaVersion": "scenario-composition-v0",
            "overrides": [
                {"path": "supportResources.2.quantity", "valueType": "number", "value": 9}
            ],
        }

        payload = self.api.run_lite_mesa_analysis(
            project,
            analysis_type="spare_shortfall",
            settings={"samples": 1, "seed": 20260705},
        )

        self.assertEqual(payload["status"], "session_complete")
        compiled_project = self.adapter.compile_calls[-1][0]
        self.assertNotIn("scenarioComposition", compiled_project)
        self.assertEqual(compiled_project["supportResources"][2]["quantity"], 9)

        compile_result = self.adapter.compile_scenario_with_gate(compiled_project, model_family="aircraft_support_v1")
        simulation_inputs = compile_result["scenario"]["simulation_inputs"]
        resource_product_id = compiled_project["components"][0]["productId"]
        self.assertEqual(
            simulation_inputs["support_network"]["nodes"][0]["inventory"][resource_product_id],
            9,
        )

    def test_aircraft_support_v1_duration_stops_at_last_explicit_periodic_mission_day(self) -> None:
        scenario = self.adapter.compile_scenario(
            periodic_three_day_aircraft_support_project("project-three-day-duration"),
            model_family="aircraft_support_v1",
        )

        inputs = scenario["simulation_inputs"]

        self.assertEqual(inputs["time"]["duration_minutes"], 3 * 24 * 60)
        self.assertEqual(inputs["mission_profile"]["duration_minutes"], 3 * 24 * 60)

    def test_lite_mesa_mission_and_downtime_use_last_periodic_mission_day_duration(self) -> None:
        project = periodic_three_day_aircraft_support_project("project-three-day-lite-analysis")

        mission_payload = self.api.run_lite_mesa_analysis(
            copy.deepcopy(project),
            analysis_type="mission_reliability",
            settings={"samples": 1, "seed": 20260705},
        )
        downtime_payload = self.api.run_lite_mesa_analysis(
            copy.deepcopy(project),
            analysis_type="downtime_factors",
            settings={"samples": 1, "seed": 20260705},
        )

        self.assertEqual(mission_payload["aggregate_metrics"]["simulation_days"], 3.0)
        self.assertEqual(downtime_payload["aggregate_metrics"]["simulation_days"], 3.0)
        self.assertEqual(sorted({row["dayIndex"] for row in mission_payload["wave_rows"]}), [1, 2, 3])
        self.assertEqual(mission_payload["period_duration_days"], 3.0)
        self.assertEqual(mission_payload["period_total_samples"], 1)
        self.assertEqual(
            mission_payload["successful_samples"] + mission_payload["period_failed_samples"],
            mission_payload["period_total_samples"],
        )

    def test_lite_mesa_downtime_event_snapshots_use_model_event_log_snapshots(self) -> None:
        snapshots = _lite_mesa_downtime_event_snapshots(
            [
                {
                    "sample_index": 0,
                    "seed": 20260621,
                    "events": [
                        {
                            "time": 42,
                            "event": "spare_shortage",
                            "message": "repair blocked by hydraulic pump shortage",
                            "details": {"failure_mode": "must-not-project"},
                            "snapshot": {
                                "aircraft_state": {
                                    "summary": {"available_aircraft": 1, "failed_count": 1, "repairing_count": 1},
                                    "aircraft": [{
                                        "tail_number": "J15-101",
                                        "state": "maintenance",
                                        "failure_mode": "must-not-project-from-state",
                                    }],
                                },
                                "support_resources": [
                                    {
                                        "resource_id": "carrier-deck",
                                        "name": "航母飞行甲板",
                                        "personnel_in_use": 1,
                                        "personnel_capacity": 2,
                                        "equipment_in_use": 1,
                                        "equipment_capacity": 2,
                                        "inventory": {"液压泵": 0},
                                    }
                                ],
                                "spare_shortages": [
                                    {
                                        "spare_type": "液压泵",
                                        "required_quantity": 1,
                                        "available_quantity": 0,
                                        "job_id": "repair-J15-101",
                                    }
                                ],
                                "active_jobs": [
                                    {
                                        "job_id": "repair-J15-101",
                                        "kind": "repair",
                                        "state": "waiting",
                                        "task": "更换液压泵",
                                        "tail_number": "J15-101",
                                    }
                                ],
                            },
                        }
                    ],
                    "frames": [
                        {
                            "simulation_time": 60,
                            "events": [{"event_type": "state_frame", "message": "state frame sampled"}],
                        }
                    ],
                }
            ],
            4,
        )

        self.assertEqual(len(snapshots), 1)
        self.assertEqual(snapshots[0]["source"], "model_event_log")
        self.assertEqual(snapshots[0]["event_type"], "spare_shortage")
        self.assertEqual(snapshots[0]["event_label"], "备件短缺")
        self.assertEqual(snapshots[0]["simulation_time"], 42.0)
        self.assertEqual(snapshots[0]["spare_shortages"][0]["spare_type"], "液压泵")
        self.assertEqual(snapshots[0]["job_node"]["job_id"], "repair-J15-101")
        self.assertNotIn("failure_mode", str(snapshots[0]))
        self.assertNotIn("must-not-project", str(snapshots[0]))

    def test_lite_mesa_mission_reliability_excludes_legacy_metrics_and_missing_period_is_explicit(self) -> None:
        result = _lite_mesa_mission_reliability_result(
            {"data": {"mission_success_probability": 0.8, "sortie_rate": 0.4}},
            [
                {"seed": 1, "metrics": {"mission_success_rate": 0.8, "sortie_rate": 0.4, "ready_rate": 0.5, "failed_sorties": 2}},
                {"seed": 2, "metrics": {"mission_success_rate": 0.6, "sortie_rate": 0.3, "ready_rate": 0.75, "failed_sorties": 3}},
            ],
            {"maxTimeWindow": 12},
        )

        self.assertEqual(result["metrics"], [
            ["出动架次率", "0.400"],
            ["波次成功率", "80%"],
            ["整周期任务可靠度", "0%"],
            ["任务周期", "--"],
            ["仿真总次数", "2"],
            ["成功次数", "0"],
        ])
        self.assertNotIn("任务失败次数", {label for label, _value in result["metrics"]})

    def test_lite_mesa_mission_reliability_rows_preserve_every_sample_wave(self) -> None:
        result = _lite_mesa_mission_reliability_result(
            {"data": {"mission_success_probability": 0.8, "sortie_rate": 0.4}},
            [
                {
                    "seed": 1,
                    "metrics": {"failed_sorties": 0, "ready_rate": 1},
                    "mission_wave_reliability": [
                        {"dayIndex": 1, "waveIndex": 1, "plannedSorties": 2, "launchedSorties": 2, "successfulSorties": 1, "plannedWaves": 1, "successfulWaves": 1, "missionSuccessRate": 1, "sortieRate": 1},
                        {"dayIndex": 1, "waveIndex": 2, "plannedSorties": 2, "launchedSorties": 1, "successfulSorties": 0, "plannedWaves": 1, "successfulWaves": 0, "missionSuccessRate": 0, "sortieRate": 0.5},
                    ],
                },
                {
                    "seed": 2,
                    "metrics": {"failed_sorties": 0, "ready_rate": 1},
                    "mission_wave_reliability": [
                        {"dayIndex": 1, "waveIndex": 1, "plannedSorties": 6, "launchedSorties": 4, "successfulSorties": 0, "plannedWaves": 1, "successfulWaves": 0, "missionSuccessRate": 0, "sortieRate": 4 / 6},
                    ],
                },
            ],
            {"maxTimeWindow": 1},
        )

        self.assertEqual(result["rows"], result["wave_rows"])
        self.assertEqual([row["waveKey"] for row in result["rows"]], ["d1-w1", "d1-w2", "d1-w1"])
        self.assertEqual([row["sampleIndex"] for row in result["rows"]], [0, 0, 1])
        self.assertEqual([row["sampleLabel"] for row in result["rows"]], ["样本 1", "样本 1", "样本 2"])
        self.assertAlmostEqual(result["rows"][0]["plannedSorties"], 2)
        self.assertAlmostEqual(result["rows"][0]["successfulSorties"], 1)
        self.assertAlmostEqual(result["rows"][0]["plannedWaves"], 1)
        self.assertAlmostEqual(result["rows"][0]["successfulWaves"], 1)
        self.assertAlmostEqual(result["rows"][0]["meanMissionSuccessRate"], 1)
        self.assertAlmostEqual(result["rows"][0]["meanSortieRate"], 1)
        self.assertEqual(result["rows"][1]["meanMissionSuccessRate"], 0)
        self.assertEqual(result["rows"][2]["meanMissionSuccessRate"], 0)
        self.assertTrue(all(0 <= row["meanMissionSuccessRate"] <= 1 for row in result["rows"]))
        self.assertNotIn("seed", result["rows"][0])

    def test_task_reliability_result_fields_keep_percent_and_period_contract(self) -> None:
        fields = build_task_reliability_result_fields(
            sortie_rate=0.81234,
            wave_success_rate="87.5%",
            period_completion_probability=0.923,
            period_duration_days=None,
        )

        self.assertEqual([field["display_value"] for field in fields], ["0.812", "87.5%", "92.3%", "--"])
        self.assertEqual(format_reliability_percent("92.3%"), "92.3%")

    def test_task_reliability_result_fields_define_half_even_boundary_display_values(self) -> None:
        fields = build_task_reliability_result_fields(
            sortie_rate=0.8125,
            wave_success_rate=0.8,
            period_completion_probability=0.9225,
            period_duration_days=2.125,
        )

        self.assertEqual([field["display_value"] for field in fields], ["0.812", "80%", "92.2%", "2.12 天"])

    def test_task_reliability_result_fields_use_decimal_half_even_at_binary_negative_and_ratio_boundaries(self) -> None:
        fields = build_task_reliability_result_fields(
            sortie_rate=-0.8125,
            wave_success_rate=0,
            period_completion_probability=1,
            period_duration_days=2.675,
        )

        self.assertEqual([field["display_value"] for field in fields], ["-0.812", "0%", "100%", "2.68 天"])
        self.assertEqual([field["unit"] for field in fields], ["", "%", "%", "天"])

    def test_lite_mesa_spare_shortfall_reports_transport_delay_hours_and_repair_cancellations(self) -> None:
        result = _lite_mesa_spare_shortfall_result(
            {
                "data": [
                    {
                        "aircraft_model": "J-15",
                        "spare_type": "航电模块",
                        "fill_rate": 0.55,
                        "risk_level": "high",
                    }
                ]
            },
            {
                "planned_sorties": 28,
                "shortage_events": 6222,
                "mean_transport_delay": 1.5,
                "spare_fill_rate": 0.55,
            },
            [
                {"metrics": {"cancelled_sorties": 2}},
                {"metrics": {"cancelled_sorties": 3}},
            ],
        )

        self.assertEqual(
            result["metrics"],
            [
                ["发生缺件备件", "1"],
                ["平均备件延误时间(h)", "1.50"],
                ["最高缺件备件", "航电模块"],
                ["因维修延误导致的任务取消次数", "5"],
            ],
        )
        self.assertEqual(result["rows"][0]["meanTransportDelayHours"], 1.5)
        self.assertNotIn("shortage", result["rows"][0])

    def test_lite_mesa_spare_pages_scope_rows_to_aircraft_airport_support_inventory(self) -> None:
        project = small_aircraft_support_project("project-lite-spare-scope")
        project["airports"] = [
            {"id": "carrier-deck", "name": "航母飞行甲板", "supportNodeId": "carrier-deck"},
            {"id": "forward-sea-base", "name": "前出海上保障点", "supportNodeId": "forward-sea-base"},
        ]
        project["combatUnit"]["members"][0]["airport"] = "航母飞行甲板"
        project["combatUnit"]["members"][0]["deploymentLocation"] = "前出海上保障点"
        project["products"].extend([
            {"id": "product-engine-spare", "name": "发动机备件", "kind": "spare"},
            {"id": "product-hydraulic-spare", "name": "液压备件", "kind": "spare"},
            {"id": "product-avionics-spare", "name": "航电模块", "kind": "spare"},
            {"id": "product-forward-spare", "name": "前出备件", "kind": "spare"},
            {"id": "product-stock-spare", "name": "仓库备件", "kind": "spare"},
        ])
        project["supportNodes"] = [
            {
                "id": "carrier-deck",
                "name": "基地",
                "personnelCapacity": 4,
                "equipmentCapacity": 4,
                "inventory": {"product-engine-spare": 4, "product-hydraulic-spare": 5, "product-avionics-spare": 6},
            },
            {
                "id": "forward-sea-base",
                "name": "中继",
                "personnelCapacity": 4,
                "equipmentCapacity": 4,
                "inventory": {"product-forward-spare": 9},
            },
            {
                "id": "carrier-stock",
                "name": "仓库",
                "personnelCapacity": 4,
                "equipmentCapacity": 4,
                "inventory": {"product-stock-spare": 12},
            },
        ]
        project["supportOrganization"] = {}
        project["components"].extend([
            {"id": "engine-spare", "name": "发动机控制模块", "productId": "product-engine-spare", "parentId": "whole-aircraft", "aircraftModel": "J-15", "productType": "LRU", "failureDistribution": {"distributionType": "exponential", "parameters": "lambda=0.01"}},
            {"id": "hydraulic-spare", "name": "液压执行器", "productId": "product-hydraulic-spare", "parentId": "whole-aircraft", "aircraftModel": "J-15", "productType": "LRU", "failureDistribution": {"distributionType": "exponential", "parameters": "lambda=0.01"}},
            {"id": "avionics-spare", "name": "航电模块", "productId": "product-avionics-spare", "parentId": "whole-aircraft", "aircraftModel": "J-15", "productType": "LRU", "failureDistribution": {"distributionType": "exponential", "parameters": "lambda=0.01"}},
        ])

        shortfall = self.api.run_lite_mesa_analysis(
            project,
            analysis_type="spare_shortfall",
            settings={"samples": 1, "seed": 20260705},
        )
        carry = self.api.run_lite_mesa_analysis(
            project,
            analysis_type="carry_list",
            settings={"samples": 1, "seed": 20260705},
        )

        expected_types = ["发动机备件", "液压备件", "航电模块"]
        self.assertEqual([row["spareType"] for row in shortfall["rows"]], expected_types)
        self.assertEqual([row["spareType"] for row in carry["rows"]], expected_types)
        self.assertTrue(all(row["productId"] for row in shortfall["rows"] + carry["rows"]))
        self.assertTrue(all("aircraftModel" in row for row in carry["rows"]))
        self.assertTrue(all("lifeLimited" in row and "lifeLandings" in row and "lifeHours" in row for row in carry["rows"]))
        self.assertIn(["备件满足率下限", "0.90"], carry["metrics"])
        self.assertNotIn("aircraft_support_v1_spares", {row["spareType"] for row in shortfall["rows"] + carry["rows"]})
        self.assertNotIn("前出备件", {row["spareType"] for row in shortfall["rows"] + carry["rows"]})
        self.assertNotIn("仓库备件", {row["spareType"] for row in shortfall["rows"] + carry["rows"]})

    def test_lite_mesa_carry_list_uses_ratio_of_totals_and_handles_zero_carried_total(self) -> None:
        settings = {"missionConfidenceTarget": 0.9}
        result = _lite_mesa_carry_list_result(
            {
                "data": [
                    {"product_id": "spare-a", "recommended_quantity": 1, "used_quantity": 1, "carried_quantity": 1},
                    {"product_id": "spare-b", "recommended_quantity": 9, "used_quantity": 1, "carried_quantity": 9},
                ]
            },
            {"shortage_events": 0},
            [{}],
            settings,
        )

        self.assertEqual(result["spare_used_total"], 2)
        self.assertEqual(result["spare_carried_total"], 10)
        self.assertAlmostEqual(result["overall_spare_utilization"], 0.2)
        self.assertIn(["总体备件利用率", "20.00%"], result["metrics"])
        self.assertAlmostEqual(result["rows"][1]["utilization"], 1 / 9)

        zero_total = _lite_mesa_carry_list_result(
            {"data": [{"product_id": "spare-zero", "recommended_quantity": 0, "used_quantity": 0, "carried_quantity": 0}]},
            {"shortage_events": 0},
            [{}],
            settings,
        )
        self.assertIsNone(zero_total["overall_spare_utilization"])
        self.assertEqual(zero_total["overall_spare_utilization_status"], "zero_carried")
        self.assertIsNone(zero_total["rows"][0]["utilization"])
        self.assertIn(["总体备件利用率", "--"], zero_total["metrics"])

        zero_used = _lite_mesa_carry_list_result(
            {"data": [{"product_id": "spare-idle", "recommended_quantity": 4, "used_quantity": 0, "carried_quantity": 4}]},
            {"shortage_events": 0},
            [{}],
            settings,
        )
        self.assertIn(["总体备件利用率", "0.00%"], zero_used["metrics"])

        over_capacity = _lite_mesa_carry_list_result(
            {"data": [{"product_id": "spare-hot", "recommended_quantity": 1, "used_quantity": 3, "carried_quantity": 1}]},
            {"shortage_events": 0},
            [{}],
            settings,
        )
        self.assertIn(["总体备件利用率", "300.00%"], over_capacity["metrics"])

        missing_raw = _lite_mesa_carry_list_result(
            {"data": [{"product_id": "legacy", "recommended_quantity": 1, "utilization": 0.5}]},
            {"shortage_events": 0},
            [{}],
            settings,
        )
        self.assertIsNone(missing_raw["overall_spare_utilization"])
        self.assertEqual(missing_raw["overall_spare_utilization_status"], "data_unavailable")
        self.assertIn(["总体备件利用率", "数据不可用"], missing_raw["metrics"])

        empty_projection = _lite_mesa_carry_list_result(
            {"data": []},
            {"spare_consumed_total": 5, "shortage_events": 2},
            [{}],
            settings,
        )
        self.assertEqual(empty_projection["rows"], [])
        self.assertIsNone(empty_projection["spare_used_total"])
        self.assertIsNone(empty_projection["spare_carried_total"])
        self.assertEqual(empty_projection["overall_spare_utilization_status"], "data_unavailable")
        self.assertIn(["总体备件利用率", "数据不可用"], empty_projection["metrics"])

    def test_aircraft_support_spare_projection_deduplicates_minute_shortage_events(self) -> None:
        projections = self.adapter._aircraft_support_v1_analysis_projections(
            {"planned_sorties": 4, "spare_fill_rate": 1.0, "spare_utilization": 0.0},
            "base-artifact",
            samples=[
                {
                    "sample_index": 0,
                    "events": [
                        {
                            "event": "spare_shortage",
                            "details": {
                                "job_id": "job-001",
                                "aircraft_model": "J-15",
                                "resource_id": "carrier-deck",
                                "spare_type": "航电模块",
                                "required_quantity": 1,
                            },
                        },
                        {
                            "event": "spare_shortage",
                            "details": {
                                "job_id": "job-001",
                                "aircraft_model": "J-15",
                                "resource_id": "carrier-deck",
                                "spare_type": "航电模块",
                                "required_quantity": 1,
                            },
                        },
                        {
                            "event": "spare_consumed",
                            "details": {
                                "job_id": "job-001",
                                "aircraft_model": "J-15",
                                "resource_id": "carrier-deck",
                                "spare_type": "航电模块",
                                "quantity": 1,
                            },
                        },
                    ],
                }
            ],
            simulation_inputs={
                "mission_profile": {
                    "airports": [{"id": "carrier-deck", "name": "航母飞行甲板", "supportNodeId": "carrier-deck"}],
                },
                "aircraft": {"assets": [{"tail_number": "J15-001", "model": "J-15", "airport": "航母飞行甲板"}]},
                "support_network": {
                    "nodes": [
                        {
                            "id": "carrier-deck",
                            "name": "基地",
                            "inventory": {"航电模块": 6},
                        }
                    ]
                },
            },
        )

        shortfall_row = projections["spare_shortfall"]["data"][0]
        carry_row = projections["carry_list"]["data"][0]
        self.assertEqual(shortfall_row["demand_count"], 1)
        self.assertEqual(shortfall_row["filled_count"], 1)
        self.assertEqual(shortfall_row["shortage_count"], 1)
        self.assertEqual(carry_row["shortage_count"], 1)
        self.assertEqual(carry_row["used_quantity"], 1)
        self.assertEqual(carry_row["carried_quantity"], 7)
        self.assertAlmostEqual(carry_row["utilization"], 1 / 7)

    def test_aircraft_support_carry_capacity_uses_only_successful_samples_after_partial_failure(self) -> None:
        projections = self.adapter._aircraft_support_v1_analysis_projections(
            {"planned_sorties": 3, "requested_sample_count": 3, "failed_sample_count": 1},
            "base-artifact",
            samples=[{"sample_index": 0, "events": []}, {"sample_index": 1, "events": []}],
            simulation_inputs={
                "aircraft": {"assets": [{"tail_number": "J15-001", "model": "J-15"}]},
                "equipment_tree": {
                    "components": [
                        {"aircraft_model": "J-15", "product_id": "spare-a", "product_type": "LRU"}
                    ]
                },
                "support_network": {"nodes": [{"id": "base", "inventory": {"spare-a": 4}}]},
                "product_catalog": {"products": [{"id": "spare-a", "name": "备件 A"}]},
            },
        )

        carry_row = projections["carry_list"]["data"][0]
        self.assertEqual(carry_row["recommended_quantity"], 4)
        self.assertEqual(carry_row["carried_quantity"], 8)
        self.assertNotEqual(carry_row["carried_quantity"], 12)

    def test_aircraft_support_spare_projection_keeps_the_aircraft_model_for_each_spare(self) -> None:
        projections = self.adapter._aircraft_support_v1_analysis_projections(
            {"planned_sorties": 4, "spare_fill_rate": 0.5, "spare_utilization": 0.5},
            "base-artifact",
            samples=[
                {
                    "sample_index": 0,
                    "events": [
                        {
                            "event": "spare_shortage",
                            "details": {
                                "job_id": "job-j15",
                                "aircraft_model": "J-15",
                                "resource_id": "carrier-deck",
                                "spare_type": "发动机备件",
                                "required_quantity": 1,
                            },
                        },
                        {
                            "event": "spare_shortage",
                            "details": {
                                "job_id": "job-j35",
                                "aircraft_model": "J-35",
                                "resource_id": "carrier-deck",
                                "spare_type": "雷达备件",
                                "required_quantity": 1,
                            },
                        },
                    ],
                }
            ],
            simulation_inputs={
                "mission_profile": {
                    "airports": [{"id": "carrier-deck", "name": "航母飞行甲板", "supportNodeId": "carrier-deck"}],
                },
                "aircraft": {
                    "assets": [
                        {"tail_number": "J15-001", "model": "J-15", "airport": "航母飞行甲板"},
                        {"tail_number": "J35-001", "model": "J-35", "airport": "航母飞行甲板"},
                    ],
                },
                "support_network": {
                    "nodes": [
                        {
                            "id": "carrier-deck",
                            "name": "基地",
                            "inventory": {"发动机备件": 0, "雷达备件": 0},
                        }
                    ]
                },
            },
        )

        rows = projections["spare_shortfall"]["data"]
        self.assertEqual(
            [(row["aircraft_model"], row["spare_type"]) for row in rows],
            [("J-15", "发动机备件"), ("J-35", "雷达备件")],
        )

    def test_aircraft_support_spare_projection_uses_only_modeled_aircraft_spare_pairs(self) -> None:
        projections = self.adapter._aircraft_support_v1_analysis_projections(
            {"planned_sorties": 4, "spare_fill_rate": 1.0, "spare_utilization": 0.0},
            "base-artifact",
            simulation_inputs={
                "aircraft": {
                    "assets": [
                        {"tail_number": "J15-001", "model": "J-15"},
                        {"tail_number": "J35-001", "model": "J-35"},
                    ],
                },
                "equipment_tree": {
                    "components": [
                        {"id": "j15-engine", "aircraft_model": "J-15", "product_id": "product-j15-engine", "product_name": "发动机备件"},
                        {"id": "j35-radar", "aircraft_model": "J-35", "product_id": "product-j35-radar", "product_name": "雷达备件"},
                        {"id": "j15-zero", "aircraft_model": "J-15", "product_id": "product-j15-zero", "product_name": "零库存备件"},
                    ],
                },
                "support_network": {
                    "nodes": [{
                        "id": "carrier-deck",
                        "inventory": {"product-j15-engine": 2, "product-j35-radar": 3, "product-unknown": 9},
                        "product_names": {"product-j15-engine": "发动机备件", "product-j35-radar": "雷达备件", "product-unknown": "未知备件"},
                    }],
                },
            },
        )

        rows = projections["spare_shortfall"]["data"]
        self.assertEqual(
            [(row["aircraft_model"], row["spare_type"]) for row in rows],
            [("J-15", "发动机备件"), ("J-35", "雷达备件"), ("J-15", "零库存备件")],
        )
        self.assertEqual(
            [row["product_id"] for row in rows],
            ["product-j15-engine", "product-j35-radar", "product-j15-zero"],
        )
        self.assertNotIn("全部机型", {row["aircraft_model"] for row in rows})
        carry_by_product = {row["product_id"]: row for row in projections["carry_list"]["data"]}
        self.assertEqual(carry_by_product["product-j15-engine"]["utilization"], 0.0)
        self.assertEqual(carry_by_product["product-j35-radar"]["utilization"], 0.0)
        self.assertIsNone(carry_by_product["product-j15-zero"]["utilization"])

    def test_run_service_submits_aircraft_support_v1_formal_monte_carlo_run(self) -> None:
        created = self._create_imported_sample_project()
        project = copy.deepcopy(created["project"])
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "m9.7.4 aircraft support monte carlo",
                "steps": 2,
                "projectJson": copy.deepcopy(project),
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 2,
                        "sweep": {
                            "failureRates": [0.01],
                            "spareMultipliers": [1.0],
                            "supportCapacities": [2],
                        },
                    }
                },
            },
        )
        service = RunService(self.repository, self.adapter, self.api.output_dir)

        submitted = service.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "monte_carlo",
                "formal_run": True,
            }
        )
        status = service.get_run_status(submitted["run_id"])
        result = self.api.get_run_result(submitted["run_id"])
        manifest = self.api.get_run_artifacts(submitted["run_id"])
        chain = self.api.get_run_chain(submitted["run_id"])
        base_artifact = self._artifact_by_kind(manifest, "monte_carlo_base")
        state_artifact = self._artifact_by_kind(manifest, "visualization_state_series")
        base_payload = json.loads((Path(self.api.output_dir) / base_artifact["path"]).read_text(encoding="utf-8"))
        state_payload = json.loads((Path(self.api.output_dir) / state_artifact["path"]).read_text(encoding="utf-8"))
        projection_artifacts = [
            artifact for artifact in manifest["artifacts"] if artifact["kind"].startswith("analysis_projection_")
        ]

        self.assertEqual(submitted["status"], "succeeded")
        self.assertEqual(submitted["phase"], "completed")
        self.assertEqual(submitted["model_family"], "aircraft_support_v1")
        self.assertEqual(submitted["run_type"], "monte_carlo")
        self.assertEqual(status["run_type"], "monte_carlo")
        self.assertEqual(status["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(result["model_family"], "aircraft_support_v1")
        self.assertEqual(result["run_id"], submitted["run_id"])
        self.assertEqual(base_payload["model_family"], "aircraft_support_v1")
        self.assertEqual(base_payload["sample_count"], 2)
        self.assertNotIn("m9_7_4_pending_fields", base_payload["sampling_contract"])
        self.assertIn("m9_7_4_closed_field_policy", base_payload["sampling_contract"])
        self.assertEqual(base_payload["logs_summary"]["completed_samples"], 2)
        self.assertEqual(base_payload["logs_summary"]["failed_samples"], 0)
        self.assertEqual(len(projection_artifacts), 4)
        self.assertTrue(all(artifact["source_artifact_id"] == base_artifact["artifact_id"] for artifact in projection_artifacts))
        self.assertEqual(state_payload["model_family"], "aircraft_support_v1")
        self.assertTrue(all("sample_index" in frame and "sample_step" in frame for frame in state_payload["frames"]))
        self.assertEqual(chain["scenario_id"], submitted["scenario_id"])
        self.assertEqual(chain["result_summary_id"], submitted["result_summary_id"])
        self.assertEqual(chain["artifact_manifest_id"], submitted["artifact_manifest_id"])
        self.assertEqual(len(self.adapter.compile_calls), 1)
        self.assertEqual(self.adapter.compile_calls[0][1], "aircraft_support_v1")
        self.assertEqual(len(self.adapter.monte_carlo_run_calls), 1)
        self.assertEqual(self.adapter.monte_carlo_run_calls[0]["scenario"]["simulation_model"]["family"], "aircraft_support_v1")

    def test_run_service_blocks_invalid_legacy_support_organization_tree(self) -> None:
        created = self._create_imported_sample_project()
        project = copy.deepcopy(created["project"])
        project["supportOrganization"] = {"tree": [{"id": "carrier-wing-support"}]}

        with self.assertRaises(BackendApiError) as raised:
            self.api.save_project(project)

        self.assertEqual(raised.exception.code, "invalid_project")

    def test_formal_run_persists_modeling_import_validation_scope_in_scenario_provenance(self) -> None:
        import_package = self._reduced_scope_import_package_without_support_domain()
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])
        created = self.api.create_project_from_modeling_import_as_system(import_package["importId"])
        project = created["project"]
        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "reduced scope validation run", "projectJson": copy.deepcopy(project)},
        )

        submitted = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
                "formal_run": True,
            }
        )

        manifest = self.api.get_run_artifacts(submitted["run_id"])
        compiled_artifact = self._artifact_by_kind(manifest, "compiled_scenario")
        compiled_payload = json.loads((Path(self.api.output_dir) / compiled_artifact["path"]).read_text(encoding="utf-8"))
        provenance = compiled_payload["compiled_from"]["mapping_provenance"]

        self.assertNotIn("validation_level", provenance)
        self.assertIn("supportResources", provenance["disabled_domains"])
        self.assertIn("supportActivities", provenance["disabled_domains"])
        self.assertEqual(provenance["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(provenance["experiment_plan_id"], plan["experiment_plan_id"])

    def test_public_import_templates_run_through_formal_backend_api_e2e(self) -> None:
        template_cases = [
            ("canonical_platform_case.json", False),
        ]
        required_monte_carlo_artifacts = M7_MONTE_CARLO_ARTIFACT_KINDS

        for template_name, expected_not_applicable in template_cases:
            with self.subTest(template=template_name):
                import_package = self._public_import_template(template_name)
                validation = validate_modeling_import_package(import_package)
                self.assertTrue(validation["ok"])
                self.assertEqual(validation["issues"], [])
                self.assertNotIn("validationLevel", validation)
                self.assertEqual(validation["usedTables"], import_package["usedTables"])

                self.api.save_modeling_import_as_system(import_package)
                self.api.publish_modeling_import_as_system(import_package["importId"])
                created = self.api.create_project_from_modeling_import_as_system(import_package["importId"])
                project = created["project"]
                saved_project = created["savedProject"]
                snapshot = created["modelingSnapshot"]

                single_plan = self.api.create_experiment_plan(
                    saved_project["project_id"],
                    {
                        "name": f"{template_name} single run",
                        "steps": 4,
                        "projectJson": copy.deepcopy(project),
                        "modeling_snapshot_id": snapshot["snapshot_id"],
                    },
                )
                single_run = self.api.submit_run(
                    {
                        "project_id": saved_project["project_id"],
                        "experiment_plan_id": single_plan["experiment_plan_id"],
                        "model_family": "aircraft_support_v1",
                        "run_type": "single",
                        "formal_run": True,
                    }
                )
                self.assertEqual(single_run["status"], "succeeded")
                single_chain = self.api.get_run_chain(single_run["run_id"])
                self.assertEqual(single_chain["project_id"], saved_project["project_id"])
                self.assertEqual(single_chain["modeling_snapshot_id"], snapshot["snapshot_id"])
                self.assertEqual(single_chain["experiment_plan_id"], single_plan["experiment_plan_id"])
                self.assertEqual(single_chain["run_id"], single_run["run_id"])
                self.assertEqual(single_chain["result_summary_id"], single_run["result_summary_id"])
                self.assertEqual(single_chain["artifact_manifest_id"], single_run["artifact_manifest_id"])
                single_stream = self.api.subscribe_run_state_stream(single_run["run_id"])
                self.assertIn("state_frame", [event["event_type"] for event in single_stream["events"]])

                monte_carlo_plan = self.api.create_experiment_plan(
                    saved_project["project_id"],
                    {
                        "name": f"{template_name} monte carlo run",
                        "steps": 4,
                        "projectJson": copy.deepcopy(project),
                        "modeling_snapshot_id": snapshot["snapshot_id"],
                        "analysisRequests": {
                            **copy.deepcopy(import_package["objects"].get("analysisRequests") or {}),
                            "largeSample": {
                                "enabled": True,
                                "samples": 1,
                                "sweep": {
                                    "failureRates": [0.05],
                                    "spareMultipliers": [1.0],
                                    "supportCapacities": [2],
                                },
                            },
                        },
                    },
                )
                monte_carlo_run = self.api.submit_run(
                    {
                        "project_id": saved_project["project_id"],
                        "experiment_plan_id": monte_carlo_plan["experiment_plan_id"],
                        "model_family": "aircraft_support_v1",
                        "run_type": "monte_carlo",
                        "formal_run": True,
                    }
                )
                self.assertEqual(monte_carlo_run["status"], "succeeded")
                monte_carlo_chain = self.api.get_run_chain(monte_carlo_run["run_id"])
                self.assertEqual(monte_carlo_chain["project_id"], saved_project["project_id"])
                self.assertEqual(monte_carlo_chain["modeling_snapshot_id"], snapshot["snapshot_id"])
                self.assertEqual(monte_carlo_chain["experiment_plan_id"], monte_carlo_plan["experiment_plan_id"])
                self.assertEqual(monte_carlo_chain["run_id"], monte_carlo_run["run_id"])
                self.assertEqual(monte_carlo_chain["result_summary_id"], monte_carlo_run["result_summary_id"])
                self.assertEqual(monte_carlo_chain["artifact_manifest_id"], monte_carlo_run["artifact_manifest_id"])

                manifest = self.api.get_run_artifacts(monte_carlo_run["run_id"])
                artifact_kinds = {artifact["kind"] for artifact in manifest["artifacts"]}
                self.assertTrue(required_monte_carlo_artifacts <= artifact_kinds)
                stream = self.api.subscribe_run_state_stream(monte_carlo_run["run_id"])
                self.assertIn("state_frame", [event["event_type"] for event in stream["events"]])
                self.assertEqual(stream["events"][-1]["payload"]["kind"], "visualization_state_series")

                spare_shortfall = self._artifact_payload(manifest, "analysis_projection_spare_shortfall")
                downtime_factors = self._artifact_payload(manifest, "analysis_projection_downtime_factors")
                expected_applicability = "not_applicable" if expected_not_applicable else "applicable"
                self.assertEqual(spare_shortfall["applicability"]["status"], expected_applicability)
                self.assertEqual(downtime_factors["applicability"]["status"], expected_applicability)
                if expected_not_applicable:
                    self.assertIn("supportResources", spare_shortfall["applicability"]["disabled_domains"])
                    self.assertIn("supportActivities", downtime_factors["applicability"]["disabled_domains"])

    def test_run_service_failed_compile_run_has_downloadable_log_artifact(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "blocked unsupported family"})

        submitted = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "unsupported_family",
                "run_type": "single",
            }
        )
        manifest = self.api.get_run_artifacts(submitted["run_id"])
        log_artifacts = [artifact for artifact in manifest["artifacts"] if artifact["kind"] == "log"]

        self.assertEqual(submitted["status"], "failed")
        self.assertEqual(len(log_artifacts), 1)
        self.assertRegex(log_artifacts[0]["sha256"], r"^[0-9a-f]{64}$")
        self.assertGreater(log_artifacts[0]["size_bytes"], 0)
        log_path = Path(self.api.output_dir) / log_artifacts[0]["path"]
        log_data = log_path.read_bytes()
        log_payload = json.loads(log_data.decode("utf-8"))
        self.assertEqual(hashlib.sha256(log_data).hexdigest(), log_artifacts[0]["sha256"])
        self.assertEqual(len(log_data), log_artifacts[0]["size_bytes"])
        self.assertEqual(log_payload["run_id"], submitted["run_id"])
        self.assertEqual(log_payload["status"], "failed")
        self.assertEqual(log_payload["events"][0]["event"], "compile_gate_failed")
        self.assertEqual(log_payload["events"][0]["error"]["code"], "unsupported_model_family")
        self.assertTrue(log_payload["events"][0]["issues"])
        self.assertEqual(log_payload["events"][0]["provenance"]["model_family"], "unsupported_family")

    def test_backend_api_submit_run_uses_m6_status_envelope(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "submit run", "steps": 1})

        submitted = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        status = self.api.get_run_status(submitted["run_id"])

        self.assertEqual(submitted["status"], "succeeded")
        self.assertEqual(submitted["phase"], "completed")
        self.assertEqual(status["run_id"], submitted["run_id"])
        self.assertEqual(status["experiment_plan_id"], plan["experiment_plan_id"])

    def test_formal_run_records_modeling_import_reference_and_blocks_republish(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])
        created = self.api.create_project_from_modeling_import_as_system(import_package["importId"])
        plan = self.api.create_experiment_plan(
            created["savedProject"]["project_id"],
            {"name": "lineage reference", "steps": 1, "projectJson": created["project"]},
        )

        submitted = self.api.submit_run(
            {
                "project_id": created["savedProject"]["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
                "formal_run": True,
            }
        )

        stored = self.api.get_modeling_import(import_package["importId"])
        self.assertEqual(stored["publishedPackage"]["lifecycle"]["referencedRunIds"], [submitted["run_id"]])
        self.repository.record_modeling_import_run_reference(import_package["importId"], submitted["run_id"])
        stored_again = self.api.get_modeling_import(import_package["importId"])
        self.assertEqual(stored_again["publishedPackage"]["lifecycle"]["referencedRunIds"], [submitted["run_id"]])

        with self.assertRaises(BackendApiError) as ctx:
            self.api.publish_modeling_import_as_system(import_package["importId"])
        self.assertEqual(ctx.exception.code, "published_import_referenced")
        self.assertEqual(ctx.exception.details["import_id"], import_package["importId"])

    def test_failed_formal_run_does_not_record_modeling_import_reference(self) -> None:
        failing_api = BackendApi(self.repository, FailingRunAdapter(), output_dir=Path(self.tempdir.name) / "failing")
        import_package = self._fixture("modeling_import_project.json")
        failing_api.save_modeling_import_as_system(import_package)
        failing_api.publish_modeling_import_as_system(import_package["importId"])
        created = failing_api.create_project_from_modeling_import_as_system(import_package["importId"])
        plan = failing_api.create_experiment_plan(
            created["savedProject"]["project_id"],
            {"name": "failed formal lineage", "steps": 1, "projectJson": created["project"]},
        )

        submitted = failing_api.submit_run(
            {
                "project_id": created["savedProject"]["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
                "formal_run": True,
            }
        )
        stored = failing_api.get_modeling_import(import_package["importId"])

        self.assertEqual(submitted["status"], "failed")
        self.assertEqual(stored["publishedPackage"]["lifecycle"].get("referencedRunIds"), [])
        republished = failing_api.publish_modeling_import_as_system(import_package["importId"])
        self.assertEqual(republished["lifecycle"]["state"], "published")

    def test_backend_api_submit_run_rejects_project_plan_mismatch(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "mismatch", "steps": 1})
        other_project = copy.deepcopy(project)
        other_project["project_id"] = "project-other"
        other_project["scenarioId"] = "other-aircraft-support-contract-demo"
        other_saved = self.api.save_project(other_project)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.submit_run(
                {
                    "project_id": other_saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "single",
                }
            )

        self.assertEqual(ctx.exception.code, "project_plan_mismatch")

    def test_backend_api_submit_run_rejects_formal_aviation_support_single_run(self) -> None:
        created = self._create_imported_sample_project()
        project = created["project"]
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "backend aviation formal", "steps": 2, "projectJson": copy.deepcopy(project)},
        )

        with self.assertRaises(BackendApiError) as ctx:
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aviation_support",
                    "run_type": "single",
                    "formal_run": True,
                }
            )

        self.assertEqual(ctx.exception.code, "retired_model_family")
        self.assertEqual(ctx.exception.details["model_family"], "aviation_support")
        self.assertEqual(ctx.exception.details["replacement_model_family"], "aircraft_support_v1")
        self.assertEqual(self.adapter.compile_calls, [])
        self.assertEqual(self.adapter.run_calls, [])

    def test_backend_api_formal_run_still_rejects_non_imported_sample_project(self) -> None:
        project = self._fixture("aircraft_support_v1_project.json")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "aviation formal gate", "steps": 1})

        with self.assertRaises(BackendApiError) as ctx:
            self.api.submit_run(
                {
                    "project_id": saved["project_id"],
                    "experiment_plan_id": plan["experiment_plan_id"],
                    "model_family": "aircraft_support_v1",
                    "run_type": "single",
                    "formal_run": True,
                }
            )

        self.assertEqual(ctx.exception.code, "formal_run_requires_imported_sample")
        self.assertEqual(ctx.exception.details["project_id"], saved["project_id"])
        self.assertIsNone(ctx.exception.details["source_import_id"])
        self.assertEqual(self.adapter.compile_calls, [])
        self.assertEqual(self.adapter.run_calls, [])

    def test_backend_api_start_simulation_run_delegates_to_m6_run_service(self) -> None:
        created = self._create_imported_sample_project()
        project = created["project"]
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "compat", "steps": 1, "projectJson": copy.deepcopy(project)},
        )

        run = self.api.start_simulation_run(saved["project_id"], plan["experiment_plan_id"])
        status = self.api.get_run_status(run["run_id"])

        self.assertEqual(run["status"], "succeeded")
        self.assertEqual(run["model_family"], "aircraft_support_v1")
        self.assertEqual(status["phase"], "completed")
        self.assertEqual(status["run_id"], run["run_id"])
        self.assertEqual(status["experiment_plan_id"], plan["experiment_plan_id"])

    def test_run_service_persists_failed_status_when_executor_fails_after_scenario_compile(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "failed executor", "steps": 1})
        failing_adapter = FailingRunAdapter()
        service = RunService(self.repository, failing_adapter, self.api.output_dir)

        submitted = service.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        stored = self.api.get_run(submitted["run_id"])
        artifacts = self.api.get_run_artifacts(submitted["run_id"])

        self.assertEqual(submitted["status"], "failed")
        self.assertEqual(submitted["phase"], "failed")
        self.assertEqual(submitted["progress"], 0)
        self.assertEqual(submitted["error"]["code"], "executor_failed")
        self.assertEqual(stored["status"], "failed")
        self.assertEqual(artifacts["run_id"], submitted["run_id"])
        self.assertEqual([artifact["kind"] for artifact in artifacts["artifacts"]], ["log"])

    def test_run_chain_preserves_snapshot_and_plan_after_project_resave(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")

        first_saved = self.api.save_project(project)
        first_snapshot = self.api.create_modeling_snapshot(first_saved["project_id"])
        first_plan = self.api.create_experiment_plan(first_saved["project_id"], {"name": "same config", "steps": 1})
        first_run = self.api.submit_run(
            {
                "project_id": first_saved["project_id"],
                "experiment_plan_id": first_plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        changed_project = copy.deepcopy(project)
        changed_project.setdefault("projectInfo", {})["name"] = "changed after first run"
        second_saved = self.api.save_project(changed_project)
        second_snapshot = self.api.create_modeling_snapshot(second_saved["project_id"])
        second_plan = self.api.create_experiment_plan(second_saved["project_id"], {"name": "same config", "steps": 1})
        second_run = self.api.submit_run(
            {
                "project_id": second_saved["project_id"],
                "experiment_plan_id": second_plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        first_chain = self.api.get_run_chain(first_run["run_id"])
        second_chain = self.api.get_run_chain(second_run["run_id"])

        self.assertNotEqual(first_snapshot["snapshot_id"], second_snapshot["snapshot_id"])
        self.assertNotEqual(first_plan["experiment_plan_id"], second_plan["experiment_plan_id"])
        self.assertEqual(first_chain["modeling_snapshot_id"], first_snapshot["snapshot_id"])
        self.assertEqual(first_chain["experiment_plan_id"], first_plan["experiment_plan_id"])
        self.assertEqual(second_chain["modeling_snapshot_id"], second_snapshot["snapshot_id"])
        self.assertEqual(second_chain["experiment_plan_id"], second_plan["experiment_plan_id"])
        self.assertEqual(first_chain["run_id"], first_run["run_id"])
        self.assertEqual(second_chain["run_id"], second_run["run_id"])

    def test_experiment_plan_can_bind_explicit_current_snapshot_after_project_change(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        old_snapshot = self.api.create_modeling_snapshot(saved["project_id"])

        changed_project = copy.deepcopy(project)
        changed_project.setdefault("projectInfo", {})["name"] = "current run input after old snapshot"
        self.api.save_project(changed_project)
        current_snapshot = self.api.create_modeling_snapshot(saved["project_id"])

        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "current run input",
                "steps": 1,
                "projectJson": changed_project,
                "modeling_snapshot_id": current_snapshot["snapshot_id"],
            },
        )
        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        chain = self.api.get_run_chain(run["run_id"])

        self.assertNotEqual(old_snapshot["snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(plan["modeling_snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(chain["modeling_snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(plan["config"]["projectJson"]["projectInfo"]["name"], "current run input after old snapshot")
        self.assertNotIn("experiment", plan["config"]["projectJson"])
        self.assertNotIn("modeling_snapshot_id", plan["config"])

    def test_create_experiment_plan_preserves_seed_policy_and_scenario_composition(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        slim_project = strip_project_sweep(project)

        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "composed branch",
                "steps": 4,
                "samples": 9,
                "seed": 909,
                "seedPolicy": {"mode": "fixed", "baseSeed": 909},
                "stopPolicy": {
                    "schemaVersion": "stop-policy-v0",
                    "mode": "and",
                    "conditions": [{"type": "duration"}, {"type": "failure"}],
                },
                "scenarioComposition": {
                    "schemaVersion": "scenario-composition-v0",
                    "overrides": [
                        {"path": "supportResources.2.quantity", "valueType": "number", "value": 12}
                    ],
                },
                "analysisRequests": {
                    "largeSample": {
                        "enabled": True,
                        "samples": 9,
                        "sweep": {
                            "failureRates": [0.06],
                            "spareMultipliers": [1],
                            "supportCapacities": [2],
                        },
                    }
                },
                "projectJson": {
                    **copy.deepcopy(slim_project),
                    "seedPolicy": {"mode": "fixed", "baseSeed": 909},
                    "stopPolicy": {
                        "schemaVersion": "stop-policy-v0",
                        "mode": "and",
                        "conditions": [{"type": "duration"}, {"type": "failure"}],
                    },
                    "scenarioComposition": {
                        "schemaVersion": "scenario-composition-v0",
                        "overrides": [
                            {"path": "supportResources.2.quantity", "valueType": "number", "value": 12}
                        ],
                    },
                    "experiment": {"name": "composed branch", "steps": 4, "samples": 9, "seed": 909},
                    "analysisRequests": {"largeSample": {"enabled": True, "samples": 9}},
                    "monteCarlo": {"failureRates": [0.06]},
                },
            },
        )

        self.assertEqual(plan["config"]["seedPolicy"], {"mode": "fixed", "baseSeed": 909})
        self.assertEqual(
            plan["config"]["stopPolicy"],
            {
                "schemaVersion": "stop-policy-v0",
                "mode": "and",
                "conditions": [{"type": "duration"}, {"type": "failure"}],
            },
        )
        self.assertEqual(
            plan["config"]["scenarioComposition"]["overrides"][0]["path"],
            "supportResources.2.quantity",
        )
        self.assertEqual(plan["config"]["analysisRequests"]["largeSample"]["samples"], 9)
        self.assertEqual(plan["config"]["projectJson"]["supportResources"][2]["quantity"], 12)
        self.assertNotIn("experiment", plan["config"]["projectJson"])
        self.assertNotIn("analysisRequests", plan["config"]["projectJson"])
        self.assertNotIn("monteCarlo", plan["config"]["projectJson"])
        self.assertNotIn("seedPolicy", plan["config"]["projectJson"])
        self.assertNotIn("stopPolicy", plan["config"]["projectJson"])
        self.assertNotIn("scenarioComposition", plan["config"]["projectJson"])

    def test_experiment_plan_config_branch_does_not_mutate_source_project(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)

        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "branch before mutation",
                "steps": 4,
                "assumptions": {
                    "support_capacity": 2,
                },
            },
        )
        plan["config"]["name"] = "branch after mutation"
        plan["config"]["assumptions"]["support_capacity"] = 99
        self.repository.upsert_experiment_plan(plan)

        stored_project = self.api.get_project(saved["project_id"])

        self.assertEqual(stored_project, strip_project_sweep(project))
        self.assertNotIn("experiment", stored_project)
        self.assertNotIn("assumptions", stored_project)

    def test_single_run_compiles_from_experiment_plan_project_branch(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        branch_project = copy.deepcopy(project)
        branch_project["experiment"]["seed"] = 99
        branch_project["components"][0]["failureDistribution"] = {
            "distributionType": "exponential",
            "parameters": "lambda=0.21",
        }
        branch_project["components"][0].setdefault("rms", {})["prediction"] = {"mtbfHours": 100}
        branch_project["resultSummary"] = {"mission_success_rate": 1}
        branch_project = strip_project_sweep(branch_project)
        equipment_resource = next(resource for resource in branch_project["supportResources"] if resource["type"] == "equipment")
        equipment_resource["quantity"] = 8

        saved = self.api.save_project(project)
        snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "single input branch",
                "steps": 5,
                "seed": 99,
                "projectJson": branch_project,
            },
        )

        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        compiled_project, model_family = self.adapter.compile_calls[-1]
        compiled_scenario, steps, run_id = self.adapter.run_calls[-1]
        provenance = compiled_scenario["compiled_from"]["mapping_provenance"]

        self.assertEqual(model_family, "aircraft_support_v1")
        self.assertNotIn("experiment", compiled_project)
        self.assertNotIn("resultSummary", compiled_project)
        self.assertNotIn("rms", compiled_project["components"][0])
        self.assertNotIn("resultSummary", self.adapter.compile_runtime_configs[-1]["projectJson"])
        self.assertNotIn("rms", self.adapter.compile_runtime_configs[-1]["projectJson"]["components"][0])
        clean_export = provenance["clean_project_export"]
        self.assertEqual(clean_export["target"], "aircraft_support_v1")
        self.assertIn("resultSummary", clean_export["stripped_fields"])
        self.assertEqual(self.adapter.compile_runtime_configs[-1]["seed"], 99)
        self.assertEqual(compiled_project["components"][0]["failureDistribution"]["parameters"], "lambda=0.21")
        self.assertEqual(next(resource for resource in compiled_project["supportResources"] if resource["type"] == "equipment")["quantity"], 8)
        self.assertEqual(compiled_scenario["simulation_inputs"]["seed"], 99)
        self.assertEqual(compiled_scenario["simulation_inputs"]["equipment_tree"]["components"][0]["failure_rate"], 0.21)
        self.assertEqual(compiled_scenario["simulation_inputs"]["support_network"]["nodes"][0]["equipment_capacity"], 8)
        self.assertEqual(steps, 5)
        self.assertEqual(run_id, run["run_id"])
        self.assertEqual(provenance["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(provenance["modeling_snapshot_id"], snapshot["snapshot_id"])
        self.assertEqual(self.api.get_project(saved["project_id"]), strip_project_sweep(project))

    def test_create_experiment_plan_binds_explicit_or_latest_modeling_snapshot(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        first_snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        project.setdefault("projectInfo", {})["name"] = "latest snapshot source"
        self.api.save_project(project)
        latest_snapshot = self.api.create_modeling_snapshot(saved["project_id"])

        explicit_plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "explicit snapshot",
                "steps": 1,
                "modeling_snapshot_id": first_snapshot["snapshot_id"],
            },
        )
        latest_plan = self.api.create_experiment_plan(saved["project_id"], {"name": "latest snapshot", "steps": 1})
        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": explicit_plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        self.assertEqual(explicit_plan["modeling_snapshot_id"], first_snapshot["snapshot_id"])
        self.assertEqual(latest_plan["modeling_snapshot_id"], latest_snapshot["snapshot_id"])
        self.assertEqual(self.api.get_run_chain(run["run_id"])["modeling_snapshot_id"], first_snapshot["snapshot_id"])

    def test_update_experiment_plan_preserves_original_id(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "edit target", "steps": 1})
        project.setdefault("projectInfo", {})["name"] = "edited source"
        self.api.save_project(project)
        latest_snapshot = self.api.create_modeling_snapshot(saved["project_id"])

        updated = self.api.update_experiment_plan(
            saved["project_id"],
            plan["experiment_plan_id"],
            {"name": "edit target", "steps": 9},
        )
        plans = self.api.list_experiment_plans(saved["project_id"])["experiment_plans"]

        self.assertEqual(updated["experiment_plan_id"], plan["experiment_plan_id"])
        self.assertEqual(updated["config"]["steps"], 9)
        self.assertEqual(updated["modeling_snapshot_id"], latest_snapshot["snapshot_id"])
        self.assertEqual([item["experiment_plan_id"] for item in plans], [plan["experiment_plan_id"]])

    def test_submit_run_after_project_edit_uses_new_explicit_modeling_snapshot(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")
        saved = self.api.save_project(project)
        old_snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        project["components"][0]["failureDistribution"] = {
            "distributionType": "exponential",
            "parameters": "lambda=0.33",
        }
        saved = self.api.save_project(project)
        current_snapshot = self.api.create_modeling_snapshot(saved["project_id"])
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {
                "name": "current explicit snapshot",
                "steps": 3,
                "seed": 606,
                "projectJson": project,
                "modeling_snapshot_id": current_snapshot["snapshot_id"],
            },
        )

        run = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        compiled_project, _model_family = self.adapter.compile_calls[-1]
        compiled_scenario, _steps, _run_id = self.adapter.run_calls[-1]
        provenance = compiled_scenario["compiled_from"]["mapping_provenance"]

        self.assertNotEqual(old_snapshot["snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(plan["modeling_snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(self.api.get_run_chain(run["run_id"])["modeling_snapshot_id"], current_snapshot["snapshot_id"])
        self.assertEqual(provenance["modeling_snapshot_id"], current_snapshot["snapshot_id"])
        self.assertNotIn("experiment", compiled_project)
        self.assertEqual(compiled_scenario["simulation_inputs"]["seed"], 606)
        self.assertEqual(compiled_project["components"][0]["failureDistribution"]["parameters"], "lambda=0.33")
        self.assertEqual(compiled_scenario["simulation_inputs"]["equipment_tree"]["components"][0]["failure_rate"], 0.33)

    def test_backend_api_delegates_submit_run_without_owning_lifecycle_lock(self) -> None:
        shared_lock = threading.Lock()
        api = BackendApi(self.repository, self.adapter, output_dir=Path(self.tempdir.name), run_lifecycle_lock=shared_lock)
        observed: dict[str, Any] = {}
        self.assertIs(api.run_service._run_lock, shared_lock)

        class FakeRunService:
            def submit_run(self, request: dict[str, Any]) -> dict[str, Any]:
                observed["outer_lock_held"] = shared_lock.locked()
                observed["request"] = copy.deepcopy(request)
                return {"run_id": "run-fake", "status": "succeeded"}

        api.run_service = FakeRunService()  # type: ignore[assignment]

        result = api.submit_run({"project_id": "project-fake", "experiment_plan_id": "plan-fake", "model_family": "aircraft_support_v1"})

        self.assertEqual(result["run_id"], "run-fake")
        self.assertEqual(observed["outer_lock_held"], False)
        self.assertFalse(hasattr(api, "_run_lock"))

    def test_run_service_uses_injected_lock_for_all_lifecycle_mutations(self) -> None:
        observations: list[tuple[str, bool]] = []

        class RecordingLock:
            def __init__(self) -> None:
                self._locked = False

            def __enter__(self) -> "RecordingLock":
                self._locked = True
                return self

            def __exit__(self, _exc_type: object, _exc: object, _traceback: object) -> None:
                self._locked = False

            def locked(self) -> bool:
                return self._locked

        lock = RecordingLock()

        class FakeRepository:
            def delete_experiment_plan_with_runs(
                self,
                project_id: str,
                experiment_plan_id: str,
                *,
                actor_user_id: str,
            ) -> dict[str, Any]:
                observations.append(("delete_plan", lock.locked()))
                return {"project_id": project_id, "experiment_plan_id": experiment_plan_id, "deleted": True}

            def archive_run_with_audit(self, run_id: str, *, actor_user_id: str) -> dict[str, Any]:
                observations.append(("archive", lock.locked()))
                return {"run_id": run_id, "lifecycle_status": "archived"}

            def soft_delete_run_with_audit(self, run_id: str, *, actor_user_id: str) -> dict[str, Any]:
                observations.append(("delete", lock.locked()))
                return {"run_id": run_id, "lifecycle_status": "deleted"}

            def control_run_with_audit(self, run_id: str, action: str, *, actor_user_id: str) -> dict[str, Any]:
                observations.append(("control", lock.locked()))
                return {"run_id": run_id, "status": "cancelled", "control": {"action": action}}

        service = RunService(FakeRepository(), self.adapter, self.api.output_dir, run_lifecycle_lock=lock)  # type: ignore[arg-type]

        def submit_unlocked(_request: dict[str, Any]) -> dict[str, Any]:
            observations.append(("submit", lock.locked()))
            return {"run_id": "run-lock", "status": "succeeded"}

        service._submit_run_unlocked = submit_unlocked  # type: ignore[method-assign]

        service.submit_run({"project_id": "project-lock", "experiment_plan_id": "plan-lock", "model_family": "aircraft_support_v1"})
        service.delete_experiment_plan("project-lock", "plan-lock", actor_user_id="user-admin")
        service.archive_run("run-lock", actor_user_id="user-admin")
        service.soft_delete_run("run-lock", actor_user_id="user-admin")
        service.control_run("run-lock", "cancel", actor_user_id="user-admin")

        self.assertEqual(
            observations,
            [
                ("submit", True),
                ("delete_plan", True),
                ("archive", True),
                ("delete", True),
                ("control", True),
            ],
        )

    def test_repeated_current_runs_create_distinct_run_chains(self) -> None:
        project = small_aircraft_support_project("project-aircraft-support-contract-001")

        saved = self.api.save_project(project)
        plan = self.api.create_experiment_plan(saved["project_id"], {"name": "repeatable current run", "steps": 1})
        first = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )
        second = self.api.submit_run(
            {
                "project_id": saved["project_id"],
                "experiment_plan_id": plan["experiment_plan_id"],
                "model_family": "aircraft_support_v1",
                "run_type": "single",
            }
        )

        self.assertNotEqual(first["run_id"], second["run_id"])
        self.assertNotEqual(first["scenario_id"], second["scenario_id"])
        self.assertNotEqual(first["result_summary_id"], second["result_summary_id"])
        self.assertNotEqual(first["artifact_manifest_id"], second["artifact_manifest_id"])
        self.assertEqual(self.api.get_run_chain(first["run_id"])["run_id"], first["run_id"])
        self.assertEqual(self.api.get_run_chain(second["run_id"])["run_id"], second["run_id"])

    def test_aviation_support_start_simulation_run_shortcut_is_retired(self) -> None:
        created = self._create_imported_sample_project()
        project = created["project"]
        saved = created["savedProject"]
        plan = self.api.create_experiment_plan(
            saved["project_id"],
            {"name": "aviation start compat", "steps": 1, "projectJson": copy.deepcopy(project)},
        )

        with self.assertRaises(BackendApiError) as ctx:
            self.api.start_simulation_run(
                saved["project_id"],
                plan["experiment_plan_id"],
                model_family="aviation_support",
            )

        self.assertEqual(ctx.exception.code, "retired_model_family")
        self.assertEqual(ctx.exception.details["model_family"], "aviation_support")
        self.assertEqual(ctx.exception.details["replacement_model_family"], "aircraft_support_v1")
        self.assertEqual(self.adapter.compile_calls, [])
        self.assertEqual(self.adapter.run_calls, [])

    def test_modeling_import_api_validates_saves_and_publishes_package_as_system(self) -> None:
        import_package = self._fixture("modeling_import_project.json")

        validation = self.api.validate_modeling_import(import_package)
        saved = self.api.save_modeling_import_as_system(import_package)
        published = self.api.publish_modeling_import_as_system(import_package["importId"])
        stored = self.api.get_modeling_import(import_package["importId"])

        self.assertTrue(validation["ok"])
        self.assertEqual(validation["status"], "valid")
        self.assertEqual(validation["issues"], [])
        self.assertEqual(saved["import_id"], import_package["importId"])
        self.assertEqual(saved["project_id"], import_package["projectId"])
        self.assertEqual(saved["validation_status"], "valid")
        self.assertEqual(published["lifecycle"]["state"], "published")
        self.assertEqual(stored["validation"]["status"], "valid")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(
            [(event["action"], event["outcome"], event["actor_user_id"]) for event in events],
            [
                ("modeling_import.save", "allowed", None),
                ("modeling_import.publish", "allowed", None),
            ],
        )
        self.assertTrue(all(event["details"].get("actor") == "system" for event in events))

    def test_project_data_template_api_lists_published_project_templates(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        draft_package = self._fixture("modeling_import_project.json")
        draft_package["importId"] = "import-draft-only"
        draft_package["projectId"] = "project-draft-only"
        draft_package["objects"]["projectInfo"]["name"] = "草稿模板"

        self.api.save_modeling_import_as_system(draft_package)
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        templates = self.api.list_project_data_templates(state="published")

        self.assertEqual([template["template_id"] for template in templates["templates"]], [import_package["importId"]])
        self.assertEqual(templates["templates"][0]["source_import_id"], import_package["importId"])
        self.assertEqual(templates["templates"][0]["template_type"], "project_data")
        self.assertEqual(templates["templates"][0]["project_id"], import_package["projectId"])
        self.assertEqual(templates["templates"][0]["name"], import_package["objects"]["projectInfo"]["name"])
        self.assertEqual(templates["templates"][0]["validation_status"], "valid")
        self.assertGreaterEqual(templates["templates"][0]["object_counts"]["missionProfiles"], 1)
        self.assertNotIn("schema_version", templates["templates"][0])
        self.assertNotIn("version", templates["templates"][0])
        self.assertNotIn("lifecycle_state", templates["templates"][0])

    def test_modeling_import_api_rejects_missing_actor_for_save(self) -> None:
        import_package = self._fixture("modeling_import_project.json")

        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_modeling_import(import_package)

        self.assertEqual(ctx.exception.code, "unauthorized")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(events[-1]["action"], "modeling_import.save")
        self.assertEqual(events[-1]["outcome"], "denied")
        self.assertIsNone(events[-1]["actor_user_id"])
        self.assertEqual(events[-1]["details"]["reason"], "missing_actor")

    def test_modeling_import_api_rejects_missing_actor_for_publish(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.publish_modeling_import(import_package["importId"])

        self.assertEqual(ctx.exception.code, "unauthorized")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(events[-1]["action"], "modeling_import.publish")
        self.assertEqual(events[-1]["outcome"], "denied")
        self.assertIsNone(events[-1]["actor_user_id"])
        self.assertEqual(events[-1]["details"]["reason"], "missing_actor")

    def test_m4_data_admin_can_save_publish_and_audit_modeling_import(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        session = self.api.login("data", "data")

        saved = self.api.save_modeling_import(import_package, actor_user_id=session["user"]["user_id"])
        published = self.api.publish_modeling_import(
            import_package["importId"],
            actor_user_id=session["user"]["user_id"],
        )

        self.assertEqual(session["user"]["role"], "数据管理员")
        self.assertEqual(saved["status"], "draft")
        self.assertEqual(published["lifecycle"]["state"], "published")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(
            [(event["action"], event["outcome"]) for event in events],
            [
                ("modeling_import.save", "allowed"),
                ("modeling_import.publish", "allowed"),
            ],
        )
        self.assertEqual({event["actor_user_id"] for event in events}, {session["user"]["user_id"]})
        login_events = self.repository.list_audit_events()
        login_event = next(event for event in login_events if event["action"] == "auth.login")
        self.assertNotEqual(login_event["resource_id"], session["session"]["token"])
        self.assertRegex(login_event["resource_id"], r"^session-[0-9a-f]{16}$")

    def test_modeling_import_to_project_preserves_support_activity_jobs_object_surface(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["objects"]["supportActivityJobs"] = [
            {"activityCode": "BA-001", "workName": "检查雷达", "durationMinutes": 30},
            {"activityCode": "BA-002", "workName": "挂载雷达", "durationMinutes": 45},
        ]

        project = modeling_import_to_project(import_package)

        jobs_by_code = {job["activityCode"]: job for job in project["supportActivityJobs"]}
        self.assertEqual(jobs_by_code["BA-001"], {"activityCode": "BA-001", "workName": "检查雷达", "durationMinutes": 30})
        self.assertEqual(jobs_by_code["BA-002"], {"activityCode": "BA-002", "workName": "挂载雷达", "durationMinutes": 45})

    def test_create_project_from_modeling_import_saves_project_and_snapshot(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        created = self.api.create_project_from_modeling_import_as_system(import_package["importId"])
        second_created = self.api.create_project_from_modeling_import_as_system(import_package["importId"])

        self.assertEqual(created["sourceImport"]["import_id"], import_package["importId"])
        self.assertEqual(created["project"]["project_id"], import_package["projectId"])
        self.assertEqual(created["project"]["missionProfile"]["sourceImportId"], import_package["importId"])
        self.assertNotIn("equipment", created["project"])
        self.assertGreaterEqual(len(created["project"]["components"]), 8)
        avionics = next(component for component in created["project"]["components"] if component.get("id") == "j15-avionics")
        self.assertTrue(
            avionics.get("aircraftModel") == "J-15"
            and avionics.get("parentId") == "aircraft-root"
            and avionics.get("productType") == "LRU"
            and avionics.get("productId")
        )
        avionics_product = next(product for product in created["project"]["products"] if product["id"] == avionics["productId"])
        self.assertEqual(avionics_product["name"], "航电系统")
        self.assertNotIn("spareType", avionics)
        self.assertGreaterEqual(len(created["project"]["missionProfile"]["compositeTasks"]), 2)
        self.assertGreaterEqual(len(created["project"]["missionProfile"]["periodicTasks"]), 1)
        self.assertNotIn("basicMission", created["project"])
        self.assertGreaterEqual(len(created["project"]["basicMissions"]), 2)
        self.assertTrue(all(basic.get("id") for basic in created["project"]["basicMissions"]))
        self.assertEqual(created["project"]["airports"], ["A"])
        self.assertNotIn("missionPhases", created["project"])
        self.assertTrue(all(len(basic.get("missionPhases", [])) >= 3 for basic in created["project"]["basicMissions"]))
        self.assertGreaterEqual(len(created["project"]["combatUnit"]["members"]), 4)
        self.assertGreaterEqual(len(created["project"]["supportNodes"]), 3)
        self.assertTrue(any(
            resource["type"] == "spare" and resource["name"] == "航电模块"
            for resource in created["project"]["supportResources"]
        ))
        activity_types = {activity["activityType"] for activity in created["project"]["supportActivities"]}
        self.assertTrue({"飞行前保障", "修复性维修", "预防性维修", "后勤保障"}.issubset(activity_types))
        self.assertTrue(all("jobs" not in activity for activity in created["project"]["supportActivities"]))
        self.assertGreaterEqual(len(created["project"]["supportActivities"][0]["activityCodes"]), 2)
        self.assertGreaterEqual(len(created["project"]["supportActivityJobs"]), 2)
        self.assertEqual(created["project"]["reliabilityBlockDiagram"], import_package["objects"]["reliabilityBlockDiagram"])
        self.assertNotIn("experiment", created["project"])
        self.assertNotIn("analysisRequests", created["project"])
        self.assertEqual(created["savedProject"]["project_id"], import_package["projectId"])
        self.assertEqual(created["modelingSnapshot"]["project"]["project_id"], import_package["projectId"])
        self.assertEqual(self.api.get_project(import_package["projectId"])["project_id"], import_package["projectId"])
        self.assertEqual(self.api.get_project(import_package["projectId"])["projectInfo"]["name"], "导入示例项目")
        self.assertNotEqual(second_created["savedProject"]["project_id"], created["savedProject"]["project_id"])
        self.assertRegex(second_created["savedProject"]["project_id"], rf"^{import_package['projectId']}-copy-[0-9]+$")
        self.assertEqual(second_created["project"]["project_id"], second_created["savedProject"]["project_id"])
        self.assertEqual(second_created["project"]["missionProfile"]["sourceImportId"], import_package["importId"])
        self.assertEqual(second_created["project"]["projectInfo"]["name"], "导入示例项目 副本 2")
        self.assertEqual(self.api.get_project(import_package["projectId"])["projectInfo"]["name"], "导入示例项目")
        self.assertEqual(second_created["modelingSnapshot"]["project"]["project_id"], second_created["savedProject"]["project_id"])
        self.assertEqual(self.api.get_project(second_created["savedProject"]["project_id"])["project_id"], second_created["savedProject"]["project_id"])
        project_ids = {entry["project_id"] for entry in self.api.list_projects()["projects"]}
        self.assertTrue({created["savedProject"]["project_id"], second_created["savedProject"]["project_id"]}.issubset(project_ids))
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        create_events = [event for event in events if event["action"] == "modeling_import.create_project"]
        self.assertEqual(len(create_events), 2)
        self.assertEqual(create_events[0]["outcome"], "allowed")
        self.assertEqual(create_events[0]["resource_id"], import_package["importId"])
        self.assertEqual(create_events[0]["details"]["project_id"], import_package["projectId"])
        self.assertEqual(create_events[0]["details"]["import_version"], 1)
        self.assertEqual(create_events[0]["details"]["actor"], "system")
        self.assertEqual(
            {event["details"]["project_id"] for event in create_events},
            {created["savedProject"]["project_id"], second_created["savedProject"]["project_id"]},
        )

    def test_modeling_import_to_project_preserves_full_authoring_surfaces(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        objects = import_package["objects"]

        project = modeling_import_to_project(import_package)

        self.assertEqual(project["projectInfo"], objects["projectInfo"])
        self.assertIsNot(project["projectInfo"], objects["projectInfo"])
        self.assertEqual(project["supportOrganization"]["tree"]["id"], objects["supportOrganization"]["tree"]["id"])
        self.assertNotIn("supportNodeId", project["supportOrganization"]["tree"]["children"][0])
        self.assertIsNot(project["supportOrganization"], objects["supportOrganization"])
        self.assertNotIn("basicMission", project)
        self.assertGreaterEqual(len(project["basicMissions"]), 2)
        self.assertTrue(all(basic.get("id") for basic in project["basicMissions"]))
        basic_ids = {basic["id"] for basic in project["basicMissions"]}
        self.assertTrue(all(
            item.get("basicMissionId") in basic_ids
            for composite in project["missionProfile"]["compositeTasks"]
            for item in composite.get("taskItems", [])
        ))
        self.assertEqual(project["airports"], ["A"])
        self.assertNotIn("carrier-deck", project["airports"])
        self.assertNotIn("experiment", project)
        self.assertNotIn("analysisRequests", project)
        self.assertEqual(project["reliabilityBlockDiagram"], objects["reliabilityBlockDiagram"])
        self.assertIsNot(project["reliabilityBlockDiagram"], objects["reliabilityBlockDiagram"])
        self.assertNotIn("monteCarlo", project)
        self.assertNotIn("monteCarlo", project["missionProfile"])
        self.assertNotIn("analysisRequests", project["missionProfile"])
        self.assertNotIn("profileType", project["missionProfile"])
        self.assertNotIn("endCondition", project["missionProfile"])
        self.assertNotIn("repeatCycleHours", project["missionProfile"])
        objects["analysisRequests"]["largeSample"]["sweep"]["failureRates"].append(0.99)
        self.assertNotIn("analysisRequests", project)

    def test_modeling_import_to_project_splits_support_nodes_resources_and_transport_policies(self) -> None:
        import_package = self._fixture("modeling_import_project.json")

        project = modeling_import_to_project(import_package)

        self.assertEqual([node["name"] for node in project["supportNodes"]], ["基地", "中继", "基层1"])
        for node in project["supportNodes"]:
            self.assertEqual(set(node), {"id", "name", "organizationNodeId"})
            self.assertTrue(str(node["id"]).startswith("support-node-"))

        resource_types = {resource["type"] for resource in project["supportResources"]}
        self.assertEqual(resource_types, {"personnel", "equipment", "spare"})
        deck_spares = [
            resource for resource in project["supportResources"]
            if resource["supportNodeName"] == "基地" and resource["type"] == "spare"
        ]
        self.assertTrue(any(resource["name"] == "航电模块" and resource["quantity"] == 6 for resource in deck_spares))
        deck_personnel = [
            resource for resource in project["supportResources"]
            if resource["supportNodeName"] == "基地" and resource["type"] == "personnel"
        ]
        self.assertEqual(deck_personnel[0]["quantity"], 5)
        deck_equipment = [
            resource for resource in project["supportResources"]
            if resource["supportNodeName"] == "基地" and resource["type"] == "equipment"
        ]
        self.assertEqual(deck_equipment[0]["quantity"], 3)

        self.assertGreaterEqual(len(project["transportPolicies"]), 3)
        first_policy = project["transportPolicies"][0]
        self.assertIn("fromOrganizationNodeId", first_policy)
        self.assertIn("toOrganizationNodeId", first_policy)
        self.assertNotIn("transportPolicies", project["supportResources"][0])
        self.assertIsInstance(project["supportOrganization"]["tree"], dict)
        self.assertIn("children", project["supportOrganization"]["tree"])

    def test_modeling_import_to_project_blocks_self_legacy_activity_transport_strategy(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["objects"] = copy.deepcopy(import_package["objects"])
        import_package["objects"].pop("transportPolicies", None)
        for mission in import_package["objects"]["missionProfiles"]:
            for basic_mission in mission.get("basicMissions", []):
                basic_mission["supportActivityName"] = "后勤保障活动方案"
        import_package["objects"]["supportActivities"] = [
            {
                "id": "logistics-plan",
                "name": "后勤保障活动方案",
                "activityName": "后勤保障活动方案",
                "equipmentId": import_package["objects"]["equipmentAssets"][0]["id"],
                "resourceId": import_package["objects"]["supportResources"][0]["id"],
                "durationHours": 1,
                "activityType": "后勤保障",
                "transportStrategies": [
                    {
                        "name": "旧调运策略",
                        "direction": "横向运输",
                        "spareType": "航电模块",
                        "triggerMode": "周期性调运",
                        "transferCycleHours": 6,
                        "from": import_package["objects"]["supportResources"][0]["id"],
                        "to": import_package["objects"]["supportResources"][1]["id"],
                        "transportTimeHours": 2,
                    }
                ],
                "organizationStrategies": [{"supportLevel": "base"}],
            }
        ]

        validation = validate_modeling_import_package(import_package)
        self.assertTrue(validation["ok"])
        with self.assertRaisesRegex(Exception, "endpoints must be different"):
            modeling_import_to_project(import_package, validation=validation)

    def test_strip_project_sweep_removes_support_node_resource_rows_and_preserves_resources(self) -> None:
        project = {
            "project_id": "project-support-boundary",
            "scenarioId": "support-boundary",
            "modelingDictionaries": {"personnelSpecialties": ["航电"]},
            "modelingImportValidation": {
                "importId": "import-support-boundary",
                "usedTables": {"supportResources": True},
                "validationLevel": "level1",
            },
            "supportOrganization": {
                "tree": {
                    "id": "support-org-root",
                    "name": "舰载保障组织",
                    "children": [
                        {"id": "carrier-deck", "name": "基地"},
                        {"id": "forward-sea-base", "name": "中继"},
                        {"id": "line-team", "name": "基层"},
                    ],
                }
            },
            "supportResourceOverrides": {"root:line-team:personnel": {"quantity": 3}},
            "deletedSupportResourceKeys": ["root:line-team:personnel"],
            "supportNodes": [
                {
                    "id": "carrier-deck",
                    "name": "基地",
                    "capacity": 4,
                    "equipmentCapacity": 3,
                    "inventory": {"航电模块": 6},
                    "lateralSupportNodes": ["基层"],
                    "nodeType": "甲板保障点",
                    "organizationStrategy": "任务优先",
                    "personnelCapacity": 5,
                    "policy": "高优先级",
                    "supportLevel": "一线保障",
                    "transportPolicies": [{"from": "line-team", "to": "carrier-deck"}],
                },
                {"id": "forward-sea-base", "name": "中继", "personnelCapacity": 2, "equipmentCapacity": 2},
                {"id": "line-team", "name": "基层", "personnelCapacity": 3, "equipmentCapacity": 3},
                {
                    "id": "carrier-stock-personnel-mech",
                    "name": "机械保障人员",
                    "organizationNodeId": "line-team",
                    "importedResourceType": "personnel",
                    "personnelModel": "机械",
                    "personnelCapacity": 3,
                },
            ],
            "supportResources": [
                {
                    "id": "personnel-1",
                    "supportNodeName": "基层",
                    "type": "personnel",
                    "name": "机械保障人员",
                    "model": "机械",
                    "quantity": 3,
                    "equipment": "J-15",
                    "equipmentId": "aircraft-type-j15",
                }
            ],
            "transportPolicies": [
                {
                    "id": "transport-1",
                    "from": "line-team",
                    "to": "carrier-deck",
                    "fromSupportNodeName": "基层",
                    "toSupportNodeName": "基地",
                    "spareName": "航电模块",
                    "capacity": 2,
                }
            ],
            "supportActivities": [
                {
                    "id": "logistics-plan",
                    "activityType": "后勤保障",
                    "useCalendarRule": True,
                    "useFlightHourRule": True,
                    "useTakeoffLandingRule": False,
                    "calendarDayFloatRatio": 0.2,
                    "runHourFloatRatio": 0.3,
                    "takeoffLandingFloatRatio": 0.4,
                    "transportStrategies": [{"from": "line-team", "to": "carrier-deck", "spareType": "航电模块"}],
                    "organizationStrategies": [{"supportLevel": "base"}],
                }
            ],
        }

        slim_project = strip_project_sweep(project)

        self.assertNotIn("supportResourceOverrides", slim_project)
        self.assertNotIn("deletedSupportResourceKeys", slim_project)
        self.assertNotIn("modelingDictionaries", slim_project)
        self.assertNotIn("validationLevel", slim_project["modelingImportValidation"])
        self.assertEqual([node["name"] for node in slim_project["supportNodes"]], ["基地", "中继", "基层"])
        self.assertTrue(all("organizationNodeId" in node for node in slim_project["supportNodes"]))
        self.assertFalse(any(node.get("id") == "carrier-stock-personnel-mech" for node in slim_project["supportNodes"]))
        self.assertEqual(slim_project["supportResources"], [
            {
                "id": "personnel-1",
                "supportNodeName": "基层",
                "type": "personnel",
                "name": "机械保障人员",
                "model": "机械",
                "quantity": 3,
                "organizationNodeId": "line-team",
            }
        ])
        self.assertEqual(slim_project["transportPolicies"], [
            {
                "id": "transport-1",
                "fromOrganizationNodeId": "line-team",
                "toOrganizationNodeId": "carrier-deck",
                "capacity": 2,
            }
        ])
        self.assertEqual(slim_project["supportActivities"], [{
            "id": "logistics-plan",
            "activityName": "logistics-plan",
            "activityType": "后勤保障",
            "planType": "后勤保障方案",
        }])

    def test_strip_project_sweep_promotes_legacy_activity_transport_strategies(self) -> None:
        project = {
            "project_id": "project-legacy-logistics",
            "scenarioId": "legacy-logistics",
            "supportNodes": [
                {"id": "base", "name": "基地"},
                {"id": "deck", "name": "甲板"},
            ],
            "supportActivities": [
                {
                    "id": "logistics-plan",
                    "activityType": "后勤保障",
                    "transportStrategies": [
                        {
                            "name": "旧调运策略",
                            "direction": "横向运输",
                            "spareType": "航电模块",
                            "triggerMode": "临界库存",
                            "criticalInventory": 2,
                            "from": "base",
                            "to": "deck",
                            "transportTimeHours": 1.5,
                        }
                    ],
                    "organizationStrategies": [{"supportLevel": "base"}],
                }
            ],
        }

        slim_project = strip_project_sweep(project)

        self.assertEqual(slim_project["transportPolicies"], [
            {
                "id": "migrated-transport-b8e701a8c223",
                "fromOrganizationNodeId": "support-node-1",
                "toOrganizationNodeId": "support-node-2",
                "name": "旧调运策略",
                "direction": "横向运输",
                "triggerMode": "临界库存",
                "criticalInventory": 2,
                "transportMode": "横向运输",
                "transportTimeHours": 1.5,
            }
        ])
        self.assertNotIn("transportStrategies", slim_project["supportActivities"][0])
        self.assertNotIn("organizationStrategies", slim_project["supportActivities"][0])

    def test_modeling_import_to_project_derives_airports_from_combat_unit_members(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        members = import_package["objects"]["missionProfiles"][0]["combatUnit"]["members"]
        members[0]["airport"] = "A"
        members[1]["airport"] = "B"
        for member in members[2:]:
            member["airport"] = "A"

        project = modeling_import_to_project(import_package)

        self.assertEqual(project["airports"], ["A", "B"])
        self.assertFalse({"carrier-deck", "forward-sea-base"} & set(project["airports"]))

    def test_modeling_import_to_project_does_not_derive_airports_from_deployment_location(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["objects"].pop("airports", None)
        mission = import_package["objects"]["missionProfiles"][0]
        mission.pop("airports", None)
        mission["combatUnit"].pop("airport", None)
        mission["combatUnit"]["deploymentLocation"] = "航母飞行甲板"
        for member in mission["combatUnit"]["members"]:
            member.pop("airport", None)
            member.pop("airportName", None)
            member.pop("deploymentAirport", None)

        project = modeling_import_to_project(import_package)

        self.assertEqual(project["airports"], [])

    def test_modeling_import_to_project_preserves_explicit_empty_collections(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        mission = import_package["objects"]["missionProfiles"][0]
        mission["combatUnit"] = {"members": []}
        import_package["objects"]["reliabilityBlockDiagram"] = {"nodes": [], "edges": []}
        import_package["objects"]["monteCarlo"] = {
            "failureRates": [],
            "spareMultipliers": [],
            "supportCapacities": [],
            "minRequiredSorties": [],
        }

        project = modeling_import_to_project(import_package)

        self.assertEqual(project["combatUnit"], {"members": []})
        self.assertEqual(project["reliabilityBlockDiagram"], {"nodes": [], "edges": []})
        self.assertNotIn("monteCarlo", project)
        self.assertNotIn("monteCarlo", project["missionProfile"])

    def test_m4_regular_user_cannot_publish_modeling_import_and_denial_is_audited(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        data_session = self.api.login("data", "data")
        user_session = self.api.login("user", "user")
        self.api.save_modeling_import(import_package, actor_user_id=data_session["user"]["user_id"])

        with self.assertRaises(BackendApiError) as ctx:
            self.api.publish_modeling_import(
                import_package["importId"],
                actor_user_id=user_session["user"]["user_id"],
            )

        self.assertEqual(ctx.exception.code, "forbidden")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(events[-1]["action"], "modeling_import.publish")
        self.assertEqual(events[-1]["outcome"], "denied")
        self.assertEqual(events[-1]["actor_user_id"], user_session["user"]["user_id"])

    def test_m4_regular_user_cannot_save_modeling_import_and_denial_is_audited(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        user_session = self.api.login("user", "user")

        with self.assertRaises(BackendApiError) as ctx:
            self.api.save_modeling_import(
                import_package,
                actor_user_id=user_session["user"]["user_id"],
            )

        self.assertEqual(ctx.exception.code, "forbidden")
        events = self.repository.list_audit_events(resource_id=import_package["importId"])
        self.assertEqual(events[-1]["action"], "modeling_import.save")
        self.assertEqual(events[-1]["outcome"], "denied")

    def test_m4_admin_can_list_create_and_update_users_with_audit(self) -> None:
        admin_session = self.api.login("admin", "admin")

        created = self.api.create_user(
            {
                "username": "planner",
                "password": "planner",
                "role": "数据管理员",
                "display_name": "规划员",
                "status": "active",
            },
            actor_user_id=admin_session["user"]["user_id"],
        )
        updated = self.api.update_user(
            created["user_id"],
            {
                "display_name": "规划员二号",
                "role": "普通用户",
                "status": "disabled",
            },
            actor_user_id=admin_session["user"]["user_id"],
        )
        deleted = self.api.delete_user(created["user_id"], actor_user_id=admin_session["user"]["user_id"])
        users = self.api.list_users(actor_user_id=admin_session["user"]["user_id"])

        self.assertEqual(created["username"], "planner")
        self.assertNotIn("password_hash", created)
        self.assertEqual(updated["display_name"], "规划员二号")
        self.assertEqual(updated["role"], "普通用户")
        self.assertEqual(updated["status"], "disabled")
        self.assertEqual(deleted["deleted"], True)
        self.assertNotIn("planner", {user["username"] for user in users["users"]})
        events = self.repository.list_audit_events(resource_id=created["user_id"])
        self.assertEqual(
            [(event["action"], event["outcome"]) for event in events],
            [
                ("users.create", "allowed"),
                ("users.update", "allowed"),
                ("users.delete", "allowed"),
            ],
        )

    def test_stage5_system_config_can_be_saved_only_by_admin_and_is_audited(self) -> None:
        admin_session = self.api.login("admin", "admin")
        data_session = self.api.login("data", "data")
        user_session = self.api.login("user", "user")
        payload = {
            "projectDataModules": [{"key": "modeling-data-source", "sheetKeys": ["equipment-system"]}],
            "granularityProfiles": [{"key": "granularity-a", "fieldKeys": ["equipment-system:mtbfHours"]}],
            "permissions": [{"feature": "项目管理", "admin": "编辑", "data": "编辑", "user": "只读"}],
            "modelingForms": {"fieldUnits": {"equipment-system:mtbfHours": "小时"}, "personnelSpecialties": ["机务"]},
        }

        missing = self.api.get_system_config("system-runtime-support")
        saved = self.api.save_system_config(
            "system-runtime-support",
            payload,
            actor_user_id=admin_session["user"]["user_id"],
        )
        loaded = self.api.get_system_config("system-runtime-support")

        self.assertEqual(missing["payload"], {})
        self.assertEqual(saved["payload"]["modelingForms"]["fieldUnits"]["equipment-system:mtbfHours"], "小时")
        self.assertEqual(loaded["payload"]["modelingForms"]["personnelSpecialties"], ["机务"])
        with self.assertRaises(BackendApiError) as data_forbidden_ctx:
            self.api.save_system_config(
                "system-runtime-support",
                payload,
                actor_user_id=data_session["user"]["user_id"],
            )
        self.assertEqual(data_forbidden_ctx.exception.code, "forbidden")
        with self.assertRaises(BackendApiError) as forbidden_ctx:
            self.api.save_system_config(
                "system-runtime-support",
                payload,
                actor_user_id=user_session["user"]["user_id"],
            )
        self.assertEqual(forbidden_ctx.exception.code, "forbidden")

        events = self.repository.list_audit_events(resource_id="system-runtime-support")
        self.assertEqual(events[-3]["action"], "system_config.save")
        self.assertEqual(events[-3]["outcome"], "allowed")
        self.assertEqual(events[-2]["action"], "system_config.save")
        self.assertEqual(events[-2]["outcome"], "denied")
        self.assertEqual(events[-1]["action"], "system_config.save")
        self.assertEqual(events[-1]["outcome"], "denied")

    def test_m4_regular_user_cannot_create_or_update_users_and_denial_is_audited(self) -> None:
        admin_session = self.api.login("admin", "admin")
        user_session = self.api.login("user", "user")
        created = self.api.create_user(
            {"username": "readonly", "password": "readonly", "role": "普通用户", "display_name": "只读用户"},
            actor_user_id=admin_session["user"]["user_id"],
        )

        with self.assertRaises(BackendApiError) as create_ctx:
            self.api.create_user(
                {"username": "blocked", "password": "blocked", "role": "普通用户"},
                actor_user_id=user_session["user"]["user_id"],
            )
        with self.assertRaises(BackendApiError) as update_ctx:
            self.api.update_user(
                created["user_id"],
                {"display_name": "不应修改"},
                actor_user_id=user_session["user"]["user_id"],
            )
        with self.assertRaises(BackendApiError) as delete_ctx:
            self.api.delete_user(
                created["user_id"],
                actor_user_id=user_session["user"]["user_id"],
            )

        self.assertEqual(create_ctx.exception.code, "forbidden")
        self.assertEqual(update_ctx.exception.code, "forbidden")
        self.assertEqual(delete_ctx.exception.code, "forbidden")
        create_events = self.repository.list_audit_events(resource_id="blocked")
        update_events = self.repository.list_audit_events(resource_id=created["user_id"])
        self.assertEqual(create_events[-1]["action"], "users.create")
        self.assertEqual(create_events[-1]["outcome"], "denied")
        self.assertEqual(update_events[-2]["action"], "users.update")
        self.assertEqual(update_events[-2]["outcome"], "denied")
        self.assertEqual(update_events[-1]["action"], "users.delete")
        self.assertEqual(update_events[-1]["outcome"], "denied")

    def test_compile_modeling_import_scenario_requires_published_valid_import(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)

        with self.assertRaises(BackendApiError) as ctx:
            self.api.compile_modeling_import_scenario(import_package["importId"])

        self.assertEqual(ctx.exception.code, "unpublished_modeling_import")

    def test_compile_modeling_import_scenario_uses_simulation_adapter(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        compiled = self.api.compile_modeling_import_scenario(import_package["importId"])

        self.assertEqual(compiled["compiled_from_import"]["import_id"], import_package["importId"])
        self.assertEqual(compiled["project"]["project_id"], import_package["projectId"])
        self.assertEqual(compiled["scenario"]["project_id"], import_package["projectId"])
        self.assertEqual(compiled["scenario"]["simulation_model"]["family"], "aircraft_support_v1")
        self.assertEqual(compiled["scenario"]["compiled_by"], "Simulation Adapter Agent")
        self.assertEqual(len(self.adapter.compile_calls), 1)
        self.assertEqual(self.adapter.compile_calls[0][1], "aircraft_support_v1")

    def test_reduced_scope_modeling_import_omits_support_domains_without_validation_failure(self) -> None:
        import_package = self._reduced_scope_import_package_without_support_domain()

        validation = validate_modeling_import_package(import_package)

        self.assertTrue(validation["ok"])
        self.assertEqual(validation["status"], "valid")
        self.assertNotIn("validationLevel", validation)
        self.assertFalse(validation["usedTables"]["supportResources"])
        self.assertFalse(validation["usedTables"]["supportActivities"])
        self.assertEqual(validation["issues"], [])
        scope_codes = [warning["code"] for warning in validation["warnings"]]
        self.assertIn("scope_not_modeled", scope_codes)
        self.assertTrue(any(warning["field_path"] == "objects.supportResources" for warning in validation["warnings"]))

    def test_declared_scope_modeling_import_rejects_declared_missing_support_domains(self) -> None:
        import_package = self._reduced_scope_import_package_without_support_domain()
        import_package["usedTables"]["supportResources"] = True
        import_package["usedTables"]["supportActivities"] = True

        validation = validate_modeling_import_package(import_package)

        self.assertFalse(validation["ok"])
        self.assertNotIn("validationLevel", validation)
        self.assertEqual(validation["usedTables"]["supportResources"], True)
        issues_by_path = {issue["field_path"]: issue for issue in validation["issues"]}
        self.assertEqual(issues_by_path["objects.supportResources"]["code"], "invalid_declared_table")
        self.assertEqual(issues_by_path["objects.supportActivities"]["code"], "invalid_declared_table")
        self.assertTrue(all(issue["severity"] == "error" for issue in validation["issues"]))

    def test_modeling_import_rejects_retired_validation_level(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["validationLevel"] = "retired"

        validation = validate_modeling_import_package(import_package)
        issues_by_path = {issue["field_path"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        self.assertNotIn("validationLevel", validation)
        self.assertEqual(issues_by_path["validationLevel"]["code"], "retired_validation_level")

    def test_modeling_import_rejects_basic_mission_support_activity_name_matching_only_legacy_name(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        mission = import_package["objects"]["missionProfiles"][0]
        legacy_name = import_package["objects"]["supportActivities"][0]["name"]
        self.assertNotEqual(legacy_name, import_package["objects"]["supportActivities"][0]["activityName"])
        mission["basicMissions"][0]["supportActivityName"] = legacy_name

        validation = validate_modeling_import_package(import_package)
        issues_by_path = {issue["field_path"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        issue = issues_by_path["objects.missionProfiles[0].basicMissions[0].supportActivityName"]
        self.assertEqual(issue["code"], "invalid_basic_mission_support_activity_name")

    def test_declared_scope_modeling_import_rejects_declared_missing_support_scope_tables(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["usedTables"] = {
            "missionProfiles": True,
            "equipmentAssets": True,
            "reliabilityBlockDiagram": True,
            "supportResources": True,
            "supportActivities": True,
            "supportOrganization": True,
            "transportPolicies": True,
        }
        import_package["objects"] = copy.deepcopy(import_package["objects"])
        import_package["objects"].pop("supportOrganization", None)
        import_package["objects"].pop("transportPolicies", None)

        validation = validate_modeling_import_package(import_package)
        issues_by_path = {issue["field_path"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        self.assertEqual(issues_by_path["objects.supportOrganization"]["code"], "invalid_declared_table")
        self.assertEqual(issues_by_path["objects.transportPolicies"]["code"], "invalid_declared_table")

    def test_declared_scope_modeling_import_rejects_declared_empty_support_scope_tables(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["usedTables"] = {
            "missionProfiles": True,
            "equipmentAssets": True,
            "reliabilityBlockDiagram": True,
            "supportResources": True,
            "supportActivities": True,
            "supportOrganization": True,
            "transportPolicies": True,
        }
        import_package["objects"] = copy.deepcopy(import_package["objects"])
        import_package["objects"]["supportResources"] = []
        import_package["objects"]["supportActivities"] = []
        import_package["objects"]["reliabilityBlockDiagram"] = {}
        import_package["objects"]["supportOrganization"] = {"tree": []}

        validation = validate_modeling_import_package(import_package)
        issues_by_path = {issue["field_path"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        self.assertEqual(issues_by_path["objects.supportResources"]["code"], "invalid_declared_table")
        self.assertEqual(issues_by_path["objects.supportActivities"]["code"], "invalid_declared_table")
        self.assertEqual(issues_by_path["objects.reliabilityBlockDiagram"]["code"], "invalid_declared_table")
        self.assertEqual(issues_by_path["objects.supportOrganization.tree"]["code"], "invalid_declared_table")

    def test_compile_modeling_import_scenario_returns_gate_envelope_for_reduced_scope_import(self) -> None:
        import_package = self._reduced_scope_import_package_without_support_domain()
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        compiled = self.api.compile_modeling_import_scenario(import_package["importId"])

        self.assertEqual(compiled["status"], "compiled")
        self.assertNotIn("validationLevel", compiled)
        self.assertFalse(compiled["usedTables"]["supportResources"])
        self.assertFalse(compiled["usedTables"]["supportActivities"])
        self.assertEqual(compiled["issues"], [])
        self.assertTrue(any(warning["code"] == "scope_not_modeled" for warning in compiled["warnings"]))
        self.assertEqual(compiled["compiled_from_import"]["import_id"], import_package["importId"])
        self.assertEqual(compiled["scenario"]["simulation_model"]["family"], "aircraft_support_v1")
        self.assertNotIn("validation_level", compiled["provenance"])
        self.assertIn("supportResources", compiled["provenance"]["disabled_domains"])
        self.assertIn("supportActivities", compiled["provenance"]["disabled_domains"])
        self.assertEqual(len(self.adapter.compile_calls), 1)
        self.assertEqual(self.adapter.compile_calls[0][1], "aircraft_support_v1")

    def test_compile_modeling_import_scenario_reuses_validation_for_project_conversion(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        with mock.patch(
            "src.spare_mvp_backend.modeling_import.validate_modeling_import_package",
            side_effect=AssertionError("duplicate modeling import validation"),
        ):
            compiled = self.api.compile_modeling_import_scenario(import_package["importId"])

        self.assertEqual(compiled["status"], "compiled")
        self.assertEqual(compiled["compiled_from_import"]["import_id"], import_package["importId"])

    def test_compile_modeling_import_scenario_rejects_retired_model_family(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        with self.assertRaises(BackendApiError) as ctx:
            self.api.compile_modeling_import_scenario(
                import_package["importId"],
                model_family="aviation_support",
            )

        self.assertEqual(ctx.exception.code, "retired_model_family")
        self.assertEqual(ctx.exception.details["model_family"], "aviation_support")
        self.assertEqual(ctx.exception.details["replacement_model_family"], "aircraft_support_v1")
        self.assertIn("aviation_support", ctx.exception.details["retired_model_families"])
        self.assertEqual(self.adapter.compile_calls, [])

    def test_modeling_import_api_reports_field_level_issues(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["objects"]["supportActivities"][0]["resourceId"] = "missing-resource"
        import_package["objects"]["equipmentAssets"][1]["quantity"] = 0

        validation = self.api.validate_modeling_import(import_package)

        self.assertFalse(validation["ok"])
        self.assertEqual(validation["status"], "invalid")
        self.assertEqual(
            sorted(issue["code"] for issue in validation["issues"]),
            ["invalid_number", "missing_reference"],
        )
        self.assertEqual(
            {issue["field_path"] for issue in validation["issues"]},
            {
                "objects.equipmentAssets[1].quantity",
                "objects.supportActivities[0].resourceId",
            },
        )

    def test_modeling_import_api_rejects_sru_parent_that_is_not_lru(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["objects"]["equipmentAssets"] = [
            {"id": "aircraft-root", "name": "整机", "quantity": 1, "productType": ""},
            None,
            {"id": "hydraulic-system", "name": "液压系统", "parentId": "aircraft-root", "quantity": 1, "productType": "SRU"},
            {"id": "hydraulic-valve", "name": "液压阀", "parentId": "hydraulic-system", "quantity": 1, "productType": "LRU"},
        ]

        validation = self.api.validate_modeling_import(import_package)
        issues_by_code = {issue["code"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        self.assertEqual(
            issues_by_code["invalid_sru_parent"]["field_path"],
            "objects.equipmentAssets[2].parentId",
        )
        self.assertEqual(issues_by_code["invalid_sru_parent"]["page"], "装备系统建模")

    def test_modeling_import_api_covers_contract_parity_issues(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["schemaVersion"] = "modeling-import-v0"
        import_package["objects"].pop("supportResources")
        import_package["objects"]["equipmentAssets"].append(
            {
                "id": "j15-radar",
                "name": "重复雷达 LRU",
                "parentId": "aircraft-root",
                "quantity": 1,
                "mtbfHours": 900,
            }
        )
        import_package["lifecycle"] = {
            "state": "published",
            "version": 1,
            "referencedRunIds": ["run-aircraft-support-contract-001"],
        }
        import_package["changes"] = [
            {
                "operation": "update",
                "objectType": "equipmentAssets",
                "objectId": "j15-radar",
                "fieldPath": "objects.equipmentAssets[id=j15-radar].name",
            }
        ]

        validation = self.api.validate_modeling_import(import_package)
        issues_by_code = {issue["code"]: issue for issue in validation["issues"]}

        self.assertFalse(validation["ok"])
        self.assertIn("invalid_schema_version", issues_by_code)
        self.assertIn("invalid_declared_table", issues_by_code)
        self.assertIn("duplicate_id", issues_by_code)
        self.assertIn("published_reference_protection", issues_by_code)
        self.assertEqual(
            issues_by_code["published_reference_protection"]["field_path"],
            "objects.equipmentAssets[id=j15-radar].name",
        )
        self.assertEqual(
            issues_by_code["published_reference_protection"]["page"],
            "装备系统建模",
        )

    def test_modeling_import_api_rejects_invalid_save_and_referenced_publish(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        invalid_package = copy.deepcopy(import_package)
        invalid_package["objects"]["missionProfiles"][0].pop("name")

        with self.assertRaises(BackendApiError) as invalid_ctx:
            self.api.save_modeling_import_as_system(invalid_package)

        self.assertEqual(invalid_ctx.exception.code, "invalid_modeling_import")
        self.assertEqual(
            invalid_ctx.exception.details["issues"][0]["field_path"],
            "objects.missionProfiles[0].name",
        )

        referenced_package = copy.deepcopy(import_package)
        referenced_package["lifecycle"] = {
            "state": "published",
            "version": 1,
            "referencedRunIds": ["run-aircraft-support-contract-001"],
        }
        self.repository.upsert_modeling_import(
            referenced_package,
            {"ok": True, "status": "valid", "issues": []},
        )

        with self.assertRaises(BackendApiError) as publish_ctx:
            self.api.publish_modeling_import_as_system(import_package["importId"])

        self.assertEqual(publish_ctx.exception.code, "published_import_referenced")

        changed_package = copy.deepcopy(import_package)
        changed_package["lifecycle"] = {"state": "draft", "version": 2, "referencedRunIds": []}
        target_index = next(
            index for index, component in enumerate(changed_package["objects"]["equipmentAssets"])
            if component["id"] == "j15-engine"
        )
        changed_package["objects"]["equipmentAssets"][target_index]["quantity"] = 3
        saved = self.api.save_modeling_import_as_system(changed_package)
        stored = self.api.get_modeling_import(import_package["importId"])

        self.assertEqual(saved["status"], "draft")
        self.assertEqual(stored["draftPackage"]["objects"]["equipmentAssets"][target_index]["quantity"], 3)
        self.assertEqual(stored["publishedPackage"]["objects"]["equipmentAssets"][target_index]["quantity"], 2)

        with self.assertRaises(BackendApiError) as republish_ctx:
            self.api.publish_modeling_import_as_system(import_package["importId"])

        self.assertEqual(republish_ctx.exception.code, "published_import_referenced")

    def test_compile_modeling_import_scenario_uses_persisted_published_snapshot_after_new_draft(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        self.api.save_modeling_import_as_system(import_package)
        self.api.publish_modeling_import_as_system(import_package["importId"])

        changed_package = copy.deepcopy(import_package)
        changed_package["lifecycle"] = {"state": "draft", "version": 2, "referencedRunIds": []}
        target_index = next(
            index for index, component in enumerate(changed_package["objects"]["equipmentAssets"])
            if component["id"] == "j15-engine"
        )
        changed_package["objects"]["equipmentAssets"][target_index]["quantity"] = 3
        self.api.save_modeling_import_as_system(changed_package)

        compiled = self.api.compile_modeling_import_scenario(import_package["importId"])

        self.assertEqual(compiled["compiled_from_import"]["import_version"], 1)
        self.assertEqual(compiled["project"]["project_version"], "import-v1")
        self.assertEqual(
            next(component for component in compiled["project"]["components"] if component["id"] == "j15-engine")["quantity"],
            2,
        )

    def test_modeling_import_api_rejects_invalid_lifecycle_before_compile(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["lifecycle"] = {
            "state": "published",
            "version": 0,
            "referencedRunIds": "run-aircraft-support-contract-001",
        }

        validation = self.api.validate_modeling_import(import_package)

        self.assertFalse(validation["ok"])
        self.assertEqual(
            [issue["code"] for issue in validation["issues"]],
            ["invalid_lifecycle_version", "invalid_lifecycle_references"],
        )

        with self.assertRaises(BackendApiError) as save_ctx:
            self.api.save_modeling_import_as_system(import_package)

        self.assertEqual(save_ctx.exception.code, "invalid_modeling_import")

    def test_modeling_import_api_accepts_json_schema_integer_version_semantics(self) -> None:
        import_package = self._fixture("modeling_import_project.json")
        import_package["lifecycle"]["version"] = 1.0

        validation = self.api.validate_modeling_import(import_package)

        self.assertTrue(validation["ok"])


if __name__ == "__main__":
    unittest.main()
