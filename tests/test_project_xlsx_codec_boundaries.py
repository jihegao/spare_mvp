from __future__ import annotations

import base64
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from src.spare_mvp_backend.api import BackendApi
from src.spare_mvp_backend.project_payload import ProjectJsonExporter
from src.spare_mvp_backend.project_xlsx import parse_project_xlsx
from src.spare_mvp_backend.project_xlsx_template import export_project_xlsx, GUIDE, FIELDS
from src.spare_mvp_backend.repository import ContractRepository, initialize_database
from src.spare_mvp_backend.xlsx_text import workbook_bytes
from src.spare_mvp_contract import SimulationAdapter

ROOT = Path(__file__).resolve().parents[1]


class ProjectExcelCodecBoundaryTest(unittest.TestCase):
    def test_binary_writer_preserves_every_newline_sequence_in_concurrent_exports(self):
        values = ['LF\nline', 'CR\rline', 'CRLF\r\nline', 'TAB\tline',
                  'CRCR\r\rline', 'mixed\r\r\n\n\r\t中文']

        def roundtrip(index):
            workbook = Workbook()
            sheet = workbook.active
            for value in values:
                sheet.append([f'{index}:{value}'])
            sheet.merge_cells('B1:C1')
            sheet['B1'] = 'merged'
            result = load_workbook(BytesIO(workbook_bytes(workbook)))
            try:
                self.assertEqual([result.active.cell(row, 1).value for row in range(1, len(values) + 1)],
                                 [f'{index}:{value}' for value in values])
                self.assertEqual(str(result.active.merged_cells), 'B1:C1')
            finally:
                result.close()

        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(roundtrip, range(8)))

    def test_cell_limit_is_checked_after_escaping_instead_of_losing_last_character(self):
        for prefix in ('@', '~', '='):
            with self.subTest(prefix=prefix):
                source = {'projectInfo': {'text': prefix + 'x' * 32765}}
                result, _, issues = parse_project_xlsx(export_project_xlsx(source))
                self.assertEqual(issues, [])
                self.assertEqual(result, source)
                with self.assertRaisesRegex(ValueError, '32767'):
                    export_project_xlsx({'projectInfo': {'text': prefix + 'x' * 32766}})
        source = {'projectInfo': {'text': 'x' * 32767}}
        self.assertEqual(parse_project_xlsx(export_project_xlsx(source))[0], source)
        with self.assertRaisesRegex(ValueError, '字段名'):
            export_project_xlsx({'projectInfo': {'x' * 32768: 'value'}})

    def test_schema_open_keys_and_control_characters_roundtrip_without_formula_coercion(self):
        source = {'projectInfo': {'': 'empty key', '=literal': 'not a formula', '@empty': None,
                                  '~key': [False, 0, '', '中文\r\n\t= @ ~'], '00': {}}}
        result, _, issues = parse_project_xlsx(export_project_xlsx(source))
        self.assertEqual(issues, [])
        self.assertEqual(result, source)

    def test_malformed_workbook_never_writes_database_and_valid_source_compiles(self):
        project = ProjectJsonExporter(repo_root=ROOT).export(json.loads(
            (ROOT / 'tests/fixtures/aircraft_support_v1_operations_project.json').read_text()))
        content = export_project_xlsx(project)
        connection = sqlite3.connect(':memory:')
        initialize_database(connection)
        before = connection.total_changes
        mutations = [
            ('formula_not_supported', lambda wb: setattr(wb[GUIDE]['B3'], 'value', '=1+1')),
            ('formula_not_supported', lambda wb: setattr(wb[FIELDS]['H2'], 'value', '=1+1')),
            ('formula_not_supported', lambda wb: setattr(wb['项目信息']['F2'], 'value', '=1+1')),
            ('unknown_sheet', lambda wb: wb.create_sheet('Injected')),
            ('invalid_cell_type', lambda wb: setattr(wb['项目信息']['E2'], 'value', 'would be lost')),
        ]
        try:
            with tempfile.TemporaryDirectory() as output:
                api = BackendApi(ContractRepository(connection), SimulationAdapter(ROOT), output_dir=output)
                baseline = api.preview_project_xlsx({'content_base64': base64.b64encode(content).decode()})
                self.assertTrue(baseline['ok'], baseline['errors'])
                for code, mutate in mutations:
                    with self.subTest(code=code):
                        workbook = load_workbook(BytesIO(content))
                        mutate(workbook)
                        altered = workbook_bytes(workbook)
                        workbook.close()
                        result = api.preview_project_xlsx({'content_base64': base64.b64encode(altered).decode()})
                        self.assertFalse(result['ok'])
                        self.assertIn(code, {error['code'] for error in result['errors']})
                        self.assertEqual(connection.total_changes, before)
        finally:
            connection.close()
