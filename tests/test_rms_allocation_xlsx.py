from __future__ import annotations

import base64
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
from threading import Thread
import unittest
from urllib import error, request

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.spare_mvp_backend.api import BackendApi, BackendApiError
from src.spare_mvp_backend.http_server import create_backend_server
from src.spare_mvp_backend.repository import ContractRepository, initialize_database
from src.spare_mvp_backend.rms_allocation_xlsx import (
    RMS_METADATA_SHEET,
    RMS_RESULT_SHEET,
    RMS_XLSX_SCHEMA_VERSION,
    RmsAllocationXlsxError,
    export_rms_allocation_xlsx,
    parse_rms_allocation_xlsx,
)
from src.spare_mvp_contract.adapter import SimulationAdapter


def rms_payload() -> dict:
    return {
        "project_id": "project-rms-001",
        "project_name": "RMS案例",
        "aircraft_model": "J-15",
        "basic_mission_id": "mission-sortie",
        "basic_mission_name": "基本出动任务",
        "mission_hours": 3,
        "mission_reliability": 0.95,
        "mttr_hours": 2,
        "plan_id": "RMS-PLAN-001",
        "plan_version": 3,
        "algorithm_version": "rms-engine-6.0.0",
        "method": "equal",
        "generated_at": "2026-07-20T12:00:00+08:00",
        "rows": [
            {
                "nodeId": "system-engine",
                "parentId": "aircraft-j15",
                "level": "系统",
                "nodeName": "动力系统",
                "model": "SYS-001",
                "installationCount": 2,
                "runningRatio": 0.8,
                "allocationShare": 0.625,
                "localAllocationShare": 0.625,
                "cumulativeAllocationShare": 0.625,
                "cumulativeInstallationCount": 2,
                "failureRate": 0.000833333333,
                "mtbfHours": 1200,
                "mttrHours": 2.5,
                "status": "已分配",
            },
            {
                "nodeId": "lru-controller",
                "parentId": "system-engine",
                "level": "LRU",
                "nodeName": "控制器",
                "model": "LRU-001",
                "installationCount": 1,
                "runningRatio": 1,
                "allocationShare": 0.375,
                "localAllocationShare": 0.6,
                "cumulativeAllocationShare": 0.375,
                "cumulativeInstallationCount": 2,
                "failureRate": 0.001,
                "mtbfHours": 1000,
                "mttrHours": 1.5,
                "status": "已分配",
            },
        ],
    }


class RmsAllocationXlsxTest(unittest.TestCase):
    def test_export_and_parse_round_trip_preserves_context_and_node_identity(self) -> None:
        download = export_rms_allocation_xlsx(rms_payload())

        workbook = load_workbook(BytesIO(download["body"]), read_only=True, data_only=False)
        self.assertEqual(workbook.sheetnames, [RMS_RESULT_SHEET, RMS_METADATA_SHEET])
        result_sheet = workbook[RMS_RESULT_SHEET]
        self.assertEqual(
            [result_sheet.cell(5, column).value for column in range(1, 7)],
            ["层级", "节点", "运行比", "失效率", "MTBF(h)", "MTTR(h)"],
        )
        self.assertEqual(result_sheet.cell(5, 7).value, "节点ID")
        self.assertEqual(result_sheet.cell(5, 8).value, "父节点ID")
        self.assertEqual(result_sheet.cell(5, 11).value, "分配份额")
        metadata = {
            row[0]: row[2]
            for row in workbook[RMS_METADATA_SHEET].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(metadata["schemaVersion"], RMS_XLSX_SCHEMA_VERSION)
        self.assertEqual(metadata["projectId"], "project-rms-001")
        self.assertEqual(metadata["basicMissionId"], "mission-sortie")
        self.assertEqual(metadata["missionReliability"], 0.95)
        self.assertEqual(metadata["planVersion"], 3)
        workbook.close()

        preview = parse_rms_allocation_xlsx(download["body"])

        self.assertTrue(preview["ok"])
        self.assertEqual(preview["errors"], [])
        self.assertEqual(preview["metadata"]["aircraftModel"], "J-15")
        self.assertEqual(preview["metadata"]["basicMissionName"], "基本出动任务")
        self.assertEqual(preview["metadata"]["planVersion"], 3)
        self.assertEqual(preview["rows"][0]["nodeId"], "system-engine")
        self.assertEqual(preview["rows"][0]["parentNodeId"], "aircraft-j15")
        self.assertEqual(preview["rows"][0]["allocationShare"], 0.625)
        self.assertEqual(preview["rows"][0]["localAllocationShare"], 0.625)
        self.assertEqual(preview["rows"][1]["cumulativeInstallationCount"], 2)
        self.assertEqual(preview["rows"][1]["mtbfHours"], 1000)

    def test_parse_reports_duplicate_node_invalid_values_and_formulas_with_locations(self) -> None:
        download = export_rms_allocation_xlsx(rms_payload())
        workbook = load_workbook(BytesIO(download["body"]))
        result_sheet = workbook[RMS_RESULT_SHEET]
        result_sheet.cell(7, 7).value = "system-engine"
        result_sheet.cell(6, 3).value = 1.5
        result_sheet.cell(6, 5).value = "=600*2"
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        preview = parse_rms_allocation_xlsx(output.getvalue())

        self.assertFalse(preview["ok"])
        errors = {error["code"]: error for error in preview["errors"]}
        self.assertEqual(errors["duplicate_node_id"]["row"], 7)
        self.assertEqual(errors["number_out_of_range"]["field"], "runningRatio")
        self.assertEqual(errors["formula_not_supported"]["column"], 5)
        self.assertEqual(errors["invalid_number"]["field"], "mtbfHours")

    def test_parse_rejects_non_rms_workbook(self) -> None:
        workbook = load_workbook(BytesIO(export_rms_allocation_xlsx(rms_payload())["body"]))
        workbook.remove(workbook[RMS_METADATA_SHEET])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        with self.assertRaisesRegex(RmsAllocationXlsxError, RMS_METADATA_SHEET):
            parse_rms_allocation_xlsx(output.getvalue())

    def test_parse_rejects_failure_rate_that_does_not_match_mtbf(self) -> None:
        download = export_rms_allocation_xlsx(rms_payload())
        workbook = load_workbook(BytesIO(download["body"]))
        result_sheet = workbook[RMS_RESULT_SHEET]
        result_sheet.cell(6, 4).value = 0.02
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        preview = parse_rms_allocation_xlsx(output.getvalue())

        self.assertFalse(preview["ok"])
        mismatch = next(
            error for error in preview["errors"]
            if error["code"] == "failure_rate_mtbf_mismatch"
        )
        self.assertEqual(mismatch["row"], 6)
        self.assertEqual(mismatch["field"], "failureRate")

    def test_inactive_row_allows_empty_mtbf_and_mttr(self) -> None:
        payload = rms_payload()
        payload["rows"] = [{
            **payload["rows"][0],
            "runningRatio": 0,
            "failureRate": 0,
            "mtbfHours": None,
            "mttrHours": None,
            "allocationShare": 0,
            "localAllocationShare": 0,
            "cumulativeAllocationShare": 0,
            "status": "未参与",
        }]

        preview = parse_rms_allocation_xlsx(export_rms_allocation_xlsx(payload)["body"])

        self.assertTrue(preview["ok"], preview["errors"])
        self.assertIsNone(preview["rows"][0]["mtbfHours"])
        self.assertIsNone(preview["rows"][0]["mttrHours"])

    def test_export_and_import_accept_engine_rows_without_retired_allocation_share(self) -> None:
        payload = rms_payload()
        for row in payload["rows"]:
            row.pop("allocationShare")

        download = export_rms_allocation_xlsx(payload)
        workbook = load_workbook(BytesIO(download["body"]))
        result_sheet = workbook[RMS_RESULT_SHEET]
        self.assertEqual(result_sheet.cell(6, 11).value, 0.625)

        # Workbooks created by the previous exporter have an empty legacy
        # allocationShare column. They must remain importable as well.
        result_sheet.cell(6, 11).value = None
        result_sheet.cell(7, 11).value = None
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        preview = parse_rms_allocation_xlsx(output.getvalue())

        self.assertTrue(preview["ok"], preview["errors"])
        self.assertEqual(preview["rows"][0]["allocationShare"], 0.625)
        self.assertEqual(preview["rows"][1]["allocationShare"], 0.375)

    def test_backend_preview_accepts_base64_and_maps_invalid_workbooks(self) -> None:
        connection = sqlite3.connect(":memory:")
        initialize_database(connection)
        with tempfile.TemporaryDirectory() as tmp:
            api = BackendApi(
                ContractRepository(connection),
                SimulationAdapter(repo_root=Path(__file__).resolve().parents[1]),
                output_dir=tmp,
            )
            content = export_rms_allocation_xlsx(rms_payload())["body"]
            preview = api.preview_rms_allocation_xlsx({
                "content_base64": base64.b64encode(content).decode("ascii"),
                "file_name": "RMS指标分配结果.xlsx",
            })
            self.assertTrue(preview["ok"])
            self.assertEqual(preview["fileName"], "RMS指标分配结果.xlsx")

            with self.assertRaises(BackendApiError) as invalid_base64:
                api.preview_rms_allocation_xlsx({"content_base64": "%%%"})
            self.assertEqual(invalid_base64.exception.code, "invalid_rms_allocation_xlsx")

            with self.assertRaises(BackendApiError) as invalid_xlsx:
                api.preview_rms_allocation_xlsx({
                    "content_base64": base64.b64encode(b"not-a-workbook").decode("ascii")
                })
            self.assertEqual(invalid_xlsx.exception.code, "invalid_rms_allocation_xlsx")
        connection.close()

    def test_http_preview_route_requires_auth_and_returns_canonical_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            server = create_backend_server(
                ("127.0.0.1", 0),
                repo_root=Path(__file__).resolve().parents[1],
                database_path=":memory:",
                output_dir=Path(tmp) / "artifacts",
            )
            thread = Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                base_url = f"http://127.0.0.1:{server.server_address[1]}/api"
                content_base64 = base64.b64encode(export_rms_allocation_xlsx(rms_payload())["body"]).decode("ascii")
                body = {"content_base64": content_base64, "file_name": "rms.xlsx"}
                with self.assertRaises(error.HTTPError) as unauthorized:
                    self._post_json(f"{base_url}/rms-allocation/import-xlsx/preview", body)
                self.assertEqual(unauthorized.exception.code, 401)

                login = self._post_json(
                    f"{base_url}/auth/login",
                    {"username": "user", "password": "user"},
                )
                preview = self._post_json(
                    f"{base_url}/rms-allocation/import-xlsx/preview",
                    body,
                    token=login["session"]["token"],
                )
                self.assertTrue(preview["ok"])
                self.assertEqual(preview["metadata"]["basicMissionId"], "mission-sortie")
                self.assertEqual(preview["rows"][0]["parentNodeId"], "aircraft-j15")
                self.assertEqual(preview["fileName"], "rms.xlsx")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)

    def test_http_export_is_xlsx_and_can_be_imported_through_preview_route(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            server = create_backend_server(
                ("127.0.0.1", 0),
                repo_root=Path(__file__).resolve().parents[1],
                database_path=":memory:",
                output_dir=Path(tmp) / "artifacts",
            )
            thread = Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                base_url = f"http://127.0.0.1:{server.server_address[1]}/api"
                login = self._post_json(
                    f"{base_url}/auth/login",
                    {"username": "user", "password": "user"},
                )
                content, headers = self._post_binary(
                    f"{base_url}/rms-allocation/export-xlsx",
                    rms_payload(),
                    token=login["session"]["token"],
                )
                self.assertTrue(content.startswith(b"PK\x03\x04"))
                self.assertEqual(
                    headers.get_content_type(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                self.assertIn("RMS", headers["content-disposition"])

                preview = self._post_json(
                    f"{base_url}/rms-allocation/import-xlsx/preview",
                    {
                        "content_base64": base64.b64encode(content).decode("ascii"),
                        "file_name": "RMS指标分配结果.xlsx",
                    },
                    token=login["session"]["token"],
                )
                self.assertTrue(preview["ok"], preview["errors"])
                self.assertEqual(preview["rows"][0]["nodeId"], "system-engine")
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)

    @staticmethod
    def _post_json(url: str, payload: dict, *, token: str = "") -> dict:
        headers = {"content-type": "application/json"}
        if token:
            headers["authorization"] = f"Bearer {token}"
        req = request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            method="POST",
            headers=headers,
        )
        opener = request.build_opener(request.ProxyHandler({}))
        with opener.open(req, timeout=10) as response:
            return json.loads(response.read().decode("utf-8"))

    @staticmethod
    def _post_binary(url: str, payload: dict, *, token: str = "") -> tuple[bytes, object]:
        headers = {"content-type": "application/json"}
        if token:
            headers["authorization"] = f"Bearer {token}"
        req = request.Request(
            url,
            data=json.dumps(payload).encode("utf-8"),
            method="POST",
            headers=headers,
        )
        opener = request.build_opener(request.ProxyHandler({}))
        with opener.open(req, timeout=10) as response:
            return response.read(), response.headers


if __name__ == "__main__":
    unittest.main()
