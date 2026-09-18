from __future__ import annotations
import base64
import copy
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from src.spare_mvp_backend.api import BackendApi
from src.spare_mvp_backend.project_payload import ProjectJsonExporter
from src.spare_mvp_backend.project_xlsx import parse_project_xlsx
from src.spare_mvp_backend.project_xlsx_template import export_project_xlsx, table_mapping, schema, GUIDE, FIELDS, EXTRAS
from src.spare_mvp_backend.repository import ContractRepository, initialize_database
from src.spare_mvp_backend.xlsx_text import workbook_bytes
from src.spare_mvp_contract.adapter import SimulationAdapter

ROOT = Path(__file__).resolve().parents[1]


def canonical_complete_project():
    """Deterministic reviewed four-domain fixture, through the authoritative exporter."""
    project = ProjectJsonExporter(repo_root=ROOT).export(json.loads((ROOT / 'exports/project-case-large.json').read_text()))
    project['projectInfo']['excelRoundTrip'] = {'zero': 0, 'enabled': False, 'blank': '', 'text': '中文\t制表\n换行\r回车',
        'reserved': ['@object', '~文字', '=不是公式'], 'emptyList': [], 'emptyObject': {},
        'largeInteger': 1234567890123456789, 'precision': 1.2345678901234567}
    return ProjectJsonExporter(repo_root=ROOT).export(project)


class ProjectExcelTemplateTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.project = canonical_complete_project()
        cls.content = export_project_xlsx(cls.project)

    def edit(self, callback):
        wb = load_workbook(BytesIO(self.content))
        callback(wb)
        try: return workbook_bytes(wb)
        finally: wb.close()

    def test_complete_four_domain_roundtrip_and_public_compile_gate(self):
        project, locations, issues = parse_project_xlsx(self.content)
        self.assertEqual(issues, [])
        self.assertEqual(project, self.project)
        Draft202012Validator(schema()).validate(project)
        clean = ProjectJsonExporter(repo_root=ROOT).export(project)
        self.assertEqual(clean, self.project)
        gate = SimulationAdapter(ROOT).compile_scenario_with_gate(clean, model_family='aircraft_support_v1')
        self.assertEqual(gate['status'], 'compiled')
        input_schema = json.loads((ROOT / 'contracts/aircraft_support_v1_input.schema.json').read_text())
        Draft202012Validator(input_schema).validate(gate['scenario']['simulation_inputs'])
        self.assertTrue(project['supportActivities'][0]['predecessors'])
        self.assertTrue(project['supportOrganization']['tree']['children'])
        self.assertTrue(project['missionProfile']['compositeTasks'][0]['taskItems'])
        self.assertTrue(locations['components[0].productId']['column'])

    def test_template_documents_every_schema_mapped_field_and_is_real_ooxml(self):
        with ZipFile(BytesIO(self.content)) as archive:
            self.assertIn('xl/workbook.xml', archive.namelist())
            # CR is encoded, not exposed to XML newline normalization (also without lxml).
            xml = b''.join(archive.read(n) for n in archive.namelist() if n.startswith('xl/worksheets/'))
            self.assertIn(b'&#13;', xml)
        wb = load_workbook(BytesIO(self.content), read_only=True)
        try:
            docs = {(r[0],r[1]) for r in list(wb[FIELDS].values)[1:]}
            for table in table_mapping():
                self.assertEqual(list(next(wb[table['sheet']].values)), table['headers'])
                for field in table['properties']: self.assertIn((table['sheet'],field), docs)
            self.assertLess(next(i for i,n in enumerate(wb.sheetnames) if n.startswith('任务')),next(i for i,n in enumerate(wb.sheetnames) if n.startswith('装备')))
            self.assertTrue(any(r[3] == 'array' for r in list(wb[EXTRAS].values)[1:]))
        finally: wb.close()

    def test_null_empty_missing_and_literal_markers_are_distinct(self):
        source = {'projectInfo': {'none': None, 'empty': '', 'list': [], 'dict': {}, 'value': '@null'}}
        result,_,errors=parse_project_xlsx(export_project_xlsx(source))
        self.assertEqual(errors,[])
        self.assertEqual(result,source)
        self.assertNotIn('absent',result['projectInfo'])

    def test_protocol_sheet_header_formula_and_unknown_sheet_fail_closed(self):
        changes = [
            ('unsupported_template_version',lambda wb: setattr(wb[GUIDE]['B1'],'value','future-v9')),
            ('unsupported_template_schema',lambda wb: setattr(wb[GUIDE]['B2'],'value','different')),
            ('missing_sheet',lambda wb: wb.remove(wb[EXTRAS])),
            ('invalid_header',lambda wb: setattr(wb['项目信息']['A1'],'value','丢失记录ID')),
            ('formula_not_supported',lambda wb: setattr(wb['项目信息']['F2'],'value','=1+1')),
            ('unknown_sheet',lambda wb: wb.create_sheet('未约定字段')),
        ]
        for code,change in changes:
            with self.subTest(code=code):
                _,_,errors=parse_project_xlsx(self.edit(change))
                error=next(e for e in errors if e['code']==code)
                self.assertTrue(all(k in error for k in ('sheet','row','column')))

    def test_duplicate_missing_parent_and_array_order_are_rejected(self):
        array_table=next(t for t in table_mapping() if t['path']==('components','[]'))
        for code,col,value in [('missing_parent',2,'missing-row'),('invalid_array_index',3,-1),('duplicate_row_id',1,'row-1')]:
            with self.subTest(code=code):
                content=self.edit(lambda wb: setattr(wb[array_table['sheet']].cell(2,col),'value',value))
                self.assertIn(code,{e['code'] for e in parse_project_xlsx(content)[2]})

    def test_preview_does_not_write_and_invokes_public_gate_with_clean_project(self):
        connection=sqlite3.connect(':memory:')
        initialize_database(connection)
        before=connection.total_changes
        with tempfile.TemporaryDirectory() as tmp:
            adapter=SimulationAdapter(ROOT)
            api=BackendApi(ContractRepository(connection),adapter,output_dir=tmp)
            with patch.object(adapter,'compile_scenario_with_gate',wraps=adapter.compile_scenario_with_gate) as gate:
                result=api.preview_project_xlsx({'content_base64':base64.b64encode(self.content).decode()})
            self.assertTrue(result['ok'],result['errors'])
            self.assertEqual(result['format_version'],'project-xlsx-v1')
            self.assertEqual(result['compile_status'],'compiled')
            self.assertEqual(connection.total_changes,before)
            gate.assert_called_once()
            self.assertEqual(gate.call_args.args[0],self.project)
        connection.close()

    def test_unknown_root_field_is_not_pruned_into_a_valid_project(self):
        project=copy.deepcopy(self.project)
        project['surpriseRoot']='must reject'
        connection=sqlite3.connect(':memory:');initialize_database(connection)
        with tempfile.TemporaryDirectory() as tmp:
            api=BackendApi(ContractRepository(connection),SimulationAdapter(ROOT),output_dir=tmp)
            result=api.preview_project_xlsx({'content_base64':base64.b64encode(export_project_xlsx(project)).decode()})
            self.assertFalse(result['ok'])
            self.assertTrue(any(e['code']=='project_schema_error' for e in result['errors']))
        connection.close()

    def test_oversized_text_and_nonfinite_numbers_cannot_be_silently_truncated(self):
        for value in ['x'*32768,float('nan'),float('inf')]:
            with self.subTest(value=type(value).__name__),self.assertRaises(ValueError):
                export_project_xlsx({'projectInfo':{'value':value}})


if __name__ == '__main__': unittest.main()
