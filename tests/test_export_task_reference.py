from __future__ import annotations

from datetime import datetime, timedelta, timezone
import hashlib
import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest import mock
from urllib import error, request

from openpyxl import load_workbook

from src.spare_mvp_backend.analysis_xlsx import (
    AnalysisXlsxLimitError,
    _validate_task_export_cell_budget,
    _validate_task_export_file_size,
    export_downtime_task_result_xlsx,
)
from src.spare_mvp_backend.errors import BackendApiError
from src.spare_mvp_backend.http_server import create_backend_server
from src.spare_mvp_backend.simulation_tasks import EngineRunner, SimulationTaskService


REPO_ROOT = Path(__file__).resolve().parents[1]


def downtime_result(events: list[dict] | None = None) -> dict:
    return {
        "analysis_type": "downtime_factors",
        "project_id": "project-stale-name-must-not-win",
        "analysis_source": {
            "kind": "experiment-plan",
            "projectName": "任务绑定项目",
            "projectId": "project-bound",
            "experimentPlanName": "冻结方案 A",
            "experimentPlanId": "plan-bound",
        },
        "event_details": events if events is not None else [
            {
                "sample_index": 0,
                "seed": 0,
                "tail_number": "A-01",
                "factor": "spare_shortage",
                "equipment_name": "动力装置",
                "start_minute": 0,
                "end_time_label": "仿真截止：DAY_2 00:00",
                "duration_minutes": 1440,
                "end_reason": "simulation_cutoff",
                "end_reason_label": "仿真截止",
                "status": "unresolved",
                "status_label": "未修复·仍等待备件",
                "end_state_label": "等待备件",
                "details": {"shortage_reason_label": "全网无可用供应节点"},
            },
            {
                "sample_index": 1,
                "seed": 101,
                "tail_number": "A-02",
                "factor": "failure",
                "start_minute": 60,
                "end_minute": 240,
                "duration_minutes": 180,
                "end_reason_label": "修复完成",
                "status": "completed",
                "status_label": "已修复",
                "end_state_label": "可用",
                "details": {"repair_completed_minute": 240},
            },
        ],
    }


class ImmediateEngine:
    def run(self, payload, _progress_callback):
        result = downtime_result()
        result["analysis_type"] = payload["analysis_type"]
        return result


class ResultEngine:
    def __init__(self, result: dict) -> None:
        self.result = result

    def run(self, payload, _progress_callback):
        result = json.loads(json.dumps(self.result))
        result["analysis_type"] = payload["analysis_type"]
        return result


class TaskReferenceExportTest(unittest.TestCase):
    def test_http_export_reads_owned_completed_task_and_ignores_client_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            server = create_backend_server(
                ("127.0.0.1", 0),
                repo_root=REPO_ROOT,
                database_path=":memory:",
                output_dir=Path(tmp) / "artifacts",
                simulation_engine=ImmediateEngine(),
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                base_url = f"http://127.0.0.1:{server.server_address[1]}/api"
                token = self._login(base_url, "data", "data")
                foreign_token = self._login(base_url, "user", "user")
                submitted = self._json(
                    base_url,
                    "POST",
                    "/simulation-tasks",
                    {"analysis_type": "downtime_factors"},
                    token,
                )
                self._wait_completed(base_url, submitted["task_id"], token)

                foreign_status, foreign_payload = self._json_error(
                    base_url,
                    "/analysis-results/export-xlsx",
                    {"task_id": submitted["task_id"], "analysis_type": "downtime_factors"},
                    foreign_token,
                )
                self.assertEqual((foreign_status, foreign_payload["code"]), (404, "simulation_task_not_found"))

                body, headers = self._download(
                    base_url,
                    "/analysis-results/export-xlsx",
                    {
                        "task_id": submitted["task_id"],
                        "analysis_type": "downtime_factors",
                        "export_date": "2026-09-22",
                        "filters": {"factors": ["spare_shortage"], "sample_indices": [0], "seeds": [0], "statuses": ["unresolved"]},
                        "summary": [["伪造汇总", 999]],
                        "detail_sections": [{"title": "伪造明细", "columns": ["值"], "rows": [["不得出现"]]}],
                    },
                    token,
                )
                self.assertIn("%E4%BB%BB%E5%8A%A1%E7%BB%91%E5%AE%9A%E9%A1%B9%E7%9B%AE", headers["content-disposition"])
                workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
                self.assertEqual(workbook.sheetnames, ["分析信息", "结果摘要", "因素排行", "停机事件明细"])
                info = dict(workbook["分析信息"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(info["项目名称"], "任务绑定项目")
                self.assertEqual(info["实验方案名称"], "冻结方案 A")
                self.assertEqual(workbook["结果摘要"]["B2"].value, 1)
                values = [cell.value for row in workbook["停机事件明细"].iter_rows() for cell in row]
                self.assertIn("样本 1", values)
                self.assertIn(0, values)
                self.assertIn("仿真截止：DAY_2 00:00", values)
                self.assertIn("未修复·仍等待备件", values)
                self.assertNotIn("不得出现", values)
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)

    def test_http_3218_event_task_reference_exceeds_one_mib_while_large_json_request_is_rejected(self) -> None:
        base_event = downtime_result()["event_details"][1]
        events = []
        for index in range(3_218):
            unique_text = "".join(
                hashlib.sha256(f"{index}:{part}".encode("ascii")).hexdigest()
                for part in range(16)
            )
            events.append({
                **base_event,
                "source_event_id": f"large-{index}",
                "sample_index": index % 7,
                "seed": index,
                "tail_number": f"L-{index:04d}",
                "start_minute": index,
                "description": unique_text,
            })
        with tempfile.TemporaryDirectory() as tmp:
            server = create_backend_server(
                ("127.0.0.1", 0),
                repo_root=REPO_ROOT,
                database_path=":memory:",
                output_dir=Path(tmp) / "artifacts",
                simulation_engine=ResultEngine(downtime_result(events)),
            )
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                base_url = f"http://127.0.0.1:{server.server_address[1]}/api"
                token = self._login(base_url, "data", "data")
                submitted = self._json(
                    base_url,
                    "POST",
                    "/simulation-tasks",
                    {"analysis_type": "downtime_factors"},
                    token,
                )
                self._wait_completed(base_url, submitted["task_id"], token)
                body, _headers = self._download(
                    base_url,
                    "/analysis-results/export-xlsx",
                    {
                        "task_id": submitted["task_id"],
                        "analysis_type": "downtime_factors",
                        "filters": {"factors": ["failure"], "sample_indices": [], "seeds": [], "statuses": []},
                    },
                    token,
                )
                self.assertGreater(len(body), 1024 * 1024)
                workbook = load_workbook(io.BytesIO(body), read_only=True)
                detail_rows = workbook["停机事件明细"].iter_rows(values_only=True)
                header = next(detail_rows)
                tail_column = header.index("机号")
                tails = [row[tail_column] for row in detail_rows]
                self.assertEqual(len(tails), 3_218)
                self.assertEqual((tails[0], tails[-1]), ("L-0000", "L-3217"))

                status, rejected = self._json_error(
                    base_url,
                    "/analysis-results/export-xlsx",
                    {"padding": "x" * (1024 * 1024)},
                    token,
                )
                self.assertEqual((status, rejected["code"]), (413, "request_too_large"))
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)

    def test_completed_task_export_rejects_type_mismatch_and_failed_or_expired_tasks(self) -> None:
        now = [datetime(2026, 9, 22, tzinfo=timezone.utc)]
        service = SimulationTaskService(
            EngineRunner(ImmediateEngine()),
            terminal_retention_seconds=60,
            clock=lambda: now[0],
        )
        submitted = service.submit("owner", {"analysis_type": "downtime_factors"})
        deadline = time.monotonic() + 2
        while service.status("owner", submitted["task_id"])["status"] == "running" and time.monotonic() < deadline:
            time.sleep(0.01)
        with self.assertRaises(BackendApiError) as mismatch:
            service.completed_result_for_export("owner", submitted["task_id"], analysis_type="mission_reliability")
        self.assertEqual(mismatch.exception.code, "analysis_export_type_mismatch")

        restarted_service = SimulationTaskService(EngineRunner(ImmediateEngine()))
        with self.assertRaises(BackendApiError) as restarted:
            restarted_service.completed_result_for_export(
                "owner", submitted["task_id"], analysis_type="downtime_factors"
            )
        self.assertEqual(restarted.exception.code, "simulation_task_not_found")

        class FailedEngine:
            def run(self, _payload, _progress_callback):
                raise BackendApiError("simulation_execution_failed", "仿真执行失败")

        failed_service = SimulationTaskService(EngineRunner(FailedEngine()))
        failed = failed_service.submit("owner", {"analysis_type": "downtime_factors"})
        failed_deadline = time.monotonic() + 2
        while failed_service.status("owner", failed["task_id"])["status"] == "running" and time.monotonic() < failed_deadline:
            time.sleep(0.01)
        with self.assertRaises(BackendApiError) as failed_export:
            failed_service.completed_result_for_export(
                "owner", failed["task_id"], analysis_type="downtime_factors"
            )
        self.assertEqual(failed_export.exception.code, "simulation_task_export_unavailable")

        now[0] += timedelta(seconds=61)
        with self.assertRaises(BackendApiError) as expired:
            service.completed_result_for_export("owner", submitted["task_id"], analysis_type="downtime_factors")
        self.assertEqual(expired.exception.code, "simulation_task_not_found")

    def test_large_complete_result_splits_every_ten_thousand_events(self) -> None:
        event = downtime_result()["event_details"][1]
        result = downtime_result([
            {
                **event,
                "source_event_id": f"event-{index}",
                "tail_number": f"A-{index:05d}",
                "seed": index,
                "start_minute": index,
            }
            for index in range(10_001)
        ])
        download = export_downtime_task_result_xlsx(
            result,
            {"analysis_type": "downtime_factors", "export_date": "2026-09-22"},
        )
        workbook = load_workbook(io.BytesIO(download["body"]), read_only=True)
        self.assertEqual(workbook.sheetnames[-2:], ["停机事件明细", "停机事件明细2"])
        exported_identities = []
        for sheet_name in ("停机事件明细", "停机事件明细2"):
            rows = workbook[sheet_name].iter_rows(values_only=True)
            header = next(rows)
            seed_column = header.index("随机种子")
            tail_column = header.index("机号")
            exported_identities.extend((row[seed_column], row[tail_column]) for row in rows)
        self.assertEqual(
            exported_identities,
            [(index, f"A-{index:05d}") for index in range(10_001)],
        )

    def test_cell_text_over_excel_limit_is_rejected_without_partial_file(self) -> None:
        event = downtime_result()["event_details"][0]
        event["equipment_name"] = "长" * 32_001
        with self.assertRaises(AnalysisXlsxLimitError):
            export_downtime_task_result_xlsx(
                downtime_result([event]),
                {"analysis_type": "downtime_factors"},
            )

    def test_event_count_and_aggregate_text_limits_reject_without_truncation(self) -> None:
        event = downtime_result()["event_details"][0]
        with mock.patch("src.spare_mvp_backend.analysis_xlsx.MAX_TASK_EXPORT_EVENTS", 2):
            exact = export_downtime_task_result_xlsx(
                downtime_result([event, dict(event)]),
                {"analysis_type": "downtime_factors"},
            )
            self.assertTrue(exact["body"].startswith(b"PK"))
            with self.assertRaisesRegex(AnalysisXlsxLimitError, "超过 2 条"):
                export_downtime_task_result_xlsx(
                    downtime_result([event, dict(event), dict(event)]),
                    {"analysis_type": "downtime_factors"},
                )

        with mock.patch("src.spare_mvp_backend.analysis_xlsx.MAX_TASK_EXPORT_CELL_TEXT_BYTES", 8):
            _validate_task_export_cell_budget([["12345678"]])
            with self.assertRaisesRegex(AnalysisXlsxLimitError, "单元格文本超过 64 MiB"):
                _validate_task_export_cell_budget([["123456789"]])

        with mock.patch("src.spare_mvp_backend.analysis_xlsx.MAX_DETAIL_COLUMNS", 40):
            _validate_task_export_cell_budget([[str(index) for index in range(40)]])
            with self.assertRaisesRegex(AnalysisXlsxLimitError, "超过 40 列"):
                _validate_task_export_cell_budget([[str(index) for index in range(41)]])

        with mock.patch("src.spare_mvp_backend.analysis_xlsx.MAX_TASK_EXPORT_FILE_BYTES", 8):
            _validate_task_export_file_size(b"12345678")
            with self.assertRaisesRegex(AnalysisXlsxLimitError, "Excel 文件超过 64 MiB"):
                _validate_task_export_file_size(b"123456789")

    def test_missing_sample_is_not_fabricated_and_blocked_error_code_is_preserved(self) -> None:
        event = dict(downtime_result()["event_details"][1])
        event.pop("sample_index")
        download = export_downtime_task_result_xlsx(
            downtime_result([event]),
            {"analysis_type": "downtime_factors"},
        )
        workbook = load_workbook(io.BytesIO(download["body"]), read_only=True)
        self.assertEqual(workbook["停机事件明细"]["A2"].value, "样本未记录")

        class BlockedEngine:
            def run(self, _payload, _progress_callback):
                return {"status": "blocked", "error_code": "all_samples_failed", "message": "全部样本失败"}

        service = SimulationTaskService(EngineRunner(BlockedEngine()))
        submitted = service.submit("owner", {"analysis_type": "downtime_factors"})
        deadline = time.monotonic() + 2
        while service.status("owner", submitted["task_id"])["status"] == "running" and time.monotonic() < deadline:
            time.sleep(0.01)
        result = service.result("owner", submitted["task_id"])
        self.assertEqual(result["error"]["code"], "all_samples_failed")

    @staticmethod
    def _login(base_url: str, username: str, password: str) -> str:
        return TaskReferenceExportTest._json(
            base_url, "POST", "/auth/login", {"username": username, "password": password}, None
        )["session"]["token"]

    @staticmethod
    def _wait_completed(base_url: str, task_id: str, token: str) -> None:
        deadline = time.monotonic() + 2
        while time.monotonic() < deadline:
            status = TaskReferenceExportTest._json(base_url, "GET", f"/simulation-tasks/{task_id}", None, token)
            if status["status"] == "completed":
                return
            time.sleep(0.01)
        raise AssertionError("task did not complete")

    @staticmethod
    def _json(base_url: str, method: str, path: str, payload: dict | None, token: str | None) -> dict:
        headers = {"content-type": "application/json"}
        if token:
            headers["authorization"] = f"Bearer {token}"
        req = request.Request(
            f"{base_url}{path}",
            data=None if payload is None else json.dumps(payload).encode("utf-8"),
            method=method,
            headers=headers,
        )
        with request.build_opener(request.ProxyHandler({})).open(req, timeout=5) as response:
            return json.loads(response.read().decode("utf-8"))

    @staticmethod
    def _download(base_url: str, path: str, payload: dict, token: str):
        req = request.Request(
            f"{base_url}{path}",
            data=json.dumps(payload).encode("utf-8"),
            method="POST",
            headers={"content-type": "application/json", "authorization": f"Bearer {token}"},
        )
        with request.build_opener(request.ProxyHandler({})).open(req, timeout=10) as response:
            return response.read(), response.headers

    @staticmethod
    def _json_error(base_url: str, path: str, payload: dict, token: str):
        try:
            TaskReferenceExportTest._download(base_url, path, payload, token)
        except error.HTTPError as exc:
            return exc.code, json.loads(exc.read().decode("utf-8"))
        raise AssertionError("request unexpectedly succeeded")


if __name__ == "__main__":
    unittest.main()
