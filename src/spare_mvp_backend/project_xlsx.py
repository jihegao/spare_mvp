"""Excel adapter for the canonical Project JSON contract."""

from __future__ import annotations

from io import BytesIO
import json
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


PROJECT_SHEET = "Project"
ARRAY_ROOTS = {
    "airports", "basicMissions", "missionPhases", "products", "components",
    "supportNodes", "supportResources", "transportPolicies", "supportActivityJobs",
    "supportActivities",
}
OBJECT_ROOTS = {"projectInfo", "missionProfile", "combatUnit", "equipment", "reliabilityBlockDiagram", "supportOrganization"}
SCALAR_ROOTS = {"schema_version", "project_id", "project_version", "scenarioId", "activeModule", "created_at", "updated_at"}
MAX_XLSX_BYTES = 12 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_XLSX_ZIP_ENTRIES = 512
MAX_XLSX_ROWS_PER_SHEET = 10_000
MAX_XLSX_COLUMNS_PER_SHEET = 256
MAX_XLSX_CELL_TEXT_BYTES = 100_000


class ProjectXlsxError(ValueError):
    pass


def parse_project_xlsx(content: bytes) -> tuple[dict[str, Any], dict[str, dict[str, Any]], list[dict[str, Any]]]:
    if len(content) > MAX_XLSX_BYTES:
        raise ProjectXlsxError(f"XLSX 文件超过 {MAX_XLSX_BYTES} bytes")
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ProjectXlsxError(f"无法读取 XLSX 工作簿: {exc}") from exc

    project: dict[str, Any] = {}
    locations: dict[str, dict[str, Any]] = {}
    issues: list[dict[str, Any]] = []
    known_sheets = {PROJECT_SHEET, *ARRAY_ROOTS}
    for sheet in workbook.worksheets:
        if sheet.max_row > MAX_XLSX_ROWS_PER_SHEET or sheet.max_column > MAX_XLSX_COLUMNS_PER_SHEET:
            issues.append(_issue(
                "sheet_too_large", sheet.title, 1, "",
                f"sheet 超过 {MAX_XLSX_ROWS_PER_SHEET} 行或 {MAX_XLSX_COLUMNS_PER_SHEET} 列限制",
            ))
            continue
        if sheet.title not in known_sheets:
            issues.append(_issue("unknown_sheet", sheet.title, 1, "", f"sheet 不属于当前 Project JSON 结构: {sheet.title}"))
            continue
        if sheet.title == PROJECT_SHEET:
            _parse_project_sheet(sheet, project, locations, issues)
        else:
            _parse_array_sheet(sheet, project, locations, issues)
    if PROJECT_SHEET not in workbook.sheetnames:
        issues.append(_issue("missing_project_sheet", PROJECT_SHEET, 1, "", "缺少 Project sheet"))
    workbook.close()
    return project, locations, issues


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ZIP_ENTRIES:
                raise ProjectXlsxError(f"XLSX ZIP 条目超过 {MAX_XLSX_ZIP_ENTRIES} 个")
            total_size = sum(max(0, entry.file_size) for entry in entries)
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ProjectXlsxError(f"XLSX 解压后超过 {MAX_XLSX_UNCOMPRESSED_BYTES} bytes")
    except BadZipFile as exc:
        raise ProjectXlsxError("XLSX 不是有效的 ZIP 工作簿") from exc


def _parse_project_sheet(sheet: Any, project: dict[str, Any], locations: dict[str, dict[str, Any]], issues: list[dict[str, Any]]) -> None:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or [str(value or "").strip() for value in rows[0][:2]] != ["field", "value"]:
        issues.append(_issue("invalid_header", sheet.title, 1, "field", "Project sheet 首行必须为 field、value"))
        return
    seen: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        field = str(row[0] or "").strip() if row else ""
        if not field and not any(value not in (None, "") for value in row):
            continue
        if not field:
            issues.append(_issue("missing_required", sheet.title, row_number, "field", "field 不能为空"))
            continue
        root = field.split(".", 1)[0]
        if root not in OBJECT_ROOTS | SCALAR_ROOTS:
            issues.append(_issue("unknown_field", sheet.title, row_number, field, f"字段不属于当前 Project JSON 结构: {field}"))
            continue
        if field in seen:
            issues.append(_issue("duplicate_field", sheet.title, row_number, field, f"字段重复: {field}"))
            continue
        seen.add(field)
        value = _cell_value(row[1] if len(row) > 1 else None, sheet.title, row_number, field, issues)
        _set_path(project, field, value)
        locations[field] = _location(sheet.title, row_number, field, value)


def _parse_array_sheet(sheet: Any, project: dict[str, Any], locations: dict[str, dict[str, Any]], issues: list[dict[str, Any]]) -> None:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        project[sheet.title] = []
        return
    headers = [str(value or "").strip() for value in rows[0]]
    if not any(headers):
        issues.append(_issue("invalid_header", sheet.title, 1, "", "数组 sheet 首行必须包含列名"))
        return
    duplicates = {header for header in headers if header and headers.count(header) > 1}
    for header in duplicates:
        issues.append(_issue("duplicate_column", sheet.title, 1, header, f"列名重复: {header}"))
    items: list[dict[str, Any]] = []
    if headers == ["value"]:
        scalar_items = []
        for row_number, row in enumerate(rows[1:], start=2):
            raw = row[0] if row else None
            if raw in (None, ""):
                continue
            value = _cell_value(raw, sheet.title, row_number, "value", issues)
            locations[f"{sheet.title}[{len(scalar_items)}]"] = _location(sheet.title, row_number, "value", value, 1)
            scalar_items.append(value)
        project[sheet.title] = scalar_items
        return
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue
        item: dict[str, Any] = {}
        item_index = len(items)
        for column_index, header in enumerate(headers, start=1):
            if not header or header in duplicates:
                continue
            raw = row[column_index - 1] if column_index <= len(row) else None
            if raw in (None, ""):
                continue
            value = _cell_value(raw, sheet.title, row_number, header, issues)
            _set_path(item, header, value)
            path = f"{sheet.title}[{item_index}].{header}"
            locations[path] = _location(sheet.title, row_number, header, value, column_index)
        items.append(item)
    project[sheet.title] = items


def _cell_value(value: Any, sheet: str, row: int, field: str, issues: list[dict[str, Any]]) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if len(text.encode("utf-8")) > MAX_XLSX_CELL_TEXT_BYTES:
        issues.append(_issue("cell_too_large", sheet, row, field, "单元格文本超过允许大小"))
        return None
    if text.startswith("="):
        issues.append(_issue("formula_not_supported", sheet, row, field, "不支持公式单元格，请导入已计算的值"))
        return None
    if text.lower() == "true": return True
    if text.lower() == "false": return False
    if text.lower() == "null": return None
    if text.startswith(("[", "{")):
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            issues.append(_issue("invalid_json_cell", sheet, row, field, f"JSON 单元格无法解析: {exc.msg}", text))
    return value


def _set_path(target: dict[str, Any], path: str, value: Any) -> None:
    cursor = target
    parts = path.split(".")
    for part in parts[:-1]:
        next_value = cursor.get(part)
        if not isinstance(next_value, dict):
            next_value = {}
            cursor[part] = next_value
        cursor = next_value
    cursor[parts[-1]] = value


def validate_import_relations(project: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    collections = [root for root in ARRAY_ROOTS if isinstance(project.get(root), list)]
    for root in collections:
        if root == "airports":
            continue
        identity_field = "activityCode" if root == "supportActivityJobs" else "id"
        seen: dict[str, int] = {}
        for index, item in enumerate(project[root]):
            if not isinstance(item, dict):
                issues.append(_path_issue("invalid_object", f"{root}[{index}]", "行必须解析为对象"))
                continue
            identity = str(item.get(identity_field) or "").strip()
            if not identity:
                issues.append(_path_issue("missing_required", f"{root}[{index}].{identity_field}", f"{identity_field} 不能为空"))
            elif identity in seen:
                issues.append(_path_issue("duplicate_id", f"{root}[{index}].{identity_field}", f"{identity_field} {identity} 与第 {seen[identity] + 1} 个对象重复", identity))
            elif identity:
                seen[identity] = index

    components = [item for item in project.get("components", []) if isinstance(item, dict)]
    parents = {str(item.get("id")): str(item.get("parentId")) for item in components if item.get("id") and item.get("parentId")}
    for component_id in parents:
        chain: set[str] = set()
        cursor = component_id
        while cursor in parents:
            if cursor in chain:
                index = next((i for i, item in enumerate(components) if str(item.get("id")) == component_id), 0)
                issues.append(_path_issue("circular_component_parent", f"components[{index}].parentId", f"装备父子关系成环，涉及 {cursor}", cursor))
                break
            chain.add(cursor)
            cursor = parents[cursor]

    basic_ids = {str(item.get("id")) for item in project.get("basicMissions", []) if isinstance(item, dict) and item.get("id")}
    profile = project.get("missionProfile") if isinstance(project.get("missionProfile"), dict) else {}
    composites = [item for item in profile.get("compositeTasks", []) if isinstance(item, dict)]
    composite_ids = {str(item.get("id")) for item in composites if item.get("id")}
    for ci, composite in enumerate(composites):
        for ti, task in enumerate(composite.get("taskItems") if isinstance(composite.get("taskItems"), list) else []):
            ref = str(task.get("basicMissionId") or "") if isinstance(task, dict) else ""
            if not ref or ref not in basic_ids:
                issues.append(_path_issue("missing_basic_mission_reference", f"missionProfile.compositeTasks[{ci}].taskItems[{ti}].basicMissionId", f"复合任务引用了不存在的基本任务 {ref or '<empty>'}", ref))
    for pi, periodic in enumerate(profile.get("periodicTasks") if isinstance(profile.get("periodicTasks"), list) else []):
        if not isinstance(periodic, dict): continue
        for ri, ref in enumerate(periodic.get("compositeTaskIds") if isinstance(periodic.get("compositeTaskIds"), list) else []):
            if str(ref) not in composite_ids:
                issues.append(_path_issue("missing_composite_task_reference", f"missionProfile.periodicTasks[{pi}].compositeTaskIds[{ri}]", f"周期任务引用了不存在的复合任务 {ref}", ref))
    return issues


def locate_issues(issues: list[dict[str, Any]], locations: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    result = []
    for issue in issues:
        next_issue = dict(issue)
        path = str(issue.get("path") or issue.get("field_path") or "")
        normalized = re.sub(r"\.([0-9]+)(?=\.|$)", r"[\1]", path)
        location = locations.get(normalized)
        if location is None:
            candidates = [(key, value) for key, value in locations.items() if normalized.startswith(key) or key.startswith(normalized)]
            location = max(candidates, key=lambda pair: len(pair[0]))[1] if candidates else None
        if location:
            next_issue.update({key: location[key] for key in ("sheet", "row", "column", "field") if key in location})
            if next_issue.get("reference_value") in (None, "") and "reference_value" in location:
                next_issue["reference_value"] = location["reference_value"]
        else:
            root = re.split(r"[.[]", normalized, maxsplit=1)[0]
            next_issue.setdefault("sheet", PROJECT_SHEET if root in OBJECT_ROOTS | SCALAR_ROOTS else root)
            next_issue.setdefault("field", normalized)
        result.append(next_issue)
    return result


def preview_counts(project: dict[str, Any]) -> dict[str, int]:
    return {root: len(project.get(root) or []) for root in sorted(ARRAY_ROOTS)}


def _path_issue(code: str, path: str, message: str, reference: Any = "") -> dict[str, Any]:
    return {"code": code, "path": path, "message": message, "reference_value": reference}


def _issue(code: str, sheet: str, row: int, field: str, message: str, reference: Any = "") -> dict[str, Any]:
    return {"code": code, "sheet": sheet, "row": row, "field": field, "message": message, "reference_value": reference}


def _location(sheet: str, row: int, field: str, value: Any, column: int = 2) -> dict[str, Any]:
    return {"sheet": sheet, "row": row, "column": column, "field": field, "reference_value": value}
