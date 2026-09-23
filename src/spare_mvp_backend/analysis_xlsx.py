"""Generate safe, standards-compliant XLSX exports for analysis snapshots."""

from __future__ import annotations

from src.spare_mvp_backend.xlsx_text import workbook_bytes

from datetime import datetime, timezone
import io
from itertools import chain
import re
import unicodedata
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SUPPORTED_ANALYSIS_TYPES = {
    "spare_shortfall": "备件短板分析",
    "carry_list": "飞机转场携行清单分析",
    "aircraft_mission_reliability": "飞机任务可靠性评估",
    "mission_reliability": "任务可靠度评估",
    "downtime_factors": "停机因素分析",
    "monte_carlo": "蒙特卡洛分析",
}
MAX_DETAIL_ROWS = 10_000
MAX_DETAIL_COLUMNS = 40
MAX_CELL_TEXT = 32_000
MAX_TASK_EXPORT_EVENTS = 100_000
MAX_TASK_EXPORT_CELL_TEXT_BYTES = 64 * 1024 * 1024
MAX_TASK_EXPORT_FILE_BYTES = 64 * 1024 * 1024
_UNSAFE_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
MAX_COLUMN_WIDTH = 40


class AnalysisXlsxError(ValueError):
    """A user-facing validation failure for an analysis export snapshot."""


class AnalysisXlsxLimitError(AnalysisXlsxError):
    """The complete filtered result cannot fit within the export limits."""


_DOWNTIME_FACTOR_LABELS = {
    "spare_shortage": "备件短缺",
    "failure": "装备故障",
    "equipment_shortage": "保障设备短缺",
    "preventive": "预防性维修",
}
_DOWNTIME_PHASE_LABELS = {
    "preflight": "飞行前保障",
    "flight": "任务执行",
    "mission": "任务执行",
    "repair": "修复性维修",
    "corrective": "修复性维修",
    "postflight": "飞行后保障",
    "preventive": "预防性维修",
    "waiting": "等待保障资源",
    "transport": "备件运输",
}


def export_downtime_task_result_xlsx(
    task_result: dict[str, Any],
    payload: dict[str, Any],
) -> dict[str, Any]:
    """Build a complete filtered downtime workbook from a retained task result."""
    if not isinstance(task_result, dict) or not isinstance(payload, dict):
        raise AnalysisXlsxError("停机分析任务结果格式无效，请重新运行后再导出。")
    result_type = str(task_result.get("analysis_type") or "").strip()
    requested_type = str(payload.get("analysis_type") or "").strip()
    if requested_type != "downtime_factors" or result_type != requested_type:
        raise AnalysisXlsxError("导出请求与任务的分析类型不一致。")

    raw_events = task_result.get("event_details")
    if not isinstance(raw_events, list):
        raise AnalysisXlsxError("停机分析任务没有可导出的完整事件结果，请重新运行。")
    filters = _downtime_export_filters(payload.get("filters"))
    events = [event for event in raw_events if isinstance(event, dict) and _downtime_event_matches(event, filters)]
    events.sort(key=lambda event: _downtime_sort_minute(event))
    if len(events) > MAX_TASK_EXPORT_EVENTS:
        raise AnalysisXlsxLimitError(
            f"筛选后的停机事件超过 {MAX_TASK_EXPORT_EVENTS} 条，请缩小筛选范围后重试。"
        )

    source = task_result.get("analysis_source") if isinstance(task_result.get("analysis_source"), dict) else {}
    project_name = _required_text(source.get("projectName") or task_result.get("project_name") or "当前项目", "项目名称")
    experiment_name = _clean_xml_text(source.get("experimentPlanName") or "当前项目").strip() or "当前项目"
    columns = [
        "样本", "随机种子", "机号", "停机因素类型", "装备/产品名称", "任务/阶段", "保障组织节点",
        "开始时间", "结束时间", "该段累计停机时长（小时）", "分段结束原因", "事件状态", "结束时状态", "短缺原因", "事件说明", "分类信息",
    ]
    ranking = _downtime_ranking(events, filters["factors"])
    total_events = len(events)
    total_hours = sum(row["downtime_hours"] for row in ranking)
    primary_factor = next((row["label"] for row in ranking if row["downtime_hours"] > 0), "--")
    exported_at = payload.get("export_date") or payload.get("exported_at")

    info_columns = ["信息项", "内容"]
    info_rows = [
        ["项目名称", project_name],
        ["分析类型", SUPPORTED_ANALYSIS_TYPES["downtime_factors"]],
        ["项目 ID", source.get("projectId") or task_result.get("project_id") or ""],
        ["实验方案名称", experiment_name],
        ["实验方案 ID", source.get("experimentPlanId") or ""],
        ["结果来源", "仿真任务保留结果"],
        ["筛选说明", _downtime_filter_description(filters)],
    ]
    summary_columns = ["指标名称", "值", "单位"]
    summary_rows = [
        ["停机事件次数", total_events, "次"],
        ["累计停机时长", round(total_hours, 6), "小时"],
        ["首要停机因素", primary_factor, ""],
    ]
    rank_columns = ["排序", "停机因素", "事件次数", "累计停机时长（小时）", "当前范围时长占比"]
    rank_rows = [
        [index, row["label"], row["event_count"], round(row["downtime_hours"], 6), row["downtime_hours"] / total_hours if total_hours else 0]
        for index, row in enumerate(ranking, start=1)
    ]
    detail_sheet_count = max(1, (len(events) + MAX_DETAIL_ROWS - 1) // MAX_DETAIL_ROWS)
    fixed_rows = [
        info_columns, *info_rows,
        summary_columns, *summary_rows,
        rank_columns, *rank_rows,
        *([columns] * detail_sheet_count),
    ]
    _validate_task_export_cell_budget(
        chain(fixed_rows, (_downtime_event_export_row(event) for event in events))
    )

    workbook = Workbook(write_only=True)
    info_sheet = workbook.create_sheet("分析信息")
    _write_only_table(info_sheet, info_columns, info_rows)
    summary_sheet = workbook.create_sheet("结果摘要")
    _write_only_table(summary_sheet, summary_columns, summary_rows)
    rank_sheet = workbook.create_sheet("因素排行")
    _write_only_table(rank_sheet, rank_columns, rank_rows)
    if events:
        for start in range(0, len(events), MAX_DETAIL_ROWS):
            sheet_index = start // MAX_DETAIL_ROWS + 1
            sheet = workbook.create_sheet("停机事件明细" if sheet_index == 1 else f"停机事件明细{sheet_index}")
            _write_only_table(
                sheet,
                columns,
                [_downtime_event_export_row(event) for event in events[start:start + MAX_DETAIL_ROWS]],
            )
    else:
        detail_sheet = workbook.create_sheet("停机事件明细")
        _write_only_table(detail_sheet, columns, [])
    workbook.properties.title = SUPPORTED_ANALYSIS_TYPES["downtime_factors"]
    workbook.properties.creator = "备件规划及任务可靠度验证评估平台"
    try:
        body = _write_only_workbook_bytes(workbook)
    except Exception as exc:
        raise AnalysisXlsxError("分析结果包含无法写入 Excel 的字符或数值，请检查后重试。") from exc
    _validate_task_export_file_size(body)
    return {
        "body": body,
        "content_type": XLSX_CONTENT_TYPE,
        "filename": analysis_export_filename(project_name, SUPPORTED_ANALYSIS_TYPES["downtime_factors"], exported_at),
    }


def export_analysis_snapshot_xlsx(payload: dict[str, Any]) -> dict[str, Any]:
    """Return an OOXML workbook built only from the supplied visible snapshot."""
    if not isinstance(payload, dict):
        raise AnalysisXlsxError("导出请求格式无效，请刷新页面后重试。")
    analysis_type = str(payload.get("analysis_type") or "").strip()
    if analysis_type not in SUPPORTED_ANALYSIS_TYPES:
        raise AnalysisXlsxError("不支持当前分析类型的 Excel 导出。")

    project_name = _required_text(payload.get("project_name"), "项目名称")
    analysis_name = _required_text(
        payload.get("analysis_name") or SUPPORTED_ANALYSIS_TYPES[analysis_type],
        "分析类型",
    )
    information = _pair_rows(payload.get("analysis_information"), "分析信息")
    summary = _summary_rows(payload.get("summary"))
    sections, omitted_rows = _detail_sections(payload.get("detail_sections"))
    if not summary and not any(section["rows"] for section in sections):
        raise AnalysisXlsxError("当前页面没有可导出的分析结果。")

    try:
        workbook = Workbook()
        info_sheet = workbook.active
        info_sheet.title = "分析信息"
        _write_table(
            info_sheet,
            ["信息项", "内容"],
            [
                ["项目名称", project_name],
                ["分析类型", analysis_name],
                *information,
                *([["导出说明", f"明细超过 {MAX_DETAIL_ROWS} 行，已截断 {omitted_rows} 行。"]] if omitted_rows else []),
            ],
        )

        summary_sheet = workbook.create_sheet("结果摘要")
        _write_table(summary_sheet, ["指标名称", "值", "单位"], summary)

        detail_sheet = workbook.create_sheet("结果明细")
        if not sections:
            _write_table(detail_sheet, ["结果"], [["当前页面无明细表，仅导出结果摘要。"]])
        else:
            row_index = 1
            for section_index, section in enumerate(sections):
                if section_index:
                    row_index += 1
                detail_sheet.cell(row=row_index, column=1, value=_safe_cell_value(section["title"]))
                detail_sheet.cell(row=row_index, column=1).font = Font(bold=True, size=12)
                row_index += 1
                row_index = _write_table(
                    detail_sheet,
                    section["columns"],
                    section["rows"],
                    start_row=row_index,
                )
            _fit_columns(detail_sheet)

        workbook.properties.title = analysis_name
        workbook.properties.creator = "备件规划及任务可靠度验证评估平台"
        output = io.BytesIO()
        output.write(workbook_bytes(workbook))
    except Exception as exc:
        raise AnalysisXlsxError("分析结果包含无法写入 Excel 的字符或数值，请检查后重试。") from exc
    filename = (
        monte_carlo_export_filename(
            project_name,
            payload.get("experiment_name") or "当前项目",
            payload.get("export_date") or payload.get("exported_at"),
        )
        if analysis_type == "monte_carlo"
        else analysis_export_filename(project_name, analysis_name, payload.get("exported_at"))
    )
    return {
        "body": output.getvalue(),
        "content_type": XLSX_CONTENT_TYPE,
        "filename": filename,
    }


def analysis_export_filename(project_name: str, analysis_name: str, exported_at: Any = None) -> str:
    timestamp = _export_timestamp(exported_at).strftime("%Y%m%d-%H%M%S")
    project_segment = _safe_filename_segment(project_name, "项目")
    analysis_segment = _safe_filename_segment(analysis_name, "分析结果")
    return f"{project_segment}-{analysis_segment}-{timestamp}.xlsx"


def monte_carlo_export_filename(project_name: str, experiment_name: Any, export_date: Any = None) -> str:
    date_stamp = _export_timestamp(export_date).strftime("%Y%m%d")
    project_segment = _safe_filename_segment(project_name, "项目")
    experiment_segment = _safe_filename_segment(experiment_name, "当前项目")
    return f"{project_segment}-{experiment_segment}-蒙特卡洛分析-{date_stamp}.xlsx"


def _downtime_export_filters(value: Any) -> dict[str, set[Any] | None]:
    if value is None:
        value = {}
    if not isinstance(value, dict):
        raise AnalysisXlsxError("停机事件筛选条件格式无效。")

    def string_set(key: str) -> set[str] | None:
        if key not in value:
            return None
        raw = value.get(key)
        if not isinstance(raw, list):
            raise AnalysisXlsxError(f"筛选条件 {key} 必须为数组。")
        return {str(item).strip() for item in raw if str(item).strip()}

    factors = string_set("factors")
    if factors is not None and not factors.issubset(_DOWNTIME_FACTOR_LABELS):
        raise AnalysisXlsxError("停机因素筛选包含不支持的值。")
    sample_indices: set[int] | None = None
    if "sample_indices" in value:
        raw_samples = value.get("sample_indices")
        if not isinstance(raw_samples, list):
            raise AnalysisXlsxError("筛选条件 sample_indices 必须为数组。")
        sample_indices = set()
        for item in raw_samples:
            if isinstance(item, bool):
                raise AnalysisXlsxError("样本筛选必须使用非负整数。")
            try:
                parsed = int(item)
            except (TypeError, ValueError) as exc:
                raise AnalysisXlsxError("样本筛选必须使用非负整数。") from exc
            if parsed < 0 or str(item).strip() not in {str(parsed), f"{parsed}.0"}:
                raise AnalysisXlsxError("样本筛选必须使用非负整数。")
            sample_indices.add(parsed)
    seeds = string_set("seeds")
    statuses = string_set("statuses")
    if sample_indices == set():
        sample_indices = None
    if seeds == set():
        seeds = None
    if statuses == set():
        statuses = None
    return {
        "factors": factors,
        "sample_indices": sample_indices,
        "seeds": seeds,
        "statuses": statuses,
    }


def _downtime_event_matches(event: dict[str, Any], filters: dict[str, set[Any] | None]) -> bool:
    factor = _first_text(event.get("factor"), event.get("event_type"), event.get("eventType"), event.get("reason"))
    sample_index = _event_sample_index(event)
    seed = _first_text(event.get("seed"), event.get("random_seed"), event.get("randomSeed"))
    status_values = {_downtime_event_status(event), _first_text(event.get("status_label"), event.get("statusLabel"))}
    return (
        (filters["factors"] is None or factor in filters["factors"])
        and (filters["sample_indices"] is None or sample_index in filters["sample_indices"])
        and (filters["seeds"] is None or seed in filters["seeds"])
        and (filters["statuses"] is None or bool(status_values & filters["statuses"]))
    )


def _downtime_ranking(events: list[dict[str, Any]], selected_factors: set[Any] | None) -> list[dict[str, Any]]:
    factor_order = list(_DOWNTIME_FACTOR_LABELS)
    included = set(factor_order) if selected_factors is None else set(selected_factors)
    rows = []
    for factor in factor_order:
        if factor not in included:
            continue
        factor_events = [event for event in events if _first_text(event.get("factor"), event.get("event_type"), event.get("reason")) == factor]
        rows.append({
            "factor": factor,
            "label": _DOWNTIME_FACTOR_LABELS[factor],
            "event_count": len(factor_events),
            "downtime_hours": sum(_downtime_duration_hours(event) for event in factor_events),
        })
    return sorted(rows, key=lambda row: (-row["downtime_hours"], factor_order.index(row["factor"])))


def _downtime_event_export_row(event: dict[str, Any]) -> list[Any]:
    factor = _first_text(event.get("factor"), event.get("event_type"), event.get("eventType"), event.get("reason"))
    details = event.get("details") if isinstance(event.get("details"), dict) else {}
    start_value = _first_present(event, "start_time_label", "startTimeLabel", "start_minute", "startMinute", "start_time", "startTime")
    end_value = _first_present(event, "end_time_label", "endTimeLabel", "end_minute", "endMinute", "end_time", "endTime")
    tail_number = _first_text(event.get("tail_number"), event.get("tailNumber"))
    status = _first_text(event.get("status_label"), event.get("statusLabel"), _downtime_event_status(event))
    seed = _first_present(event, "seed", "random_seed", "randomSeed")
    sample_index = _event_sample_index(event)
    return [
        f"样本 {sample_index + 1}" if sample_index is not None else "样本未记录",
        seed if seed is not None else "未记录",
        tail_number or "未记录机号",
        _DOWNTIME_FACTOR_LABELS.get(factor, "其他停机因素"),
        _first_text(event.get("equipment_name"), event.get("equipmentName"), event.get("aircraft_type"), event.get("aircraftType"), tail_number) or "暂无数据",
        _downtime_task_phase(event),
        _first_text(event.get("support_node_name"), event.get("supportNodeName")) or "未记录保障组织",
        _downtime_time_label(start_value),
        _downtime_end_time_label(event, end_value),
        round(_downtime_duration_hours(event), 6),
        _first_text(event.get("end_reason_label"), event.get("endReasonLabel"), event.get("segment_end_reason"), event.get("segmentEndReason"), event.get("end_reason"), event.get("endReason")) or "未记录",
        status or "未记录",
        _first_text(event.get("end_state_label"), event.get("endStateLabel"), status) or "未记录",
        _first_text(event.get("shortage_reason_label"), event.get("shortageReasonLabel"), event.get("shortage_reason"), event.get("shortageReason"), details.get("shortage_reason_label"), details.get("shortage_reason"), details.get("reason")) or "--",
        _downtime_event_description(event, factor, tail_number),
        "；".join(f"{label}：{value}" for label, value in _downtime_specific_details(factor, details)),
    ]


def _event_sample_index(event: dict[str, Any]) -> int | None:
    value = _first_present(event, "sample_index", "sampleIndex")
    if isinstance(value, bool) or value is None:
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return None
    if parsed < 0 or str(value).strip() not in {str(parsed), f"{parsed}.0"}:
        return None
    return parsed


def _downtime_sort_minute(event: dict[str, Any]) -> float:
    value = _first_present(event, "start_minute", "startMinute", "start_time", "startTime")
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _downtime_event_status(event: dict[str, Any]) -> str:
    return _first_text(
        event.get("end_status"), event.get("endStatus"), event.get("segment_status"),
        event.get("segmentStatus"), event.get("status"), event.get("status_label"),
        event.get("statusLabel"),
    )


def _downtime_duration_hours(event: dict[str, Any]) -> float:
    for key in ("duration_hours", "durationHours"):
        try:
            value = float(event.get(key))
            if value >= 0:
                return value
        except (TypeError, ValueError):
            pass
    for key in ("duration_minutes", "durationMinutes"):
        try:
            value = float(event.get(key))
            if value >= 0:
                return value / 60.0
        except (TypeError, ValueError):
            pass
    return 0.0


def _downtime_task_phase(event: dict[str, Any]) -> str:
    mission = _first_text(event.get("mission_name"), event.get("missionName"), event.get("task_label"), event.get("taskLabel"))
    mission_id = _first_text(event.get("mission_id"), event.get("missionId"), event.get("task_id"), event.get("taskId"))
    phase = _first_text(event.get("mission_phase_name"), event.get("missionPhaseName"), event.get("phase_name"), event.get("phaseName"), event.get("task_name"), event.get("taskName"))
    details = event.get("details") if isinstance(event.get("details"), dict) else {}
    if details.get("maintenance_method") == "non_replacement" and re.search(r"(?:LRU|换件|更换)", phase, re.IGNORECASE):
        phase = "原位维修"
    elif phase:
        phase = _DOWNTIME_PHASE_LABELS.get(phase.lower(), phase if re.search(r"[\u3400-\u9fff]", phase) else "其他阶段")
    if not mission and not mission_id:
        return f"任务外事件；阶段：{phase}" if phase else "不在任务阶段"
    task_label = mission or "任务名称未解析"
    return f"{task_label}；阶段：{phase or '未记录'}"


def _downtime_specific_details(factor: str, details: dict[str, Any]) -> list[tuple[str, str]]:
    def value(raw: Any, fallback: str = "暂无数据") -> str:
        text = _first_text(raw)
        return text or fallback

    if factor == "spare_shortage":
        return [
            ("备件", value(details.get("spare_name") or details.get("spare_type"))),
            ("需求", value(details.get("required_quantity"))),
            ("可用", value(details.get("available_quantity"))),
            ("短缺", value(details.get("shortage_quantity"))),
            ("短缺原因", value(details.get("shortage_reason_label") or details.get("shortage_reason_code"))),
            ("到货/等待结束", _downtime_time_label(details.get("arrival_minute") if details.get("arrival_minute") is not None else details.get("wait_end_minute"))),
        ]
    if factor == "failure":
        method = _first_text(details.get("maintenance_method_label"))
        if not method:
            if details.get("maintenance_method") == "replacement":
                method = "换件维修"
            elif details.get("maintenance_method") == "non_replacement":
                method = "原位维修"
            else:
                method = "换件维修" if details.get("requires_spare") is True else "维修方式未记录"
        return [
            ("故障部件", value(details.get("component_name"))),
            ("故障发生", _downtime_time_label(details.get("failure_minute"))),
            ("维修方式", method),
            ("修复完成", _downtime_time_label(details.get("repair_completed_minute"))),
        ]
    if factor == "equipment_shortage":
        wait_minutes = details.get("wait_minutes")
        try:
            wait_label = f"{float(wait_minutes):g} 分钟" if float(wait_minutes) >= 0 else "暂无数据"
        except (TypeError, ValueError):
            wait_label = "暂无数据"
        return [
            ("保障设备", value(details.get("equipment_name") or details.get("equipment_model"))),
            ("需求", value(details.get("required_quantity"))),
            ("可用", value(details.get("available_quantity"))),
            ("短缺", value(details.get("shortage_quantity"))),
            ("等待时长", wait_label),
        ]
    if factor == "preventive":
        return [
            ("维修项目", value(details.get("maintenance_item") or details.get("maintenance_type"))),
            ("触发条件", _preventive_trigger_label(details.get("trigger_condition") or details.get("trigger_type"))),
            ("计划开始", _downtime_time_label(details.get("planned_start_minute"))),
            ("实际完成", _downtime_time_label(details.get("completed_minute"))),
        ]
    return [("事件信息", "未识别的停机事件")]


def _preventive_trigger_label(raw: Any) -> str:
    triggers = [item.strip() for item in str(raw or "").split(",") if item.strip()]
    if not triggers:
        return "未记录触发条件"
    labels = {"calendar_time": "日历时间到期", "flight_hours": "飞行小时到期", "landings": "起落次数到期"}
    return "、".join(labels.get(item.lower(), item if re.search(r"[\u3400-\u9fff]", item) else "其他触发条件") for item in triggers)


def _downtime_event_description(event: dict[str, Any], factor: str, tail_number: str) -> str:
    explicit = _first_text(event.get("description"), event.get("event_description"), event.get("eventDescription"))
    if explicit:
        return explicit
    subject = f"飞机{tail_number}" if tail_number else "当前装备"
    descriptions = {
        "failure": "装备发生故障，当前不可用并等待修复",
        "equipment_shortage": "保障设备不足，当前作业正在等待资源",
        "spare_shortage": "所需备件短缺，当前作业正在等待补给",
        "preventive": "装备正在执行预防性维修",
    }
    return subject + descriptions.get(factor, "记录到未识别的停机事件")


def _downtime_time_label(value: Any) -> str:
    if value is None or value == "":
        return "暂无时间"
    if isinstance(value, str) and (value.startswith("DAY_") or "仿真截止" in value):
        return _clean_xml_text(value).strip()
    try:
        minute = round(float(value))
    except (TypeError, ValueError):
        return _clean_xml_text(value).strip() or "暂无时间"
    if minute < 0:
        return "暂无时间"
    day = minute // 1440 + 1
    minute_of_day = minute % 1440
    return f"DAY_{day} {minute_of_day // 60:02d}:{minute_of_day % 60:02d}"


def _downtime_end_time_label(event: dict[str, Any], value: Any) -> str:
    label = _downtime_time_label(value)
    end_reason = _first_text(event.get("end_reason"), event.get("endReason"))
    if end_reason == "simulation_cutoff" and "仿真截止" not in label:
        return f"仿真截止：{label}"
    return label


def _downtime_filter_description(filters: dict[str, set[Any] | None]) -> str:
    labels = []
    for key, label in (("factors", "因素"), ("sample_indices", "样本"), ("seeds", "种子"), ("statuses", "状态")):
        values = filters[key]
        if values is not None:
            labels.append(f"{label}：{', '.join(str(value) for value in sorted(values, key=str)) or '无'}")
    return "；".join(labels) or "全部事件"


def _validate_task_export_cell_budget(rows: Any) -> None:
    total_bytes = 0
    for row in rows:
        if len(row) > MAX_DETAIL_COLUMNS:
            raise AnalysisXlsxLimitError(f"导出列数超过 {MAX_DETAIL_COLUMNS} 列限制。")
        for value in row:
            text = "" if value is None else _ILLEGAL_XML_CHARACTERS.sub("", str(value))
            if len(text) > MAX_CELL_TEXT:
                raise AnalysisXlsxLimitError("停机事件包含超过 Excel 单元格长度限制的文本，请缩小或清理结果后重试。")
            total_bytes += len(text.encode("utf-8"))
            if total_bytes > MAX_TASK_EXPORT_CELL_TEXT_BYTES:
                raise AnalysisXlsxLimitError("筛选结果的单元格文本超过 64 MiB，请缩小筛选范围后重试。")


def _validate_task_export_file_size(body: bytes) -> None:
    if len(body) > MAX_TASK_EXPORT_FILE_BYTES:
        raise AnalysisXlsxLimitError("生成的 Excel 文件超过 64 MiB，请缩小筛选范围后重试。")


def _write_only_table(sheet: Any, columns: list[Any], rows: list[list[Any]]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.freeze_panes = "A2"
    header = []
    for value in columns:
        cell = WriteOnlyCell(sheet, value=_safe_cell_value(value))
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(cell)
    sheet.append(header)
    for row in rows:
        sheet.append([_safe_cell_value(value) for value in row])


def _write_only_workbook_bytes(workbook: Workbook) -> bytes:
    raw = io.BytesIO()
    workbook.save(raw)
    output = io.BytesIO()
    with ZipFile(io.BytesIO(raw.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename.endswith(".xml"):
                content = content.replace(b"\r", b"&#13;")
            target.writestr(entry, content)
    return output.getvalue()


def _first_text(*values: Any) -> str:
    for value in values:
        if value is None:
            continue
        text = _ILLEGAL_XML_CHARACTERS.sub("", str(value)).strip()
        if text:
            return text
    return ""


def _first_present(mapping: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in mapping and mapping[key] is not None and mapping[key] != "":
            return mapping[key]
    return None


def _pair_rows(value: Any, label: str) -> list[list[Any]]:
    if value is None:
        return []
    if not isinstance(value, list):
        raise AnalysisXlsxError(f"{label}格式无效。")
    rows: list[list[Any]] = []
    for row in value:
        if not isinstance(row, list) or len(row) != 2:
            raise AnalysisXlsxError(f"{label}格式无效。")
        row_label = _safe_cell_value(row[0])
        row_value = _safe_cell_value(row[1])
        if row_label == "parallelCoresError":
            if str(row_value).strip():
                rows.append(["导出说明", f"并行核心数配置异常：{row_value}"])
            continue
        rows.append([row_label, row_value])
    return rows


def _summary_rows(value: Any) -> list[list[Any]]:
    if value is None:
        return []
    if not isinstance(value, list):
        raise AnalysisXlsxError("结果摘要格式无效。")
    rows: list[list[Any]] = []
    for row in value:
        if not isinstance(row, list) or not 2 <= len(row) <= 3:
            raise AnalysisXlsxError("结果摘要格式无效。")
        rows.append([
            _safe_cell_value(row[0]),
            _safe_cell_value(row[1]),
            _safe_cell_value(row[2] if len(row) > 2 else ""),
        ])
    return rows


def _detail_sections(value: Any) -> tuple[list[dict[str, Any]], int]:
    if value is None:
        return [], 0
    if not isinstance(value, list):
        raise AnalysisXlsxError("结果明细格式无效。")
    sections: list[dict[str, Any]] = []
    remaining_rows = MAX_DETAIL_ROWS
    omitted_rows = 0
    for raw_section in value:
        if not isinstance(raw_section, dict):
            raise AnalysisXlsxError("结果明细格式无效。")
        title = _required_text(raw_section.get("title"), "明细标题")
        columns = raw_section.get("columns")
        rows = raw_section.get("rows")
        if not isinstance(columns, list) or not columns or len(columns) > MAX_DETAIL_COLUMNS:
            raise AnalysisXlsxError(f"“{title}”的列定义无效或超过 {MAX_DETAIL_COLUMNS} 列限制。")
        if not isinstance(rows, list):
            raise AnalysisXlsxError(f"“{title}”的明细格式无效。")
        safe_columns = [_safe_cell_value(column) for column in columns]
        safe_rows: list[list[Any]] = []
        accepted = min(len(rows), remaining_rows)
        for row in rows[:accepted]:
            if not isinstance(row, list) or len(row) != len(columns):
                raise AnalysisXlsxError(f"“{title}”的明细列数与表头不一致。")
            safe_rows.append([_safe_cell_value(cell) for cell in row])
        omitted_rows += len(rows) - accepted
        remaining_rows -= accepted
        sections.append({"title": title, "columns": safe_columns, "rows": safe_rows})
    return sections, omitted_rows


def _write_table(sheet, columns: list[Any], rows: list[list[Any]], *, start_row: int = 1) -> int:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for column_index, value in enumerate(columns, start=1):
        cell = sheet.cell(row=start_row, column=column_index, value=_safe_cell_value(value))
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[start_row].height = max(sheet.row_dimensions[start_row].height or 0, 30)
    for row_index, row in enumerate(rows, start=start_row + 1):
        needs_extra_height = False
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_safe_cell_value(value))
            visual_width = _text_display_width(cell.value)
            should_wrap = "\n" in str(cell.value or "") or visual_width > MAX_COLUMN_WIDTH
            cell.alignment = Alignment(vertical="top", wrap_text=should_wrap)
            needs_extra_height = needs_extra_height or should_wrap
        if needs_extra_height:
            sheet.row_dimensions[row_index].height = 45
    sheet.freeze_panes = sheet.freeze_panes or f"A{start_row + 1}"
    _fit_columns(sheet)
    return start_row + len(rows) + 1


def _fit_columns(sheet) -> None:
    for column_index in range(1, sheet.max_column + 1):
        width = max(
            (_text_display_width(sheet.cell(row=row_index, column=column_index).value) for row_index in range(1, sheet.max_row + 1)),
            default=8,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(MAX_COLUMN_WIDTH, max(12, width + 2))


def _text_display_width(value: Any) -> int:
    """Approximate Excel's visible width, counting full-width CJK glyphs twice."""
    lines = str(value or "").expandtabs(4).splitlines() or [""]
    return max(
        sum(
            0
            if unicodedata.combining(character)
            else 2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
            for character in line
        )
        for line in lines
    )


def _safe_cell_value(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return value
    text = _clean_xml_text(value).strip()[:MAX_CELL_TEXT]
    if text.startswith(_FORMULA_PREFIXES):
        return f"'{text}"
    return text


def _required_text(value: Any, label: str) -> str:
    text = _clean_xml_text(value).strip()
    if not text:
        raise AnalysisXlsxError(f"缺少{label}，无法导出。")
    return text


def _safe_filename_segment(value: Any, fallback: str) -> str:
    text = _UNSAFE_FILENAME.sub("_", _clean_xml_text(value)).strip(" .")
    text = re.sub(r"\s+", " ", text)[:60].strip(" .")
    return text or fallback


def _clean_xml_text(value: Any) -> str:
    return _ILLEGAL_XML_CHARACTERS.sub("", str(value or ""))


def _export_timestamp(value: Any) -> datetime:
    text = str(value or "").strip()
    if text:
        try:
            if re.fullmatch(r"\d{8}", text):
                return datetime.strptime(text, "%Y%m%d").replace(tzinfo=timezone.utc)
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            pass
    return datetime.now(timezone.utc)
