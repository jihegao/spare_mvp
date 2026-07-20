"""Round-trip XLSX adapter for editable RMS allocation result snapshots."""

from __future__ import annotations

from io import BytesIO
import math
import re
from typing import Any

from openpyxl import Workbook, load_workbook

from src.spare_mvp_backend.project_xlsx import MAX_XLSX_BYTES, validate_xlsx_archive


RMS_XLSX_SCHEMA_VERSION = "rms-allocation-xlsx-v1"
RMS_RESULT_SHEET = "RMS分配结果"
RMS_METADATA_SHEET = "RMS元数据"

METADATA_FIELDS = (
    ("schemaVersion", "格式版本", ("schema_version",)),
    ("projectId", "项目ID", ("project_id",)),
    ("projectName", "项目名称", ("project_name",)),
    ("aircraftModel", "飞机型号", ("aircraft_model",)),
    ("basicMissionId", "基本任务ID", ("basic_mission_id",)),
    ("basicMissionName", "基本任务名称", ("basic_mission_name",)),
    ("missionHours", "任务时长(h)", ("mission_hours",)),
    ("missionReliability", "整机任务可靠度", ("mission_reliability",)),
    ("mttrHours", "整机MTTR(h)", ("mttr_hours",)),
    ("planId", "方案ID", ("plan_id",)),
    ("planVersion", "方案版本", ("plan_version",)),
    ("algorithmVersion", "算法版本", ("algorithm_version",)),
    ("method", "计算方法", ()),
    ("generatedAt", "生成时间", ("generated_at",)),
)

RESULT_COLUMNS = (
    ("level", "层级"),
    ("nodeName", "节点"),
    ("runningRatio", "运行比"),
    ("failureRate", "失效率"),
    ("mtbfHours", "MTBF(h)"),
    ("mttrHours", "MTTR(h)"),
    ("nodeId", "节点ID"),
    ("parentNodeId", "父节点ID"),
    ("model", "型号"),
    ("installationCount", "安装数"),
    ("allocationShare", "分配份额"),
    ("localAllocationShare", "本级分配份额"),
    ("cumulativeAllocationShare", "累计分配份额"),
    ("cumulativeInstallationCount", "累计安装数"),
    ("status", "状态"),
    ("verificationStatus", "校验状态"),
)

_REQUIRED_METADATA = (
    "projectId", "aircraftModel", "basicMissionId", "missionHours",
    "missionReliability", "mttrHours", "planId", "planVersion",
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class RmsAllocationXlsxError(ValueError):
    """Raised when an RMS workbook cannot be read safely."""


def export_rms_allocation_xlsx(payload: dict[str, Any]) -> dict[str, Any]:
    metadata = _metadata_from_payload(payload)
    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = RMS_RESULT_SHEET

    # Keep the legacy three-line summary and row-5 result header stable.
    result_sheet.append(["项目", _safe_cell_value(metadata.get("projectName"))])
    result_sheet.append(["计算方法", _safe_cell_value(metadata.get("method"))])
    result_sheet.append(["生成时间", _safe_cell_value(metadata.get("generatedAt"))])
    result_sheet.append([])
    result_sheet.append([label for _, label in RESULT_COLUMNS])
    for source_row in payload.get("rows") or []:
        row = source_row if isinstance(source_row, dict) else {}
        result_sheet.append([
            _result_cell_value(field, _result_source_value(row, field))
            for field, _ in RESULT_COLUMNS
        ])

    metadata_sheet = workbook.create_sheet(RMS_METADATA_SHEET)
    metadata_sheet.append(["field", "label", "value"])
    for field, label, _ in METADATA_FIELDS:
        metadata_sheet.append([field, label, _safe_cell_value(metadata.get(field))])

    output = BytesIO()
    try:
        workbook.save(output)
    except Exception as exc:
        raise RmsAllocationXlsxError(f"无法写入 RMS Excel: {exc}") from exc
    finally:
        workbook.close()
    return {
        "body": output.getvalue(),
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "filename": "rms-allocation-result.xlsx",
    }


def parse_rms_allocation_xlsx(content: bytes) -> dict[str, Any]:
    if len(content) > MAX_XLSX_BYTES:
        raise RmsAllocationXlsxError(f"XLSX 文件超过 {MAX_XLSX_BYTES} bytes")
    try:
        validate_xlsx_archive(content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except RmsAllocationXlsxError:
        raise
    except Exception as exc:
        raise RmsAllocationXlsxError(f"无法读取 RMS XLSX 工作簿: {exc}") from exc

    try:
        if RMS_METADATA_SHEET not in workbook.sheetnames:
            raise RmsAllocationXlsxError(f"缺少 {RMS_METADATA_SHEET} sheet")
        if RMS_RESULT_SHEET not in workbook.sheetnames:
            raise RmsAllocationXlsxError(f"缺少 {RMS_RESULT_SHEET} sheet")
        metadata, issues = _parse_metadata_sheet(workbook[RMS_METADATA_SHEET])
        rows, row_issues = _parse_result_sheet(workbook[RMS_RESULT_SHEET])
        issues.extend(row_issues)
        issues.extend(_validate_metadata(metadata))
        issues.extend(_validate_result_rows(rows))
        return {
            "ok": not issues,
            "schemaVersion": RMS_XLSX_SCHEMA_VERSION,
            "metadata": metadata,
            "rows": rows,
            "errors": issues,
        }
    finally:
        workbook.close()


def _metadata_from_payload(payload: dict[str, Any]) -> dict[str, Any]:
    metadata_source = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
    metadata: dict[str, Any] = {}
    for field, _, aliases in METADATA_FIELDS:
        value = _first_present(metadata_source, (field, *aliases))
        if value is None:
            value = _first_present(payload, (field, *aliases))
        metadata[field] = value
    # The workbook format version is server-owned and must not be confused with
    # an allocation plan or result envelope schema supplied by the caller.
    metadata["schemaVersion"] = RMS_XLSX_SCHEMA_VERSION
    return metadata


def _parse_metadata_sheet(sheet: Any) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    rows = list(sheet.iter_rows(values_only=False))
    if not rows or [str(cell.value or "").strip() for cell in rows[0][:3]] != ["field", "label", "value"]:
        raise RmsAllocationXlsxError(f"{RMS_METADATA_SHEET} sheet 首行必须为 field、label、value")
    supported = {field for field, _, _ in METADATA_FIELDS}
    metadata: dict[str, Any] = {}
    issues: list[dict[str, Any]] = []
    for row_number, cells in enumerate(rows[1:], start=2):
        field = str(cells[0].value or "").strip() if cells else ""
        if not field and not any(cell.value not in (None, "") for cell in cells):
            continue
        if field not in supported:
            issues.append(_issue("unknown_metadata_field", RMS_METADATA_SHEET, row_number, field, f"不支持的元数据字段: {field or '<empty>'}"))
            continue
        if field in metadata:
            issues.append(_issue("duplicate_metadata_field", RMS_METADATA_SHEET, row_number, field, f"元数据字段重复: {field}"))
            continue
        value_cell = cells[2] if len(cells) > 2 else None
        if value_cell is not None and value_cell.data_type == "f":
            issues.append(_issue("formula_not_supported", RMS_METADATA_SHEET, row_number, field, "不支持公式单元格"))
            metadata[field] = None
        else:
            metadata[field] = value_cell.value if value_cell is not None else None
    metadata.setdefault("schemaVersion", "")
    return metadata, issues


def _parse_result_sheet(sheet: Any) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows = list(sheet.iter_rows(values_only=False))
    if len(rows) < 5:
        raise RmsAllocationXlsxError(f"{RMS_RESULT_SHEET} sheet 缺少结果表头")
    headers = [str(cell.value or "").strip() for cell in rows[4]]
    expected = [label for _, label in RESULT_COLUMNS]
    if headers[:len(expected)] != expected:
        raise RmsAllocationXlsxError(f"{RMS_RESULT_SHEET} sheet 第 5 行表头与 {RMS_XLSX_SCHEMA_VERSION} 不匹配")

    parsed_rows: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for row_number, cells in enumerate(rows[5:], start=6):
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        item: dict[str, Any] = {}
        for column_index, (field, _) in enumerate(RESULT_COLUMNS, start=1):
            cell = cells[column_index - 1] if column_index <= len(cells) else None
            if cell is not None and cell.data_type == "f":
                issues.append(_issue("formula_not_supported", RMS_RESULT_SHEET, row_number, field, "不支持公式单元格", column_index))
                item[field] = None
            else:
                item[field] = cell.value if cell is not None else None
        parsed_rows.append(item)
    return parsed_rows, issues


def _validate_metadata(metadata: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if str(metadata.get("schemaVersion") or "").strip() != RMS_XLSX_SCHEMA_VERSION:
        issues.append(_issue("unsupported_schema_version", RMS_METADATA_SHEET, 2, "schemaVersion", f"格式版本必须为 {RMS_XLSX_SCHEMA_VERSION}"))
    metadata_rows = {field: index + 2 for index, (field, _, _) in enumerate(METADATA_FIELDS)}
    for field in _REQUIRED_METADATA:
        value = metadata.get(field)
        if value is None or (isinstance(value, str) and not value.strip()):
            issues.append(_issue("missing_required_metadata", RMS_METADATA_SHEET, metadata_rows[field], field, f"缺少必填元数据: {field}"))
    plan_version = _finite_number(metadata.get("planVersion"))
    if metadata.get("planVersion") not in (None, "") and (plan_version is None or plan_version < 1 or not plan_version.is_integer()):
        issues.append(_issue("invalid_plan_version", RMS_METADATA_SHEET, metadata_rows["planVersion"], "planVersion", "planVersion 必须为大于等于 1 的整数"))
    elif plan_version is not None:
        metadata["planVersion"] = int(plan_version)
    numeric_metadata = {
        "missionHours": (0.0, None, False),
        "missionReliability": (0.0, 1.0, False),
        "mttrHours": (0.0, None, True),
    }
    for field, (minimum, maximum, inclusive_minimum) in numeric_metadata.items():
        value = _finite_number(metadata.get(field))
        if value is None:
            continue
        above_minimum = value >= minimum if inclusive_minimum else value > minimum
        below_maximum = True if maximum is None else value < maximum
        if not above_minimum or not below_maximum:
            interval = f"({minimum:g}, {maximum:g})" if maximum is not None else (f">= {minimum:g}" if inclusive_minimum else f"> {minimum:g}")
            issues.append(_issue("metadata_number_out_of_range", RMS_METADATA_SHEET, metadata_rows[field], field, f"{field} 必须在 {interval} 范围内"))
        else:
            metadata[field] = value
    return issues


def _validate_result_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    seen_node_ids: set[str] = set()
    if not rows:
        return [_issue("missing_result_rows", RMS_RESULT_SHEET, 6, "nodeId", "至少需要一行 RMS 分配结果")]
    numeric_rules = {
        "runningRatio": (0.0, 1.0, True),
        "failureRate": (0.0, None, True),
        "mtbfHours": (0.0, None, False),
        "mttrHours": (0.0, None, True),
        "installationCount": (1.0, None, True),
        "allocationShare": (0.0, 1.0, True),
        "localAllocationShare": (0.0, 1.0, True),
        "cumulativeAllocationShare": (0.0, 1.0, True),
        "cumulativeInstallationCount": (1.0, None, True),
    }
    for index, row in enumerate(rows):
        row_number = index + 6
        node_id = str(row.get("nodeId") or "").strip()
        row["nodeId"] = node_id
        row["parentNodeId"] = str(row.get("parentNodeId") or "").strip()
        if not node_id:
            issues.append(_issue("missing_node_id", RMS_RESULT_SHEET, row_number, "nodeId", "节点ID不能为空", 7))
        elif node_id in seen_node_ids:
            issues.append(_issue("duplicate_node_id", RMS_RESULT_SHEET, row_number, "nodeId", f"节点ID重复: {node_id}", 7))
        else:
            seen_node_ids.add(node_id)
        if not str(row.get("nodeName") or "").strip():
            issues.append(_issue("missing_node_name", RMS_RESULT_SHEET, row_number, "nodeName", "节点名称不能为空", 2))
        inactive = _result_row_is_inactive(row)
        optional_numeric_fields = {
            "localAllocationShare", "cumulativeAllocationShare", "cumulativeInstallationCount"
        }
        for field, (minimum, maximum, inclusive_minimum) in numeric_rules.items():
            value = _finite_number(row.get(field))
            column = next(i for i, (candidate, _) in enumerate(RESULT_COLUMNS, start=1) if candidate == field)
            if value is None:
                if field in optional_numeric_fields or (inactive and field in {"mtbfHours", "mttrHours"}):
                    row[field] = None
                    continue
                issues.append(_issue("invalid_number", RMS_RESULT_SHEET, row_number, field, f"{field} 必须为有限数值", column))
                continue
            if (value < minimum if inclusive_minimum else value <= minimum) or (maximum is not None and value > maximum):
                interval = f"[{minimum:g}, {maximum:g}]" if maximum is not None else (f">= {minimum:g}" if inclusive_minimum else f"> {minimum:g}")
                issues.append(_issue("number_out_of_range", RMS_RESULT_SHEET, row_number, field, f"{field} 必须在 {interval} 范围内", column))
            else:
                row[field] = int(value) if field == "installationCount" and value.is_integer() else value
        for field in ("installationCount", "cumulativeInstallationCount"):
            count = _finite_number(row.get(field))
            if count is not None and not count.is_integer():
                column = next(i for i, (candidate, _) in enumerate(RESULT_COLUMNS, start=1) if candidate == field)
                issues.append(_issue("invalid_installation_count", RMS_RESULT_SHEET, row_number, field, f"{field} 必须为整数", column))
        failure_rate = _finite_number(row.get("failureRate"))
        mtbf_hours = _finite_number(row.get("mtbfHours"))
        if not inactive and failure_rate is not None and mtbf_hours is not None and mtbf_hours > 0:
            expected_failure_rate = 1 / mtbf_hours
            tolerance = max(1e-9, expected_failure_rate * 1e-6)
            if abs(failure_rate - expected_failure_rate) > tolerance:
                issues.append(_issue(
                    "failure_rate_mtbf_mismatch",
                    RMS_RESULT_SHEET,
                    row_number,
                    "failureRate",
                    "失效率必须等于 1 / MTBF(h)",
                    4,
                ))
    return issues


def _result_cell_value(field: str, value: Any) -> Any:
    if field in {
        "runningRatio", "failureRate", "mtbfHours", "mttrHours", "installationCount",
        "allocationShare", "localAllocationShare", "cumulativeAllocationShare",
        "cumulativeInstallationCount",
    }:
        if value in (None, ""):
            return None
        number = _finite_number(value)
        return number if number is not None else _safe_cell_value(value)
    return _safe_cell_value(value)


def _result_source_value(row: dict[str, Any], field: str) -> Any:
    if field == "parentNodeId":
        return row.get("parentNodeId") if "parentNodeId" in row else row.get("parentId")
    if field == "localAllocationShare" and field not in row:
        return row.get("allocationShare")
    return row.get(field)


def _result_row_is_inactive(row: dict[str, Any]) -> bool:
    if _finite_number(row.get("runningRatio")) == 0:
        return True
    status = str(row.get("status") or "").strip().casefold()
    return status in {"未参与", "inactive", "not-participating", "not_participating"}


def _safe_cell_value(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return value
    text = _ILLEGAL_XML_CHARACTERS.sub("", str(value)).strip()[:100_000]
    return f"'{text}" if text.startswith(_FORMULA_PREFIXES) else text


def _finite_number(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _first_present(source: dict[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        if key in source:
            return source[key]
    return None


def _issue(code: str, sheet: str, row: int, field: str, message: str, column: int = 3) -> dict[str, Any]:
    return {"code": code, "sheet": sheet, "row": row, "column": column, "field": field, "message": message}
